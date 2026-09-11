"""DOCX (and other OOXML) document-context packaging.

Extracts page text plus the one structural fact its allowlisted operation needs: images
still missing alt text (WCAG 1.1.1), via `formats.office.images.undescribed_images` —
the exact same pure walk the production 1.1.1 detector and `apply_alt.apply_alt_text`
share, so a locator minted here is guaranteed to resolve at application time.
"""
from __future__ import annotations

import zipfile
from dataclasses import dataclass, field
from io import BytesIO

import docx

from experiments.document_wide_ai.application.production_adapters import office_undescribed_images
from experiments.document_wide_ai.contracts.v1 import ExtractionIssue, sha256_hex

EXTRACTOR_VERSION = "docx-extractor.v1"


@dataclass(frozen=True)
class DocxImage:
    locator: str  # "{part}#{docPr name | r:embed id}"
    part_name: str
    name: str


@dataclass(frozen=True)
class PackagedDocx:
    extractor_version: str
    paragraph_count: int
    text_context: str
    undescribed_images: tuple[DocxImage, ...]
    extraction_issues: tuple[ExtractionIssue, ...] = field(default_factory=tuple)

    def image_by_locator(self, locator: str) -> DocxImage | None:
        for img in self.undescribed_images:
            if img.locator == locator:
                return img
        return None


def fingerprint_missing() -> str:
    """Every undescribed image starts from the same precondition: no alt text yet."""
    return sha256_hex(b"<missing>")


def package_docx(source_bytes: bytes, *, max_text_chars: int) -> PackagedDocx:
    issues: list[ExtractionIssue] = []
    text_parts: list[str] = []
    paragraph_count = 0

    try:
        d = docx.Document(BytesIO(source_bytes))
        for p in d.paragraphs:
            if p.text:
                text_parts.append(p.text)
        paragraph_count = len(d.paragraphs)
        for table in d.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text:
                        text_parts.append(cell.text)
    except Exception as exc:
        issues.append(ExtractionIssue(kind="extraction_failed", detail=f"python-docx open failed: {exc}"))

    text_context = "\n".join(text_parts)
    if len(text_context) > max_text_chars:
        text_context = text_context[:max_text_chars]
        issues.append(
            ExtractionIssue(kind="extraction_truncated", detail=f"text truncated to {max_text_chars} chars")
        )

    images: list[DocxImage] = []
    try:
        with zipfile.ZipFile(BytesIO(source_bytes)) as zf:
            entries = {n: zf.read(n) for n in zf.namelist()}
        for rec in office_undescribed_images(entries):
            images.append(DocxImage(locator=rec["locator"], part_name=rec["part"], name=rec["name"]))
    except Exception as exc:
        issues.append(ExtractionIssue(kind="extraction_failed", detail=f"image extraction failed: {exc}"))

    return PackagedDocx(
        extractor_version=EXTRACTOR_VERSION,
        paragraph_count=paragraph_count,
        text_context=text_context,
        undescribed_images=tuple(images),
        extraction_issues=tuple(issues),
    )
