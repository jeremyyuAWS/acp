"""Apply validator-approved edits to a FRESH copy of the source bytes.

Never touches the input bytes. Deterministic ordering (by edit_id). An application
error or a detected regression (page/paragraph count changed, the saved file no longer
opens) rejects the WHOLE candidate — no partial candidate is ever returned as if it
were successful (PRD §5).
"""
from __future__ import annotations

import zipfile
from dataclasses import dataclass
from io import BytesIO

import docx
import pikepdf
import pypdf

from experiments.document_wide_ai.application.allowlist import (
    SET_OFFICE_IMAGE_ALT_TEXT,
    SET_PDF_FIELD_ACCESSIBLE_NAME,
    SET_PDF_FIGURE_ALT_TEXT,
)
from experiments.document_wide_ai.application.production_adapters import (
    apply_office_alt_text,
    apply_pdf_field_name,
    apply_pdf_figure_alt,
)
from experiments.document_wide_ai.contracts.v1 import DocumentFormat, ProposedEdit


@dataclass(frozen=True)
class AppliedEditRecord:
    edit_id: str
    locator: str
    before: str
    after: str


@dataclass(frozen=True)
class ApplicationResult:
    candidate_bytes: bytes | None
    applied: tuple[AppliedEditRecord, ...]
    not_applied: tuple[tuple[str, str], ...]  # (edit_id, reason)
    candidate_rejected_reason: str | None

    @property
    def candidate_rejected(self) -> bool:
        return self.candidate_rejected_reason is not None


def _locator_str(edit: ProposedEdit) -> str:
    loc = edit.locator
    if loc.format == DocumentFormat.PDF:
        return loc.element_ref
    return f"{loc.part_name}#{loc.element_ref}"


def apply_edits(
    source_bytes: bytes, document_format: DocumentFormat, edits: list[ProposedEdit]
) -> ApplicationResult:
    ordered = sorted(edits, key=lambda e: e.edit_id)
    if document_format == DocumentFormat.PDF:
        return _apply_pdf(source_bytes, ordered)
    if document_format == DocumentFormat.DOCX:
        return _apply_docx(source_bytes, ordered)
    return ApplicationResult(
        candidate_bytes=None,
        applied=(),
        not_applied=tuple((e.edit_id, "unsupported_format") for e in ordered),
        candidate_rejected_reason=f"unsupported format: {document_format}",
    )


def _apply_pdf(source_bytes: bytes, edits: list[ProposedEdit]) -> ApplicationResult:
    adapters = {SET_PDF_FIELD_ACCESSIBLE_NAME: apply_pdf_field_name,
                SET_PDF_FIGURE_ALT_TEXT: apply_pdf_figure_alt}
    supported = [e for e in edits if e.operation in adapters]
    unsupported = [e for e in edits if e.operation not in adapters]

    try:
        before_pages = len(pypdf.PdfReader(BytesIO(source_bytes)).pages)
    except Exception as exc:
        return ApplicationResult(
            candidate_bytes=None,
            applied=(),
            not_applied=tuple((e.edit_id, "source_unreadable") for e in edits),
            candidate_rejected_reason=f"source PDF failed to open before applying: {exc}",
        )

    values = {_locator_str(e): e.proposed_value for e in supported}
    edit_by_locator = {_locator_str(e): e for e in supported}

    try:
        new_bytes, applied_rows, unresolved_locators = source_bytes, [], []
        for operation, adapter in adapters.items():
            operation_values = {_locator_str(e): e.proposed_value for e in supported if e.operation == operation}
            if operation_values:
                new_bytes, rows, unresolved = adapter(new_bytes, operation_values)
                applied_rows.extend(rows)
                unresolved_locators.extend(unresolved)
    except Exception as exc:
        return ApplicationResult(
            candidate_bytes=None,
            applied=(),
            not_applied=tuple((e.edit_id, "adapter_raised") for e in edits),
            candidate_rejected_reason=f"PDF adapter raised: {exc}",
        )

    applied = tuple(
        AppliedEditRecord(
            edit_id=edit_by_locator[row["locator"]].edit_id,
            locator=row["locator"],
            before=row["before"],
            after=row["after"],
        )
        for row in applied_rows
        if row["locator"] in edit_by_locator
    )
    applied_locators = {r.locator for r in applied}
    not_applied = [(edit_by_locator[loc].edit_id, "locator_not_resolved") for loc in unresolved_locators if loc in edit_by_locator]
    not_applied += [(e.edit_id, "unsupported_operation") for e in unsupported]
    not_applied += [
        (e.edit_id, "not_applied")
        for e in supported
        if _locator_str(e) not in applied_locators and e.edit_id not in {i for i, _ in not_applied}
    ]

    try:
        after_pages = len(pypdf.PdfReader(BytesIO(new_bytes)).pages)
    except Exception as exc:
        return ApplicationResult(
            candidate_bytes=None,
            applied=(),
            not_applied=tuple((e.edit_id, "regression") for e in edits),
            candidate_rejected_reason=f"candidate PDF failed to reopen after applying: {exc}",
        )
    if after_pages != before_pages:
        return ApplicationResult(
            candidate_bytes=None,
            applied=(),
            not_applied=tuple((e.edit_id, "regression") for e in edits),
            candidate_rejected_reason=f"page count changed: {before_pages} -> {after_pages}",
        )

    return ApplicationResult(
        candidate_bytes=new_bytes, applied=applied, not_applied=tuple(not_applied), candidate_rejected_reason=None
    )


def _apply_docx(source_bytes: bytes, edits: list[ProposedEdit]) -> ApplicationResult:
    supported = [e for e in edits if e.operation == SET_OFFICE_IMAGE_ALT_TEXT]
    unsupported = [e for e in edits if e.operation != SET_OFFICE_IMAGE_ALT_TEXT]

    try:
        before_parts = set(zipfile.ZipFile(BytesIO(source_bytes)).namelist())
        before_paragraphs = len(docx.Document(BytesIO(source_bytes)).paragraphs)
    except Exception as exc:
        return ApplicationResult(
            candidate_bytes=None,
            applied=(),
            not_applied=tuple((e.edit_id, "source_unreadable") for e in edits),
            candidate_rejected_reason=f"source DOCX failed to open before applying: {exc}",
        )

    values = {_locator_str(e): e.proposed_value for e in supported}
    edit_by_locator = {_locator_str(e): e for e in supported}

    try:
        new_bytes, applied_rows, unresolved_locators = apply_office_alt_text(source_bytes, values)
    except Exception as exc:
        return ApplicationResult(
            candidate_bytes=None,
            applied=(),
            not_applied=tuple((e.edit_id, "adapter_raised") for e in edits),
            candidate_rejected_reason=f"apply_alt_text raised: {exc}",
        )

    applied = tuple(
        AppliedEditRecord(
            edit_id=edit_by_locator[row["locator"]].edit_id,
            locator=row["locator"],
            before=row["before"],
            after=row["after"],
        )
        for row in applied_rows
        if row["locator"] in edit_by_locator
    )
    applied_locators = {r.locator for r in applied}
    not_applied = [(edit_by_locator[loc].edit_id, "locator_not_resolved") for loc in unresolved_locators if loc in edit_by_locator]
    not_applied += [(e.edit_id, "unsupported_operation") for e in unsupported]
    already = {i for i, _ in not_applied}
    not_applied += [
        (e.edit_id, "not_applied") for e in supported if _locator_str(e) not in applied_locators and e.edit_id not in already
    ]

    try:
        after_parts = set(zipfile.ZipFile(BytesIO(new_bytes)).namelist())
        after_paragraphs = len(docx.Document(BytesIO(new_bytes)).paragraphs)
    except Exception as exc:
        return ApplicationResult(
            candidate_bytes=None,
            applied=(),
            not_applied=tuple((e.edit_id, "regression") for e in edits),
            candidate_rejected_reason=f"candidate DOCX failed to reopen after applying: {exc}",
        )
    if after_parts != before_parts:
        return ApplicationResult(
            candidate_bytes=None,
            applied=(),
            not_applied=tuple((e.edit_id, "regression") for e in edits),
            candidate_rejected_reason="package part set changed",
        )
    if after_paragraphs != before_paragraphs:
        return ApplicationResult(
            candidate_bytes=None,
            applied=(),
            not_applied=tuple((e.edit_id, "regression") for e in edits),
            candidate_rejected_reason=f"paragraph count changed: {before_paragraphs} -> {after_paragraphs}",
        )

    return ApplicationResult(
        candidate_bytes=new_bytes, applied=applied, not_applied=tuple(not_applied), candidate_rejected_reason=None
    )
