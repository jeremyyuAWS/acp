"""Saved real Office packages must contain exact approved alt and no collateral edits."""
from io import BytesIO
import zipfile
import pytest
from PIL import Image
from apply_alt import apply_alt_text
from office_alt_integrity import verify_alt_write


def package(ext):
    image = BytesIO()
    Image.new('RGB', (30, 20), 'blue').save(image, format='PNG')
    image.seek(0)
    output = BytesIO()
    if ext == 'docx':
        from docx import Document
        document = Document()
        document.add_paragraph('Keep this content')
        document.add_picture(image)
        document.save(output)
        locator = 'word/document.xml#Picture 1'
    elif ext == 'pptx':
        from pptx import Presentation
        from pptx.util import Inches
        document = Presentation()
        slide = document.slides.add_slide(document.slide_layouts[6])
        slide.shapes.add_picture(image, Inches(1), Inches(1))
        document.save(output)
        locator = 'ppt/slides/slide1.xml#Picture 1'
    else:
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image as Drawing
        document = Workbook()
        document.active['A1'] = 'Keep this content'
        document.active['B1'] = '=SUM(1,2)'
        document.active.add_image(Drawing(image), 'D4')
        document.save(output)
        locator = 'xl/drawings/drawing1.xml#Image 1'
    return output.getvalue(), locator


def mutate(data, change):
    with zipfile.ZipFile(BytesIO(data)) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    change(parts)
    output = BytesIO()
    with zipfile.ZipFile(output, 'w') as archive:
        for name, content in parts.items():
            archive.writestr(name, content)
    return output.getvalue()


@pytest.mark.parametrize('ext', ['docx', 'pptx', 'xlsx'])
def test_exact_readback_and_real_package_preservation(ext):
    before, locator = package(ext)
    values = {locator: 'Approved blue rectangle & "label"'}
    after, applied, unresolved = apply_alt_text(before, values)
    assert len(applied) == 1 and not unresolved
    assert verify_alt_write(before, after, values)
    assert not verify_alt_write(before, before, values)
    assert not verify_alt_write(before, after, {locator: 'Wrong value'})
    assert not verify_alt_write(before, after, {locator + ' missing': values[locator]})
    assert not verify_alt_write(before, after, {})
    assert not verify_alt_write(before, b'broken zip', values)

    media = next(name for name in zipfile.ZipFile(BytesIO(after)).namelist() if '/media/' in name)
    damaged = mutate(after, lambda parts: parts.__setitem__(media, b'lost image'))
    assert not verify_alt_write(before, damaged, values)
    part = locator.split('#')[0]
    # Valid XML and still a readable archive; layout changes are nevertheless rejected.
    damaged = mutate(after, lambda parts: parts.__setitem__(part, parts[part].replace(b'cx="', b'cx="9', 1)))
    assert damaged != after
    assert not verify_alt_write(before, damaged, values)


def test_workbook_values_and_formula_cannot_disappear():
    before, locator = package('xlsx')
    values = {locator: 'Blue rectangle'}
    after, _, _ = apply_alt_text(before, values)
    damaged = mutate(after, lambda parts: parts.__setitem__('xl/worksheets/sheet1.xml',
        parts['xl/worksheets/sheet1.xml'].replace(b'SUM(1,2)', b'SUM(100,2)')))
    assert not verify_alt_write(before, damaged, values)


def test_duplicate_package_names_rejected():
    before, locator = package('xlsx')
    after, _, _ = apply_alt_text(before, {locator: 'Blue rectangle'})
    output = BytesIO(after)
    with pytest.warns(UserWarning):
        with zipfile.ZipFile(output, 'a') as archive:
            archive.writestr('xl/worksheets/sheet1.xml', '<worksheet/>')
    assert not verify_alt_write(before, output.getvalue(), {locator: 'Blue rectangle'})


# Exercise the live approved-write lane, not just the comparison helper.
from test_apply_approved_values import store, _seed, _Blob, _run_handler, SID, FILE


def test_failed_integrity_keeps_previous_copy_and_withholds_credit(store, monkeypatch):
    import apply_alt
    before, locator = package('pptx')
    item = _seed(store, names=('Picture 1',))
    store.update_hitl_item(item, 'approved', None, None)
    store.approve_proposal_values(item, ['Approved blue rectangle'])
    original_writer = apply_alt.apply_alt_text
    def damaged_writer(data, values, **kwargs):
        after, applied, unresolved = original_writer(data, values, **kwargs)
        after = mutate(after, lambda parts: parts.__setitem__('docProps/core.xml', b'<lost/>'))
        return after, applied, unresolved
    monkeypatch.setattr(apply_alt, 'apply_alt_text', damaged_writer)
    blob = _Blob(before)
    _run_handler(monkeypatch, store, blob, residual=set())
    assert blob.data == before
    assert not blob.uploads
    assert store.count_unapplied_approved_values(SID, FILE) == 1
