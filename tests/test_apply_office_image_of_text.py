"""The docx/xlsx half of ADR 0055's locator translation, and WHY it is not the pptx one.

tests/test_remediation_described_image_office_round_trip.py proves the lane end to end. This
file pins the three ways the pptx translation fails on these formats — each measured on a real
package rather than reasoned about from the pptx code — plus the invariant the new resolver is
built on.

None of this needs OCR: the translation is pure package structure, so these run in milliseconds
and stay useful on a machine where tesseract is missing and the round trip skips.
"""
from __future__ import annotations

import functools
import io
import re
import sys
import tempfile
import zipfile
from pathlib import Path

import pytest

ACP = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ACP / "api"))

pytest.importorskip("docx")
pytest.importorskip("openpyxl")
pytest.importorskip("PIL")

ALT_A = "A benefits summary card listing medical, dental and vision cover."
ALT_B = "A payroll notice card about Monday timesheets."


@functools.lru_cache(maxsize=None)
def _img(colour: str = "white") -> Path:
    """Any raster under a media/ prefix. The translation never opens the image — what it must get
    right is which PART references it — so this deliberately carries no text at all."""
    from PIL import Image
    p = Path(tempfile.mkdtemp()) / f"{colour}.png"
    Image.new("RGB", (240, 160), colour).save(p)
    return p


@functools.lru_cache(maxsize=None)
def _docx_twice() -> bytes:
    """One image, placed twice — ONE media part, ONE relationship id, TWO <wp:docPr>."""
    import docx
    from docx.shared import Inches
    doc = docx.Document()
    doc.add_paragraph("intro")
    doc.add_picture(str(_img()), width=Inches(2))
    doc.add_paragraph("outro")
    doc.add_picture(str(_img()), width=Inches(2))
    out = Path(tempfile.mkdtemp()) / "d.docx"
    doc.save(out)
    return out.read_bytes()


@functools.lru_cache(maxsize=None)
def _xlsx_two_sheets() -> bytes:
    """Two images, one per sheet — two media parts, two DRAWING parts, absolute rel targets."""
    import openpyxl
    from openpyxl.drawing.image import Image as XLImage
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "One"
    ws1.add_image(XLImage(str(_img("white"))), "B2")
    ws2 = wb.create_sheet("Two")
    ws2.add_image(XLImage(str(_img("ivory"))), "B2")
    out = Path(tempfile.mkdtemp()) / "x.xlsx"
    wb.save(out)
    return out.read_bytes()


@functools.lru_cache(maxsize=None)
def _xlsx_one_image_two_drawings() -> bytes:
    """ONE media part referenced by TWO drawing parts — what Excel writes for the same picture
    on two sheets, and what openpyxl will not produce (it duplicates the bytes instead).

    Built by rewriting the second drawing's relationship to point at the first image and dropping
    the now-orphaned media part. Surgery on the rels only: every other part is copied verbatim,
    so this is the real package shape and not a mock of it.
    """
    src = _xlsx_two_sheets()
    buf = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(src)) as zin, \
         zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            if info.filename == "xl/media/image2.png":
                continue
            data = zin.read(info.filename)
            if info.filename == "xl/drawings/_rels/drawing2.xml.rels":
                data = data.replace(b"/xl/media/image2.png", b"/xl/media/image1.png")
            zout.writestr(info, data)
    return buf.getvalue()


def _rels(data: bytes, part: str) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return z.read(part).decode("utf-8")


def _media(data: bytes) -> list[str]:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return [n for n in z.namelist() if "/media/" in n]


# ------------------------------------------------- the pptx translation does not transfer

def test_the_pptx_resolver_finds_nothing_in_a_docx_or_an_xlsx():
    """It enumerates ppt/media/ alone, so on these packages it is not wrong — it is silent.

    Silence is the failure mode that matters: the caller keeps the untranslated 'image N', apply
    _alt reports it unresolved, the described row is never marked applied, and the file can never
    certify. Nothing raises and nothing logs a mismatch.
    """
    import apply_pptx_image_of_text as pptx_aoit
    assert pptx_aoit.resolve_media_locators(_docx_twice(), ["image 1"]) == {}
    assert pptx_aoit.resolve_media_locators(_xlsx_two_sheets(), ["image 1", "image 2"]) == {}


def test_an_xlsx_relationship_target_is_absolute():
    """Measured failure 2: openpyxl writes Target="/xl/media/image1.png".

    posixpath.join drops everything before an absolute component, so the pptx canonicaliser
    yields a leading-slash path — which matches no zip namelist entry, because namelist entries
    have no leading slash. The new canonicaliser strips it.
    """
    import posixpath
    import apply_office_image_of_text as aoit
    target = re.search(r'Target="([^"]+)"',
                       _rels(_xlsx_two_sheets(), "xl/drawings/_rels/drawing1.xml.rels")).group(1)
    assert target.startswith("/")
    assert posixpath.normpath(posixpath.join("xl/drawings", target)) not in _media(
        _xlsx_two_sheets())
    assert aoit._canonical("xl/drawings", target) in _media(_xlsx_two_sheets())


def test_an_xlsx_relationship_writes_its_attributes_in_the_other_order():
    """Measured failure 3, independent of failure 2 and with the identical symptom.

    apply_pptx_image_of_text._REL_RE requires Id="…" BEFORE Target="…"; openpyxl emits Type,
    Target, Id. So even with the leading slash fixed, that pattern matches nothing.
    """
    import apply_pptx_image_of_text as pptx_aoit
    rels = _rels(_xlsx_two_sheets(), "xl/drawings/_rels/drawing1.xml.rels")
    assert rels.index('Target="') < rels.index('Id="')
    assert not pptx_aoit._REL_RE.search(rels)


def test_a_docx_places_one_image_twice_behind_one_relationship_id():
    """Measured failure 1, at the package level: the fact that makes 'part#rId' insufficient.

    The exact strings are pinned as a STABILITY contract, not as the correctness argument, and a
    bite check is why that distinction is written down. Swapping the candidate order to try the
    rId first turns this assertion red and leaves the round trip GREEN — the resolver falls back
    to the name for the second placement, so both are still described. What is load-bearing is
    that a locator is emitted PER PLACEMENT and verified against apply_alt.resolve_target; which
    fragment shape wins is a preference (a name survives an rId being renumbered, and it is the
    shape the 1.1.1 detector already mints for these parts). The per-placement claim is bitten
    by test_the_pptx_shaped_locator_would_describe_only_one_placement in the round-trip module.
    """
    import apply_office_image_of_text as aoit
    data = _docx_twice()
    assert len(_media(data)) == 1
    rids = re.findall(r'r:embed="([^"]+)"', _rels(data, "word/document.xml"))
    assert len(rids) == 2 and len(set(rids)) == 1        # two placements, ONE rId
    # …so the translation must address the ELEMENTS, and it does.
    assert aoit.resolve_media_locators(data, ["image 1"], "docx") == {
        "image 1": ["word/document.xml#Picture 1", "word/document.xml#Picture 2"]}


# ------------------------------------------------- what the new resolver produces

def test_every_emitted_locator_resolves_to_a_distinct_picture():
    """The invariant the module is built on, checked through apply_alt's OWN resolver.

    A locator is only useful if the writer resolves it back to the picture it was minted for.
    Rather than trust that, the resolver checks each candidate against resolve_target and keeps
    only what resolves — so this asserts the property end to end: distinct locators, distinct
    offsets, no two placements collapsing onto one element.
    """
    import apply_alt
    import apply_office_image_of_text as aoit
    for data, ext, locs in ((_docx_twice(), "docx", ["image 1"]),
                            (_xlsx_two_sheets(), "xlsx", ["image 1", "image 2"]),
                            (_xlsx_one_image_two_drawings(), "xlsx", ["image 1"])):
        resolved = aoit.resolve_media_locators(data, locs, ext)
        assert resolved, f"{ext}: nothing resolved"
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            entries = {n: z.read(n) for n in z.namelist()}
        seen: set[tuple[str, int]] = set()
        for targets in resolved.values():
            for t in targets:
                part, fragment = apply_alt.parse_locator(t)
                tag = apply_alt.tag_for_part(part)
                assert tag, f"{t}: part carries no alt-bearing element"
                at = apply_alt.resolve_target(entries[part].decode("utf-8"), tag, fragment)
                assert at is not None, f"{t}: apply_alt cannot resolve the locator we minted"
                assert (part, at) not in seen, f"{t}: two locators name the same element"
                seen.add((part, at))


def test_one_media_part_in_two_drawings_expands_to_both():
    """ONE LOCATOR, MANY PLACEMENTS on a workbook — the pptx lane's lesson, reached through the
    xlsx indirection where the sheet is not the alt-bearing part.

    Describing only the first drawing would leave the second carrying openpyxl's 'Picture'
    placeholder, which the 1.1.1 detector reads as junk, so the criterion would still fail and
    the credit would be withheld for a write that was otherwise right.
    """
    import apply_office_image_of_text as aoit
    data = _xlsx_one_image_two_drawings()
    assert len(_media(data)) == 1
    out = aoit.expand_media_locator_values(data, {"image 1": ALT_A}, "xlsx")
    assert out == {"xl/drawings/drawing1.xml#Image 1": ALT_A,
                   "xl/drawings/drawing2.xml#Image 1": ALT_A}


def test_each_card_maps_to_its_own_drawing():
    """Two media parts must not both resolve to the first drawing — a mistake a one-image
    workbook could never expose, and one that would put a reviewer's description on a picture
    they never saw."""
    import apply_office_image_of_text as aoit
    out = aoit.expand_media_locator_values(
        _xlsx_two_sheets(), {"image 1": ALT_A, "image 2": ALT_B}, "xlsx")
    assert out == {"xl/drawings/drawing1.xml#Image 1": ALT_A,
                   "xl/drawings/drawing2.xml#Image 1": ALT_B}


# ------------------------------------------------- the things it must NOT do

def test_ordinary_alt_locators_pass_through_untouched():
    """An approved 1.1.1 row's 'part#name' must reach apply_alt exactly as the reviewer's card
    used it — the alt lane runs this expansion over ALL its values, not only described ones."""
    import apply_office_image_of_text as aoit
    plain = {"word/document.xml#Picture 1": ALT_A}
    assert aoit.expand_media_locator_values(_docx_twice(), dict(plain), "docx") == plain
    mixed = aoit.expand_media_locator_values(
        _docx_twice(), {"image 1": ALT_A, "word/document.xml#Picture 9": ALT_B}, "docx")
    assert mixed["word/document.xml#Picture 9"] == ALT_B


def test_an_unresolvable_media_locator_is_left_alone_rather_than_guessed_at():
    """Omitted from the map, kept in the values: it reaches apply_alt, is reported unresolved,
    and shows up in the apply.unresolved log under the name the reviewer's card used. Mapping it
    to some other picture would be a reviewer's signature on prose about an image they never saw.
    """
    import apply_office_image_of_text as aoit
    data = _docx_twice()
    assert aoit.resolve_media_locators(data, ["image 9"], "docx") == {}
    assert aoit.expand_media_locator_values(data, {"image 9": ALT_A}, "docx") == {"image 9": ALT_A}
    for junk in ("image", "image 0", "image -1", "", "  "):
        assert aoit.resolve_media_locators(data, [junk], "docx") == {}


def test_an_external_relationship_is_never_treated_as_a_media_part():
    """A linked image is a URL, not a part in the package — nobody can describe it, and a
    TargetMode="External" relationship must not shadow a real one."""
    import apply_office_image_of_text as aoit
    src = _docx_twice()
    buf = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(src)) as zin, \
         zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == "word/_rels/document.xml.rels":
                data = data.replace(
                    b"</Relationships>",
                    b'<Relationship Id="rIdExt" TargetMode="External" '
                    b'Target="http://example.com/word/media/image1.png"/></Relationships>')
            zout.writestr(info, data)
    assert aoit.resolve_media_locators(buf.getvalue(), ["image 1"], "docx") == {
        "image 1": ["word/document.xml#Picture 1", "word/document.xml#Picture 2"]}


def test_pptx_is_delegated_rather_than_reimplemented():
    """One implementation per format, and the proven pptx lane keeps the one it was merged with.

    Asserted by IDENTITY of the answer, on a package built the way the pptx round trip builds
    its own, so a future divergence in either module fails here rather than in production.
    """
    pptx = pytest.importorskip("pptx")
    import apply_office_image_of_text as aoit
    import apply_pptx_image_of_text as pptx_aoit
    from pptx.util import Inches
    prs = pptx.Presentation()
    for _ in range(2):
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.add_picture(str(_img()), Inches(1), Inches(2), Inches(2), Inches(1.4))
    out = Path(tempfile.mkdtemp()) / "p.pptx"
    prs.save(out)
    data = out.read_bytes()
    assert (aoit.resolve_media_locators(data, ["image 1"], "pptx")
            == pptx_aoit.resolve_media_locators(data, ["image 1"]))
    assert (aoit.expand_media_locator_values(data, {"image 1": ALT_A}, "pptx")
            == pptx_aoit.expand_media_locator_values(data, {"image 1": ALT_A}))


def test_handlers_asks_this_module_for_every_office_format():
    """The wiring, structurally — the gap ADR 0055 left open was that handlers asked only pptx.

    Read from the source rather than by importing handlers, which drags in the scheduler and the
    DB driver (the same reason scripts/gen_capability_levels.py parses it).
    """
    import ast
    tree = ast.parse((ACP / "api" / "handlers.py").read_text())
    names = {a.name for n in ast.walk(tree) if isinstance(n, ast.ImportFrom)
             and n.module == "apply_office_image_of_text" for a in n.names}
    assert {"SUPPORTED_EXTS", "is_media_index_locator",
            "expand_media_locator_values"} <= names, (
        "handlers no longer routes the described-image locator translation through "
        "apply_office_image_of_text — docx and xlsx descriptions would reach nothing")

    import apply_office_image_of_text as aoit
    assert set(aoit.SUPPORTED_EXTS) == {"docx", "xlsx", "pptx"}
