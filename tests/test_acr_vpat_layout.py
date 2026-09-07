"""The WCAG report laid out the way the ITI template lays it out — acceptance row 13.

The template splits WCAG by conformance level: one table per level, three columns each, because
the level is the table you are in rather than a column in it. ACP printed one four-column table
with a Level column, under a heading (`WCAG 2.2 Report`) that matched no edition of the template
at all.

What this file pins is the layout AND the two things the layout must not quietly change: every
criterion still prints, and nothing here claims conformance to a level that was not evaluated.
"""
import io

import pytest

import acr_catalog
import acr_export_preview


def _projection(edition="VPAT 2.5Rev WCAG", **report):
    rows = acr_catalog.build_matrix("r-1", edition)
    return acr_export_preview.project(
        {"vpat_edition": edition, "wcag_version": "2.2", "report_title": "ACP ACR", **report},
        rows)


# ── the heading, which is the part that was wrong in all four editions at once ────────────────

def test_the_508_edition_gets_wcag_2_0_and_the_others_get_2_x():
    """The bug this slice fixes, stated as the difference that made it invisible.

    ACP built `f"WCAG {report.wcag_version} Report"` — so every edition got `WCAG 2.2 Report`,
    which is not what any of the four templates say. Being wrong IDENTICALLY everywhere is why no
    comparison between editions could surface it; the templates disagree and ACP did not.
    """
    assert _projection("VPAT 2.5Rev 508")["wcag"]["heading"] == "WCAG 2.0 Report"
    for edition in ("VPAT 2.5Rev WCAG", "VPAT 2.5Rev EU", "VPAT 2.5Rev INT"):
        assert _projection(edition)["wcag"]["heading"] == "WCAG 2.x Report", edition


def test_an_unselected_edition_falls_back_to_the_templates_own_generic_wording():
    """Not to `WCAG 2.2 Report`. `2.x` is how the template spans WCAG versions, so it stays true
    for a 2.1 report where a version-stamped heading would quietly assert 2.2."""
    projection = acr_export_preview.project(
        {"wcag_version": "2.1"}, acr_catalog.build_matrix("r-1", "VPAT 2.5Rev WCAG"))
    assert projection["wcag"]["heading"] == "WCAG 2.x Report"


# ── the per-level tables ──────────────────────────────────────────────────────────────────────

def test_the_levels_are_the_templates_three_tables_in_order():
    levels = _projection()["wcag"]["levels"]
    assert [lv["heading"] for lv in levels] == [
        "Table 1: Success Criteria, Level A",
        "Table 2: Success Criteria, Level AA",
        "Table 3: Success Criteria, Level AAA",
    ]
    assert [lv["level"] for lv in levels] == ["A", "AA", "AAA"]


def test_every_criterion_still_prints_exactly_once():
    """The property a regrouping is most likely to break, and the one that matters most.

    A criterion that vanishes from a conformance report because the layout had nowhere to put it
    is worse than an ugly layout: it reads as a standard with fewer requirements than it has.
    """
    projection = _projection()
    grouped = [r["criterion_num"] for lv in projection["wcag"]["levels"] for r in lv["rows"]]
    assert sorted(grouped) == sorted(r["criterion_num"] for r in projection["criteria"])
    assert len(grouped) == len(set(grouped)) == 55


def test_a_row_whose_level_is_outside_the_template_still_prints():
    """No silent drop. ACP's catalog is A/AA today, so the escape group is always empty — which
    is exactly when a guard earns its place, because nothing else would notice it breaking."""
    rows = acr_catalog.build_matrix("r-1", "VPAT 2.5Rev WCAG")
    rows[0] = {**rows[0], "level": "AAA-ish"}
    projection = acr_export_preview.project({"vpat_edition": "VPAT 2.5Rev WCAG"}, rows)
    printed = [r["criterion_num"] for lv in projection["wcag"]["levels"] for r in lv["rows"]]
    assert rows[0]["criterion_num"] in printed
    assert len(printed) == 55
    assert projection["wcag"]["levels"][-1]["level"] is None


def test_the_flat_criteria_list_is_still_there_for_published_snapshots():
    """PRD §17: a published snapshot is immutable, so anything rendered against `criteria` has to
    keep rendering. The grouping is additive — a new key beside it, not a replacement for it."""
    projection = _projection()
    assert len(projection["criteria"]) == 55
    assert projection["criteria"][0]["level"] in ("A", "AA")


# ── the scope statement, which is a §19 constraint and not a formatting choice ────────────────

def test_level_aaa_is_declared_unevaluated_rather_than_omitted():
    """An omitted table is a silent scope decision and reads as "nothing to report". PRD §19
    forbids a document implying a claim about a standard it did not evaluate, in either
    direction, so the absence is stated in words."""
    aaa = _projection()["wcag"]["levels"][2]
    assert aaa["rows"] == []
    assert "not evaluated" in aaa["not_evaluated"]
    assert "no conformance to them is claimed" in aaa["not_evaluated"]


def test_a_level_with_rows_carries_no_unevaluated_notice():
    for level in _projection()["wcag"]["levels"][:2]:
        assert level["rows"]
        assert "not_evaluated" not in level


# ── the rendered formats ──────────────────────────────────────────────────────────────────────

def test_the_html_drops_the_level_column_and_keeps_the_level_as_a_heading():
    html = acr_export_preview.to_html(_projection())
    assert '<th scope="col">Level</th>' not in html
    assert "<h2>WCAG 2.x Report</h2>" in html
    assert "<h3>Table 1: Success Criteria, Level A</h3>" in html
    assert html.count('<th scope="col">Criteria</th>') == 2, "Level A and Level AA"


def test_the_html_states_the_aaa_scope_without_an_empty_table():
    html = acr_export_preview.to_html(_projection())
    assert "<h3>Table 3: Success Criteria, Level AAA</h3>" in html
    assert "no conformance to them is claimed" in html


def test_the_word_document_matches_the_templates_heading_structure():
    pytest.importorskip("docx")
    import docx

    import acr_export_docx

    document = docx.Document(io.BytesIO(acr_export_docx.render(_projection("VPAT 2.5Rev INT"))))
    headings = [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]
    assert "WCAG 2.x Report" in headings
    assert "Table 1: Success Criteria, Level A" in headings
    assert "Table 2: Success Criteria, Level AA" in headings
    assert "Revised Section 508 Report" in headings
    assert "EN 301 549 Report" in headings
    # Order matters: the template prints WCAG, then 508, then EN.
    assert (headings.index("WCAG 2.x Report") < headings.index("Revised Section 508 Report")
            < headings.index("EN 301 549 Report"))


def test_the_word_wcag_tables_are_three_column_with_a_header_row():
    pytest.importorskip("docx")
    import docx

    import acr_export_docx

    document = docx.Document(io.BytesIO(acr_export_docx.render(_projection())))
    three_col = [t for t in document.tables if len(t.columns) == 3]
    assert len(three_col) == 2, "Level A and Level AA; AAA has no table"
    for table in three_col:
        assert [c.text for c in table.rows[0].cells] == [
            "Criteria", "Conformance Level", "Remarks and Explanations"]
    assert sum(len(t.rows) - 1 for t in three_col) == 55


def test_the_word_document_still_passes_acps_own_accessibility_gate():
    """The check that makes the relayout safe to ship. Splitting one table into two, under two
    new heading levels, is two more chances to emit a table with no header row or a heading level
    that skips — and row 14's gate is what refuses that document rather than serving it."""
    pytest.importorskip("docx")
    import acr_export_docx

    verdict = acr_export_docx.check(acr_export_docx.render(_projection()))
    assert verdict["ok"] is True, verdict["failures"]


# ── what a structure decision does not license ────────────────────────────────────────────────

def test_no_format_claims_to_be_a_vpat():
    """ADR 0053's Q1 is still open: the owner's 2026-09-07 decision was about REPRODUCTION, which
    is copyright, and the VPAT name is a service mark. Matching the layout must not drift into
    using the mark, and this is the test that fails if it ever does.
    """
    projection = _projection()
    assert projection["template"]["is_official_iti_template"] is False
    html = acr_export_preview.to_html(projection)
    assert "is not a VPAT" in html
    for heading in [lv["heading"] for lv in projection["wcag"]["levels"]]:
        assert "VPAT" not in heading
        assert "®" not in heading
    assert "VPAT" not in projection["wcag"]["heading"]
