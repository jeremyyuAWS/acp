import io
import zipfile
import pytest
from PIL import Image
from document_wide_office import build_office_manifest, targets
from document_wide_manifest import build_manifest, package_images
from experiments.document_wide_ai.packaging.limits import ExtractionLimits
from experiments.document_wide_ai.contracts.v1 import CONTRACT_VERSION, ProposedEdit, EditResponseEnvelope
from experiments.document_wide_ai.validation.validator import validate_edit_response
from apply_alt import apply_alt_text
from formats.office.images import undescribed_images
from test_document_wide_manifest import setup_manifest


def office_fixture(fmt):
    pixels = io.BytesIO()
    Image.new('RGB', (12, 12), 'red').save(pixels, format='PNG')
    pixels.seek(0)
    output = io.BytesIO()
    if fmt == 'pptx':
        from pptx import Presentation
        from pptx.util import Inches
        doc = Presentation()
        slide = doc.slides.add_slide(doc.slide_layouts[6])
        slide.shapes.add_picture(pixels, Inches(1), Inches(1))
        box = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(3), Inches(1))
        box.text = 'Sales context and quarterly performance'
        doc.save(output)
    else:
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image as SheetImage
        doc = Workbook()
        doc.active['A1'] = 'Sales context and quarterly performance'
        doc.active.add_image(SheetImage(pixels), 'B2')
        doc.save(output)
    return output.getvalue()


@pytest.mark.parametrize('fmt', ['pptx', 'xlsx'])
def test_real_office_context_evidence_identity_and_saved_writer(fmt, isolated_store, monkeypatch):
    data = office_fixture(fmt)
    exact = next(iter(targets(data)))
    aliases = targets(data)[exact][0]
    assessed = next(alias for alias in aliases if alias.startswith(fmt+':'))
    rows = setup_manifest(isolated_store, monkeypatch, data, 'sales.'+fmt, '1.1.1', [assessed])
    manifest = build_manifest(isolated_store, 'scan', 'sales.'+fmt, data)
    assert manifest.findings[0].finding_id == rows[0]['finding_id']
    assert 'Sales context and quarterly performance' in manifest.text_context
    assert len(package_images(data, manifest)) == 1
    finding = manifest.findings[0]
    edit = ProposedEdit('alt', (finding.finding_id,), finding.locator, 'set_office_image_alt_text',
        'Red square representing quarterly sales', '', 'Embedded image and full extracted context')
    result = validate_edit_response(manifest, EditResponseEnvelope(CONTRACT_VERSION, 'request', manifest.source_sha256, (edit,), ()))
    assert result.valid_edits == (edit,)
    corrected, applied, unresolved = apply_alt_text(data, {exact: edit.proposed_value})
    assert not unresolved and len(applied) == 1
    with zipfile.ZipFile(io.BytesIO(corrected)) as z:
        assert not undescribed_images({n:z.read(n) for n in z.namelist()})
    if fmt == 'pptx':
        from pptx import Presentation
        assert len(Presentation(io.BytesIO(corrected)).slides) == 1
    else:
        from openpyxl import load_workbook
        assert load_workbook(io.BytesIO(corrected)).active['A1'].value == 'Sales context and quarterly performance'


@pytest.mark.parametrize('fmt', ['pptx', 'xlsx'])
def test_selected_scope_and_extraction_limits(fmt):
    data = office_fixture(fmt)
    manifest = build_office_manifest(data, document_id='a.'+fmt, assessment_revision='rev', selected_criteria=('2.4.2',), limits=ExtractionLimits())
    assert not manifest.findings
    with pytest.raises(ValueError):
        build_office_manifest(data, document_id='a.'+fmt, assessment_revision='rev', selected_criteria=('1.1.1',), limits=ExtractionLimits(max_text_chars=2))


def test_pdf_language_manifest_uses_exact_assessed_target(isolated_store, monkeypatch):
    from test_pdf_structural_language import fixture, targets as language_targets
    data = fixture()
    target = language_targets(data)[0]
    rows = setup_manifest(isolated_store, monkeypatch, data, 'foreign.pdf', '3.1.2', [target.locator])
    manifest = build_manifest(isolated_store, 'scan', 'foreign.pdf', data)
    assert manifest.findings[0].finding_id == rows[0]['finding_id']
    assert manifest.findings[0].locator.element_ref == target.locator
    assert target.text in manifest.text_context
    edit = ProposedEdit('language', (rows[0]['finding_id'],), manifest.findings[0].locator,
        'set_pdf_structure_language', 'fr', '', 'Exact French passage')
    result = validate_edit_response(manifest, EditResponseEnvelope(CONTRACT_VERSION, 'request', manifest.source_sha256, (edit,), ()))
    assert result.valid_edits == (edit,)
