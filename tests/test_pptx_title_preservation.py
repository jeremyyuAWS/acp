"""Real saved PPTX title approvals must not replace body text or title formatting."""
import io
import zipfile

from lxml import etree
from pptx import Presentation
from pptx.util import Inches
import pytest

from apply_pptx_slide_titles import apply_pptx_slide_titles


def fixture(*, existing='', ambiguous=False, missing_body=False):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    title = slide.shapes.title
    title.text = existing
    title.text_frame.margin_left = Inches(.3)
    title.text_frame.word_wrap = True
    paragraph = title.text_frame.paragraphs[0]
    paragraph.font.name = 'Aptos Display'
    paragraph.font.bold = True
    paragraph.font.size = Inches(.35)
    body = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(7), Inches(1))
    body.text = 'Financial results must remain unchanged'
    # Shape order is legal and changes which shape a regex sees first.
    tree = slide.shapes._spTree
    tree.remove(body._element)
    tree.insert(2, body._element)
    if ambiguous:
        from copy import deepcopy
        tree.append(deepcopy(title._element))
    if missing_body:
        title._element.remove(title._element.txBody)
    stream = io.BytesIO()
    presentation.save(stream)
    return stream.getvalue()


def parts(data):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def test_title_after_body_fills_only_title_and_preserves_saved_layout():
    original = fixture()
    old = Presentation(io.BytesIO(original)).slides[0]
    before_body = etree.tostring(old.shapes[0]._element, method='c14n')
    before_properties = etree.tostring(old.shapes.title.text_frame._txBody.bodyPr, method='c14n')
    fixed, applied, unresolved = apply_pptx_slide_titles(original, {'slide 1': 'Q&A <2026>'})
    saved = Presentation(io.BytesIO(fixed)).slides[0]
    assert saved.shapes.title.text == 'Q&A <2026>'
    assert etree.tostring(saved.shapes[0]._element, method='c14n') == before_body
    assert etree.tostring(saved.shapes.title.text_frame._txBody.bodyPr, method='c14n') == before_properties
    paragraph = saved.shapes.title.text_frame.paragraphs[0]
    assert paragraph.font.name == 'Aptos Display'
    assert paragraph.font.bold is True
    assert applied == [{'locator': 'slide 1', 'before': '(empty title placeholder)', 'after': 'Q&A <2026>'}]
    assert unresolved == []
    old_parts, new_parts = parts(original), parts(fixed)
    assert old_parts.keys() == new_parts.keys()
    assert all(value == new_parts[name] for name, value in old_parts.items() if name != 'ppt/slides/slide1.xml')


@pytest.mark.parametrize('options', [{'existing': 'User added this title'}, {'ambiguous': True}, {'missing_body': True}])
def test_stale_or_ambiguous_or_unwritable_title_is_not_claimed_applied(options):
    original = fixture(**options)
    fixed, applied, unresolved = apply_pptx_slide_titles(original, {'slide 1': 'AI suggestion'})
    assert fixed == original
    assert applied == []
    assert unresolved == ['slide 1']


@pytest.mark.parametrize('value', ['', '   ', 'Invalid\x00XML'])
def test_empty_or_invalid_title_leaves_original_bytes(value):
    original = fixture()
    assert apply_pptx_slide_titles(original, {'slide 1': value}) == (original, [], ['slide 1'])
