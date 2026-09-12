"""Approved OCR replacements remove rasters and preserve real Office packages."""
import io
import zipfile
from pathlib import Path

import pytest
from lxml import etree as ET

from apply_office_image_replacement import apply_office_image_replacement, NS

PROVES_LANES = {('docx', '1.4.5'), ('xlsx', '1.4.5')}

TEXT = 'Benefits at a glance\nMedical dental and vision cover\nEnrollment closes on Friday'


def picture():
    from PIL import Image, ImageDraw, ImageFont
    im = Image.new('RGB', (900, 260), 'white')
    try:
        font = ImageFont.truetype('/System/Library/Fonts/Supplemental/Arial.ttf', 38)
    except OSError:
        font = ImageFont.truetype('DejaVuSans.ttf', 38)
    draw = ImageDraw.Draw(im)
    for i, line in enumerate(TEXT.splitlines()):
        draw.text((20, 15 + i * 75), line, font=font, fill='black')
    out = io.BytesIO()
    im.save(out, format='PNG')
    return out.getvalue()


def document(ext, shared=False):
    out = io.BytesIO()
    if ext == 'docx':
        from docx import Document
        doc = Document()
        doc.add_heading('Unrelated heading', 1)
        doc.add_paragraph('Unrelated body copy')
        doc.add_picture(io.BytesIO(picture()))
        if shared:
            doc.add_picture(io.BytesIO(picture()))
        doc.save(out)
    else:
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image
        wb = Workbook()
        wb.active['A1'] = 'Unrelated cell'
        wb.active['A2'] = '=SUM(1,2)'
        wb.active.add_image(Image(io.BytesIO(picture())), 'C4')
        wb.save(out)
    return out.getvalue()


def members(data):
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return {n: z.read(n) for n in z.namelist()}


def mutate(data, part, fn):
    original = members(data)
    root = ET.fromstring(original[part])
    fn(root)
    original[part] = ET.tostring(root)
    out = io.BytesIO()
    with zipfile.ZipFile(out, 'w') as z:
        for n, content in original.items():
            z.writestr(n, content)
    return out.getvalue()


@pytest.mark.parametrize('ext', ['docx', 'xlsx'])
def test_real_package_reopens_and_unrelated_bytes_survive(ext, tmp_path):
    data = document(ext)
    fixed, applied, unresolved = apply_office_image_replacement(data, ext, {'image 1': TEXT})
    assert applied == [{'locator': 'image 1', 'before': 'image of text', 'after': TEXT}]
    assert unresolved == []
    original, corrected = members(data), members(fixed)
    assert not any('/media/' in p for p in corrected)
    edited = ('word/document.xml', 'word/_rels/document.xml.rels') if ext == 'docx' else (
        'xl/drawings/drawing1.xml', 'xl/drawings/_rels/drawing1.xml.rels')
    for part, content in original.items():
        if part in corrected and part not in edited:
            assert corrected[part] == content, part
    root = ET.fromstring(corrected[edited[0]])
    if ext == 'docx':
        from docx import Document
        reopened = Document(io.BytesIO(fixed))
        assert reopened.paragraphs[0].text == 'Unrelated heading'
        assert reopened.paragraphs[1].text == 'Unrelated body copy'
        assert reopened.paragraphs[2].text == TEXT
        assert len(reopened.inline_shapes) == 0
    else:
        from openpyxl import load_workbook
        reopened = load_workbook(io.BytesIO(fixed))
        assert reopened.active['A1'].value == 'Unrelated cell'
        assert reopened.active['A2'].value == '=SUM(1,2)'
        assert root.xpath('.//a:t/text()', namespaces=NS) == TEXT.splitlines()
        assert not root.xpath('.//xdr:pic', namespaces=NS)
        before = ET.fromstring(original[edited[0]])
        for name in ('from', 'to', 'ext', 'pos'):
            assert [ET.tostring(e) for e in root.xpath('.//xdr:' + name, namespaces=NS)] == [
                ET.tostring(e) for e in before.xpath('.//xdr:' + name, namespaces=NS)]
    # Independent detector traversal, not the adapter's index, sees no raster.
    import ocr
    path = tmp_path / ('fixed.' + ext)
    path.write_bytes(fixed)
    assert list(ocr._ooxml_images(path)) == []


@pytest.mark.parametrize('ext', ['docx', 'xlsx'])
def test_real_ocr_finding_clears_after_raster_replacement(ext, tmp_path):
    import ocr
    if not ocr.is_available():
        pytest.skip('OCR not installed')
    path = tmp_path / ('original.' + ext)
    path.write_bytes(document(ext))
    assert ocr.images_of_text(path, '.' + ext), 'Fixture must actually trigger detector'
    fixed, applied, unresolved = apply_office_image_replacement(path.read_bytes(), ext, {'image 1': TEXT})
    assert applied and not unresolved
    path.write_bytes(fixed)
    assert ocr.images_of_text(path, '.' + ext) == []


@pytest.mark.parametrize('ext', ['docx', 'xlsx'])
@pytest.mark.parametrize('value', ['', ' ', '\x00bad', 'x' * 10001])
def test_invalid_transcript_is_unresolved_and_identical(ext, value):
    data = document(ext)
    assert apply_office_image_replacement(data, ext, {'image 1': value}) == (data, [], ['image 1'])


def test_reused_word_media_refuses_both_placements():
    data = document('docx', shared=True)
    assert apply_office_image_replacement(data, 'docx', {'image 1': TEXT}) == (data, [], ['image 1'])


@pytest.mark.parametrize('ext', ['docx', 'xlsx'])
def test_crop_refuses_without_changing_package(ext):
    part = 'word/document.xml' if ext == 'docx' else 'xl/drawings/drawing1.xml'
    def crop(root):
        fill = root.xpath('.//a:blip/..', namespaces=NS)[0]
        ET.SubElement(fill, '{' + NS['a'] + '}srcRect', l='10000')
    data = mutate(document(ext), part, crop)
    assert apply_office_image_replacement(data, ext, {'image 1': TEXT}) == (data, [], ['image 1'])


@pytest.mark.parametrize('ext', ['docx', 'xlsx'])
def test_missing_locator_is_unresolved(ext):
    data = document(ext)
    assert apply_office_image_replacement(data, ext, {'image 9': TEXT}) == (data, [], ['image 9'])


@pytest.fixture
def store(monkeypatch, tmp_path):
    import store as store_module
    monkeypatch.setattr(store_module, '_SQLITE_PATH', tmp_path / 'ocr.db')
    return store_module.Store()


@pytest.mark.parametrize('ext', ['docx', 'xlsx'])
def test_approved_values_reach_saved_bytes_through_production_handler(ext, monkeypatch, store):
    import ocr
    if not ocr.is_available():
        pytest.skip('OCR not installed')
    import test_remediation_verified_pptx_image_of_text as lane
    monkeypatch.setattr(lane, 'FILE', 'benefits.' + ext)
    blob = lane._Blob(document(ext))
    lane._seed(store, {'image 1': TEXT})
    lane._run_lane(monkeypatch, store, blob)
    assert blob.uploads, 'The handler must persist the corrected Office bytes'
    assert not any('/media/' in p for p in members(blob.data))
    assert store.count_unapplied_approved_values(lane.SID, lane.FILE) == 0


def test_unreplaceable_word_picture_gets_no_upload_or_credit(monkeypatch, store):
    import test_remediation_verified_pptx_image_of_text as lane
    monkeypatch.setattr(lane, 'FILE', 'benefits.docx')
    original = document('docx', shared=True)
    blob = lane._Blob(original)
    lane._seed(store, {'image 1': TEXT})
    lane._run_lane(monkeypatch, store, blob)
    assert not blob.uploads
    assert blob.data == original
    assert store.count_unapplied_approved_values(lane.SID, lane.FILE) > 0


def test_grouped_excel_picture_is_unresolved():
    part = 'xl/drawings/drawing1.xml'
    def group(root):
        pic = root.xpath('.//xdr:pic', namespaces=NS)[0]
        anchor = pic.getparent()
        container = ET.Element('{' + NS['xdr'] + '}grpSp')
        anchor.replace(pic, container)
        container.append(pic)
    data = mutate(document('xlsx'), part, group)
    assert apply_office_image_replacement(data, 'xlsx', {'image 1': TEXT}) == (data, [], ['image 1'])


def test_word_floating_picture_is_unresolved():
    part = 'word/document.xml'
    def floating(root):
        root.xpath('.//wp:inline', namespaces=NS)[0].tag = '{' + NS['wp'] + '}anchor'
    data = mutate(document('docx'), part, floating)
    assert apply_office_image_replacement(data, 'docx', {'image 1': TEXT}) == (data, [], ['image 1'])


@pytest.mark.parametrize('ext', ['docx', 'xlsx'])
def test_xml_special_characters_are_preserved_as_text(ext):
    text = 'Health & safety <important> "quoted"'
    fixed, applied, unresolved = apply_office_image_replacement(document(ext), ext, {'image 1': text})
    assert applied and not unresolved
    part = 'word/document.xml' if ext == 'docx' else 'xl/drawings/drawing1.xml'
    root = ET.fromstring(members(fixed)[part])
    query = './/w:t/text()' if ext == 'docx' else './/a:t/text()'
    assert text in root.xpath(query, namespaces=NS)
