"""Real Word note pictures must be detected, addressed and checked in written copies."""
import copy
import io
import sys
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / 'api'))
from apply_alt import apply_alt_text
from formats.docx.detectors.non_text_content import detect
from formats.office.images import undescribed_images
from remediate_office import remediate_office

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
R = 'http://schemas.openxmlformats.org/package/2006/relationships'
CT = 'http://schemas.openxmlformats.org/package/2006/content-types'


def entries(data):
    with zipfile.ZipFile(io.BytesIO(data)) as package:
        return {name: package.read(name) for name in package.namelist()}


def note_document(captions=False):
    png = io.BytesIO()
    Image.new('RGB', (96, 72), 'blue').save(png, format='PNG')
    png.seek(0)
    document = Document()
    document.add_paragraph('Body text stays unchanged.')
    document.add_picture(png)
    buffer = io.BytesIO()
    document.save(buffer)
    parts = entries(buffer.getvalue())
    root = etree.fromstring(parts['word/document.xml'])
    picture = root.find(f'.//{{{W}}}drawing').getparent().getparent()
    picture.getparent().remove(picture)
    rels = etree.fromstring(parts['word/_rels/document.xml.rels'])
    image_rels = [copy.deepcopy(rel) for rel in rels if rel.get('Type', '').endswith('/image')]
    types = etree.fromstring(parts['[Content_Types].xml'])
    body = root.find(f'{{{W}}}body')
    for index, story in enumerate(['footnotes', 'endnotes'], 1):
        singular = story[:-1]
        notes = etree.Element(f'{{{W}}}{story}', nsmap=root.nsmap)
        # Standard note separator entries keep this an ordinary usable Word package.
        for note_id, kind in [('-1', 'separator'), ('0', 'continuationSeparator')]:
            item = etree.SubElement(notes, f'{{{W}}}{singular}', {f'{{{W}}}id':note_id, f'{{{W}}}type':kind})
            run = etree.SubElement(etree.SubElement(item, f'{{{W}}}p'), f'{{{W}}}r')
            etree.SubElement(run, f'{{{W}}}{kind}')
        note = etree.SubElement(notes, f'{{{W}}}{singular}', {f'{{{W}}}id':'1'})
        paragraph = etree.SubElement(note, f'{{{W}}}p')
        run = etree.SubElement(paragraph, f'{{{W}}}r')
        etree.SubElement(run, f'{{{W}}}t').text = f'{singular} text stays unchanged.'
        note_picture = copy.deepcopy(picture)
        # Drawing IDs are unique across the package; repeated display names remain legal.
        note_picture.xpath('//*[local-name()="docPr"]')[0].set('id', str(index))
        note.append(note_picture)
        if captions:
            caption = etree.SubElement(note, f'{{{W}}}p')
            run = etree.SubElement(caption, f'{{{W}}}r')
            etree.SubElement(run, f'{{{W}}}t').text = f'Figure {index}: Blue area represents the synthetic measurement.'
        parts[f'word/{story}.xml'] = etree.tostring(notes, xml_declaration=True, encoding='UTF-8')
        own_rels = etree.Element(f'{{{R}}}Relationships', nsmap={None:R})
        for rel in image_rels:
            own_rels.append(copy.deepcopy(rel))
        parts[f'word/_rels/{story}.xml.rels'] = etree.tostring(own_rels)
        etree.SubElement(rels, f'{{{R}}}Relationship', Id=f'rIdSynthetic{story}', Type=f'http://schemas.openxmlformats.org/officeDocument/2006/relationships/{story}', Target=f'{story}.xml')
        etree.SubElement(types, f'{{{CT}}}Override', PartName=f'/word/{story}.xml', ContentType=f'application/vnd.openxmlformats-officedocument.wordprocessingml.{story}+xml')
        reference = etree.Element(f'{{{W}}}p')
        run = etree.SubElement(reference, f'{{{W}}}r')
        etree.SubElement(run, f'{{{W}}}{singular}Reference', {f'{{{W}}}id':'1'})
        body.insert(len(body)-1, reference)
    parts['word/document.xml'] = etree.tostring(root, xml_declaration=True, encoding='UTF-8')
    parts['word/_rels/document.xml.rels'] = etree.tostring(rels)
    parts['[Content_Types].xml'] = etree.tostring(types)
    out = io.BytesIO()
    with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as package:
        for name, data in parts.items():
            package.writestr(name, data)
    return out.getvalue()


def test_both_reader_visible_note_images_are_detected_and_reachable(tmp_path):
    source = note_document()
    path = tmp_path / 'notes.docx'
    path.write_bytes(source)
    assert Document(io.BytesIO(source)).paragraphs[0].text == 'Body text stays unchanged.'
    missing = undescribed_images(entries(source))
    assert {item['part'] for item in missing} == {'word/footnotes.xml', 'word/endnotes.xml'}
    assert len(detect(path)) == 2
    values = {item['locator']: f'Synthetic blue measurement shown in {item["part"]}.' for item in missing}
    written, applied, unresolved = apply_alt_text(source, values)
    assert len(applied) == 2 and not unresolved
    path.write_bytes(written)
    assert not detect(path)
    before, after = entries(source), entries(written)
    assert before['word/document.xml'] == after['word/document.xml']
    assert all(before[name] == after[name] for name in before if name.startswith('word/media/'))


def test_authored_note_captions_are_written_to_real_corrected_copy_and_rechecked(tmp_path):
    source = note_document(captions=True)
    path = tmp_path / 'authored-notes.docx'
    path.write_bytes(source)
    assert len(detect(path)) == 2
    corrected, applied, skipped = remediate_office(path, ai_enabled=False, in_scope=lambda sc: sc == '1.1.1')
    assert corrected and corrected != path and applied
    assert path.read_bytes() == source
    assert not detect(corrected)
    before, after = entries(source), entries(corrected.read_bytes())
    for story in ['footnotes', 'endnotes']:
        original = etree.fromstring(before[f'word/{story}.xml'])
        result = etree.fromstring(after[f'word/{story}.xml'])
        assert original.xpath('//*[local-name()="t"]/text()') == result.xpath('//*[local-name()="t"]/text()')
        description = result.xpath('//*[local-name()="docPr"]/@descr')
        assert len(description) == 1 and 'Blue area represents' in description[0]
    assert Document(corrected).paragraphs[0].text == 'Body text stays unchanged.'


def test_describing_one_note_keeps_the_other_note_open(tmp_path):
    source = note_document()
    missing = undescribed_images(entries(source))
    locator = next(item['locator'] for item in missing if item['part']=='word/footnotes.xml')
    written, applied, unresolved = apply_alt_text(source, {locator:'Blue synthetic measurement in the footnote.'})
    assert len(applied)==1 and not unresolved
    assert [item['part'] for item in undescribed_images(entries(written))] == ['word/endnotes.xml']
