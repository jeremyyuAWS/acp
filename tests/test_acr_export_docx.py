"""The ACR as an accessible Word document — Phase 5 groundwork (PRD §16).

WHAT IS BEING CLAIMED, AND WHAT IS DELIBERATELY NOT. This is not the ITI VPAT® template and does
not pretend to be; the document says so on its own first page and a test below asserts it. What is
claimed is narrower and checkable: the .docx ACP generates has a real heading outline, a header row
that repeats across pages, a declared document language, a title in its properties, and no colour
carrying meaning — and ACP's own Word analyser finds no failure in it.

THE PAIR THAT MATTERS is `test_the_generated_document_passes_acps_own_word_analyser` together with
`test_the_gate_is_not_vacuous`. The first alone is worthless: a checker that reports nothing on a
broken file also reports nothing on a good one, and "0 failures" would then be a fact about the
checker rather than about the document. The second feeds the same gate documents that are wrong on
purpose and requires it to say so.

WHY "NO FAIL" IS THE BAR. Measured on this repo: 10 docx registrations, every one `partial` or
`heuristic`, none `full` — and `assessment.CAN_CERTIFY_PASS` is `frozenset({Coverage.FULL})`. So
`PASS` is unreachable for any Word document by construction, and an "all PASS" export gate could
never go green however good the document was. A gate nobody can satisfy is one everybody
eventually disables. `test_no_docx_registration_can_certify_a_pass` pins the measurement, so if a
FULL-coverage docx technique ever lands, this reasoning is revisited rather than inherited.
"""
from __future__ import annotations

import sys
import zipfile
from pathlib import Path

import pytest

ACP = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ACP / "api"))

pytest.importorskip("docx")

import acr_export_docx  # noqa: E402
import acr_export_preview  # noqa: E402

REPORT = {"report_title": "ACP ACR", "product_name": "ACP by Movate", "product_version": "1.4.0",
          "wcag_version": "2.2", "vendor_name": "Movate"}
CRITERIA = [
    {"criterion_num": "1.4.3", "criterion_name": "Contrast (Minimum)", "level": "AA",
     "principle": "Perceivable", "final_status": "Supports", "remarks": "Measured at 7:1."},
    {"criterion_num": "2.4.7", "criterion_name": "Focus Visible", "level": "AA",
     "principle": "Operable", "final_status": "Does Not Support",
     "remarks": "Focus ring missing on the date picker."},
]


@pytest.fixture(scope="module")
def projection():
    return acr_export_preview.project(REPORT, CRITERIA)


@pytest.fixture(scope="module")
def rendered(projection):
    return acr_export_docx.render(projection)


def _parts(blob: bytes, tmp_path) -> dict[str, str]:
    path = tmp_path / "acr.docx"
    path.write_bytes(blob)
    with zipfile.ZipFile(path) as zf:
        return {name: zf.read(name).decode("utf-8")
                for name in ("word/document.xml", "word/styles.xml", "docProps/core.xml")}


# ── the document is a document ─────────────────────────────────────────────────

def test_it_produces_a_real_docx(rendered, tmp_path):
    assert rendered[:2] == b"PK", rendered[:8]
    path = tmp_path / "a.docx"
    path.write_bytes(rendered)
    with zipfile.ZipFile(path) as zf:
        assert "word/document.xml" in zf.namelist()


# ── the accessibility features, each asserted where Word actually stores it ────

def test_the_headings_are_real_headings_and_not_bold_paragraphs(rendered, tmp_path):
    """1.3.1. A bold paragraph looks identical on screen and is invisible to a screen reader's
    heading navigation, which is the only way a reader gets around a 55-row report."""
    doc = _parts(rendered, tmp_path)["word/document.xml"]
    assert 'w:val="Heading1"' in doc
    assert 'w:val="Heading2"' in doc


def test_the_table_header_row_repeats_across_pages(rendered, tmp_path):
    """1.3.1, and the same defect `acr_export_pdf`'s `display: table-header-group` fixes for the
    PDF. Without `<w:tblHeader/>` the column meanings appear on page one only, and a reader on
    page four is holding four headings in their head."""
    doc = _parts(rendered, tmp_path)["word/document.xml"]
    # Both tables — report metadata and the conformance table.
    assert doc.count("<w:tblHeader/>") == 2, doc.count("<w:tblHeader/>")


def test_the_document_declares_its_language(projection, tmp_path):
    """3.1.1. With no declared language a synthesiser reads an English conformance report in
    whatever voice it happens to be set to, which is noise rather than a document.

    RENDERED IN A NON-DEFAULT LANGUAGE ON PURPOSE. python-docx's default template already ships
    `<w:lang w:val="en-US" .../>` in styles.xml, so asserting that `w:lang` and `en-US` appear
    passes whether or not this renderer ever touches the styles — it measures the template. The
    first draft of this test did exactly that, and a bite check that deleted the call to
    `_set_document_language` left the whole suite green. Asking for `fr-CA` is what makes the
    assertion about this code.

    ASSERTED ON THE EXTRACTED ELEMENTS, not with `in` over the raw XML. `styles.xml` is ~350 KB,
    and a failing `assert 'x' not in <350KB string>` sent pytest's assertion introspection into a
    14-minute spin. A list of the two `w:lang` elements is both cheaper and a sharper claim.
    """
    import re

    def langs(blob):
        return re.findall(r'<w:lang w:val="([^"]+)"', _parts(blob, tmp_path)["word/styles.xml"])

    found = langs(acr_export_docx.render(projection, language="fr-CA"))
    assert found, "the renderer declared no language at all"
    # EVERY declaration, not just one: the template puts a language in docDefaults as well as on
    # the Normal style, and setting only the style leaves the file answering the question twice
    # with different values.
    assert set(found) == {"fr-CA"}, f"a stale language survived: {found}"

    assert set(langs(acr_export_docx.render(projection))) == {acr_export_docx.DOCUMENT_LANGUAGE}


def test_the_title_is_in_the_document_properties_not_only_the_body(rendered, tmp_path):
    """2.4.2. The properties title is what a screen reader and a file browser announce; a heading
    on page one is not a substitute for it."""
    core = _parts(rendered, tmp_path)["docProps/core.xml"]
    assert "<dc:title>" in core
    assert "ACP ACR" in core


def test_the_conformance_level_is_the_cell_text(rendered, tmp_path):
    """1.4.1 — nothing in this document communicates by colour. The conformance term is the cell's
    text, which is also why it survives being printed in black and white."""
    doc = _parts(rendered, tmp_path)["word/document.xml"]
    assert "Supports" in doc and "Does Not Support" in doc
    assert "w:highlight" not in doc


# ── the export gate: the pair ──────────────────────────────────────────────────

def test_the_generated_document_passes_acps_own_word_analyser(rendered):
    """PRD §16's export gate, on the real artifact. Worthless without the next test."""
    result = acr_export_docx.check(rendered)
    assert result["ok"] is True, result["failures"]
    assert result["failures"] == []


@pytest.mark.parametrize("defect", ["heading-skip", "empty-heading"])
def test_the_gate_is_not_vacuous(defect, tmp_path):
    """The same gate, on documents that are wrong on purpose.

    A checker that reports nothing on a broken file also reports nothing on a good one, so
    "0 failures" above would otherwise be a fact about the checker rather than about the document.
    Both defects here are ones `office_structure.docx_checks` names explicitly, and the empty
    heading is the nastier: it is in the outline, breaks no level sequence and has no runs to fail
    contrast on, so it passes every other check ACP has.
    """
    from docx import Document

    document = Document()
    document.add_heading("Title", level=1)
    if defect == "heading-skip":
        document.add_heading("Skipped a level", level=3)
    else:
        document.add_heading("", level=2)
    document.add_paragraph("body")

    path = tmp_path / "broken.docx"
    document.save(path)

    result = acr_export_docx.check(path.read_bytes(), tmp_dir=tmp_path / "check")
    assert result["ok"] is False, "the gate did not notice a defect it claims to detect"
    assert result["failures"], result


def test_review_findings_are_returned_rather_than_counted_as_failures():
    """A REVIEW is "a human has to look", not "ACP approved it". Counting them as failures would
    make the gate unsatisfiable; swallowing them would turn "nothing ACP can decide" into a pass —
    the shape PRD §4.4 forbids. They come back in their own list."""
    result = acr_export_docx.check(acr_export_docx.render(
        acr_export_preview.project(REPORT, CRITERIA)))
    assert set(result) == {"ok", "failures", "reviews"}
    assert isinstance(result["reviews"], list)


def test_no_docx_registration_can_certify_a_pass():
    """The measurement the "no FAIL" bar rests on, pinned so the reasoning is revisited rather
    than inherited if a FULL-coverage docx technique ever lands."""
    import formats.docx  # noqa: F401  — registration is lazy; without this the registry is empty
    import rule_registry
    from assessment import CAN_CERTIFY_PASS

    regs = [r for r in rule_registry.all_registrations() if r.fmt == "docx"]
    assert regs, "no docx registrations at all — the import above stopped working"
    assert not any(r.coverage in CAN_CERTIFY_PASS for r in regs), (
        "a docx technique now declares FULL coverage, so PASS is reachable and the export gate "
        "should be reconsidered rather than left at 'no FAIL'")


# ── what the document says about itself ────────────────────────────────────────

def test_it_says_it_is_not_a_vpat(rendered, tmp_path):
    """A .docx of a conformance report is the artifact in this product most likely to be mistaken
    for a VPAT, because that is the format a real one arrives in."""
    doc = _parts(rendered, tmp_path)["word/document.xml"]
    assert "not a VPAT" in doc
    assert "must not be submitted as one" in doc


def test_it_does_not_claim_the_pdf_renderer_s_validation(rendered, tmp_path):
    """The PDF notice cites veraPDF and PDF/UA-1. Neither says anything whatsoever about a Word
    file, and copying that wording across would be a conformance claim about a format it was never
    measured on — the cheapest possible way to make this document dishonest."""
    doc = _parts(rendered, tmp_path)["word/document.xml"]
    assert "veraPDF" not in doc
    assert "PDF/UA" not in doc
    assert "necessary and not sufficient" in doc


def test_the_totals_are_counts_and_never_a_percentage(rendered, tmp_path):
    """ADR 0016/0023 and PRD §4.4. A "% compliant" figure on a conformance report is the exact
    thing this feature exists not to produce."""
    doc = _parts(rendered, tmp_path)["word/document.xml"]
    assert "%" not in doc


# ── one projection, four renderers ─────────────────────────────────────────────

def test_the_docx_and_the_projection_cannot_disagree(projection, rendered, tmp_path):
    """It consumes `project()` like every other export, so what a reviewer approved on screen is
    what the customer receives. Asserted as the consequence rather than by reading the source."""
    doc = _parts(rendered, tmp_path)["word/document.xml"]
    missing = [c["criterion_num"] for c in projection["criteria"]
               if c["criterion_num"] not in doc]
    assert not missing, f"in the projection but not the document: {missing}"


def test_an_undecided_criterion_inherits_the_projections_placeholder(tmp_path):
    """The §9 guard lives in `acr_export_preview`, and this renderer gets it for free by consuming
    the projection. Asserted here so a future rewrite that built its own cells would fail."""
    proj = acr_export_preview.project(REPORT, [dict(CRITERIA[0], final_status=None, remarks="")])
    doc = _parts(acr_export_docx.render(proj), tmp_path)["word/document.xml"]
    assert "not yet evaluated" in doc
    # Counted as undecided rather than as any of the four terms. Asserted through the totals
    # because "Supports" appears in that line unconditionally — `_totals` emits every status with
    # its count, so `"Supports" not in doc` is false for a document with no Supports at all. An
    # earlier version of this test asserted exactly that and failed against correct behaviour.
    assert "undecided: 1" in doc


# ── the plumbing ───────────────────────────────────────────────────────────────

def test_the_filename_names_the_report_and_its_revision():
    """Both spellings of the id, for the reason `acr_export_pdf.filename_for` records: the store
    row calls it `id`, the projection calls it `report_id`, and reading only one is how every
    report in the system downloaded under a single name."""
    assert acr_export_docx.filename_for({"id": "acr_abc"}) == "acr-acr_abc.docx"
    assert acr_export_docx.filename_for({"report_id": "acr_abc"}) == "acr-acr_abc.docx"
    assert acr_export_docx.filename_for({"id": "acr_abc", "revision": 3}) == "acr-acr_abc-rev3.docx"
    assert "/" not in acr_export_docx.filename_for({"id": "../../etc/passwd"})


def test_is_available_reports_what_render_would_do(monkeypatch):
    """A caller that asks before offering a download must get the answer rendering would give.
    If these disagree the UI offers a button that always fails, or hides one that would work."""
    assert acr_export_docx.is_available() is True
    monkeypatch.setitem(sys.modules, "docx", None)
    assert acr_export_docx.is_available() is False
