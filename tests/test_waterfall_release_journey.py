"""Cross-service waterfall acceptance on synthetic documents and isolated SQLite.

Only provider HTTP and artifact storage are doubles. Durable run policy, budget,
text transport parsing, attempt/review history, title writer and missing-title
check are real. This proves title presence and controlled release readiness, not
semantic title quality, full accessibility, or a real cloud/storage deployment.
"""
from dataclasses import asdict
import hashlib
import io
import json
import sys
import time
from types import SimpleNamespace

import httpx
import pytest

from ai_attempt_history import AttemptHistory
from ai_review_chain import read_reviews
from ai_run_policy import run_context
import llm_waterfall_provider as bounded
from remediation_waterfall_view import read_waterfall
from test_approved_release_progression import MemoryBlob
from test_llm_waterfall_provider import FakeProviders, Response, result

SID = 'waterfall-release-journey'
FILE = 'synthetic-review.pptx'
TITLE = 'Quarterly accessibility review'
PROMPT = 'Source slide is about quarterly accessibility review. Suggest a short slide title.'


@pytest.fixture
def journey(isolated_store, monkeypatch, tmp_path):
    import core
    import handlers
    from office_structure import pptx_checks
    from pptx import Presentation
    from proposals import Verification

    store = isolated_store
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[5])
    output = io.BytesIO()
    deck.save(output)
    blob = MemoryBlob(output.getvalue())
    store.init_scan_run(SID, 'local', 1, '2026-09-08T00:00:00Z', 'fixture', 'hash', owner='owner')
    store.save_file_result(SID, {
        'file': FILE, 'engine': 'office', 'status': 'pass', 'score': 60,
        'compliant': 0, 'skipped_rules': 0,
        'issues': [{'ruleId': 'PPTX_TITLE_EMPTY', 'wcag': '2.4.6 Headings and Labels',
                    'severity': 'MODERATE', 'location': 'Slide 1'}],
    }, '2026-09-08T00:00:00Z')
    store.record_remediation(SID, FILE, blob_url='https://fixture.invalid/corrected-v1',
                             corrected_sha256=hashlib.sha256(blob.data).hexdigest(),
                             corrected_bytes=len(blob.data))

    def verify(data, filename, *, scan_id=None):
        assert scan_id == SID
        path = tmp_path / filename
        path.write_bytes(data)
        missing = any(row['ruleId'] == 'PPTX_TITLE_EMPTY' for row in pptx_checks(path))
        return Verification(True, {'2.4.6'} if missing else set())

    monkeypatch.setattr(core, 'store', store)
    monkeypatch.setitem(sys.modules, 'blob', blob)
    monkeypatch.setattr(handlers, '_verify_residual', verify)
    # Any accidental bypass of the explicitly injected transport fails the test.
    monkeypatch.setattr(httpx, 'post', lambda *a, **kw: pytest.fail('unexpected network transport'))
    specs = tuple(bounded.TextModelSpec('openai', name, 'fixture-price-v1', '1', '2',
        8192, 128, int(time.time()) + 3600) for name in ('fixture-first-v1', 'fixture-fallback-v1'))
    monkeypatch.setenv('ACP_BOUNDED_TEXT_MODELS_JSON', json.dumps([asdict(spec) for spec in specs]))

    def start(*, ai=1, cap='5.00', review=True, fail_usage=False, expensive=False):
        policy = {'ai': ai, 'rule_based': 2, 'ai_budget_usd': cap,
                  'ai_review': {'enabled': review}}
        execution = store.enqueue_stage_batch(SID, 'remediate', 'remediate_file', [
            {'owner': 'owner', 'scan_id': SID, 'file': FILE, 'remediation_impact_policy': policy}],
            snapshot_id='fixture-assessment', request_fingerprint='fixture-run')
        batch = execution['batch_id']
        store.seed_finding_dispositions(SID, batch)
        job = store.get_job(execution['job_ids'][0])
        payload = job['payload']
        if isinstance(payload, str):
            payload = json.loads(payload)
        calls = []

        def post(*args, **kwargs):
            model = kwargs['json']['model']
            calls.append(model)
            if fail_usage:
                raise httpx.ReadTimeout('fixture unknown usage')
            text = '' if len(calls) == 1 else TITLE if len(calls) == 2 else json.dumps({
                'verdict': 'accept', 'reason': 'Suitable for human consideration.'})
            data = result(model=model, text=text)
            if expensive:
                data['usage'] = {'prompt_tokens': 8192, 'completion_tokens': 128}
            return Response(data)

        generator = bounded.StrictTextGenerator(specs, provider_module=FakeProviders, post=post)
        monkeypatch.setattr(bounded, 'configured_generator', lambda: generator)
        return SimpleNamespace(store=store, blob=blob, handlers=handlers, verify=verify,
                               batch=batch, job=job, payload=payload, calls=calls)
    return start


def test_empty_first_fallback_review_human_approval_verified_artifact(journey):
    import providers
    from pptx import Presentation

    j = journey()
    original = j.blob.data
    with run_context(j.store, j.payload, j.job):
        draft = providers.text_generate(PROMPT)
    assert j.calls == ['fixture-first-v1', 'fixture-fallback-v1', 'fixture-first-v1']
    assert draft['text'] == TITLE and draft['approval_required'] is True
    assert draft['review']['verdict'] == 'accept'
    assert j.blob.data == original and j.blob.uploads == 0
    history = AttemptHistory(j.store._db).list_run('owner', SID, j.batch)
    assert [row['purpose'] for row in history] == ['draft', 'fallback', 'review']
    assert [row['result']['text'] for row in history[:2]] == ['', TITLE]
    assert all(row['spending_state'] == 'settled' for row in history)
    receipt = read_reviews(j.store._db, 'owner', SID, j.batch)[0]
    assert receipt['proposal_sha256'] == hashlib.sha256(TITLE.encode()).hexdigest()
    view = read_waterfall(j.store, 'owner', SID, j.batch)
    assert view['spending']['spent_units'] == 360
    assert view['spending']['held_units'] == 0
    assert [row['models'][0]['model'] for row in view['stages']] == ['fixture-first-v1', 'fixture-fallback-v1']
    assert not view['contribution_available']  # three paid calls are not three fixed findings
    with pytest.raises(LookupError):
        read_waterfall(j.store, 'intruder', SID, j.batch)

    # Production proposal persistence + exact human-approved value, not fabricated
    # model-to-finding attribution. The draft-to-queue orchestration is an explicit seam.
    item = j.store.enqueue_proposals(SID, FILE, '2.4.6', [{
        'locator': 'slide 1', 'before': '', 'proposed_value': draft['text'],
        'rationale': 'Synthetic source title suggestion', 'source': 'ai',
    }], rule_name='Headings and Labels')
    assert j.store.finding_reconciliation(SID, j.batch)['awaiting_review'] == 1
    j.store.complete_hitl_decision(item, 'approved', None, None, resolution=None,
        approved_values=[TITLE], actor='owner', detail='Synthetic acceptance fixture',
        request_id='fixture-approval', expected_version=0)
    assert j.store.finding_reconciliation(SID, j.batch)['approved_pending_verification'] == 1
    assert not j.store.mark_file_compliant_if_reviewed(SID, FILE)
    assert j.blob.data == original
    j.handlers._apply_approved_values({'owner': 'owner', 'scan_id': SID, 'file': FILE}, {})
    assert Presentation(io.BytesIO(j.blob.data)).slides[0].shapes.title.text == TITLE
    assert j.verify(j.blob.data, FILE, scan_id=SID).cleared({'2.4.6'})
    reconciliation = j.store.finding_reconciliation(SID, j.batch)
    assert reconciliation['exact'] and reconciliation['resolved_verified'] == 1
    assert reconciliation['awaiting_review'] == reconciliation['approved_pending_verification'] == 0
    assert j.store.get_hitl_item(item)['applied']
    record = j.store.get_file_record(SID, FILE)
    assert record['compliant'] and record['remediated_at']
    with j.store._db.cursor() as cur:
        j.store._db.execute(cur, 'SELECT corrected_sha256 FROM file_records WHERE scan_id=%s AND file=%s', (SID, FILE))
        assert j.store._db.fetchone(cur)['corrected_sha256'] == hashlib.sha256(j.blob.data).hexdigest()
    with run_context(j.store, j.payload, j.job):
        assert providers.text_generate(PROMPT)['text'] == TITLE
    j.handlers._apply_approved_values({'owner': 'owner', 'scan_id': SID, 'file': FILE}, {})
    assert len(j.calls) == 3 and j.blob.uploads == 1
    assert read_waterfall(j.store, 'owner', SID, j.batch)['spending']['spent_units'] == 360


@pytest.mark.parametrize('ai,cap', [(0, '5.00'), (1, '0.00')])
def test_rules_only_or_zero_budget_never_admit_ai(journey, ai, cap):
    import providers
    j = journey(ai=ai, cap=cap)
    with run_context(j.store, j.payload, j.job):
        assert providers.text_generate(PROMPT)['deferred']
    view = read_waterfall(j.store, 'owner', SID, j.batch)
    assert not view['ai_enabled']
    assert view['spending']['spent_units'] == view['spending']['held_units'] == 0
    assert j.calls == [] and j.blob.uploads == 0
    assert AttemptHistory(j.store._db).list_run('owner', SID, j.batch) == []


def test_unknown_provider_charge_stops_fallback_and_preserves_reservation(journey):
    import providers
    j = journey(fail_usage=True)
    with run_context(j.store, j.payload, j.job):
        draft = providers.text_generate(PROMPT)
        again = providers.text_generate(PROMPT)
    assert draft['reason'] == 'provider_usage_unknown'
    assert again['deferred'] and not draft.get('text')
    view = read_waterfall(j.store, 'owner', SID, j.batch)
    assert view['spending']['unknown_charges'] == 1 and view['spending']['blocked']
    assert view['spending']['held_units'] == 8448 and view['spending']['spent_units'] == 0
    assert view['stages'][1]['operations'] == 0
    assert j.calls == ['fixture-first-v1'] and j.blob.uploads == 0
    assert j.store.finding_reconciliation(SID, j.batch)['resolved_verified'] == 0


def test_budget_exhausted_after_first_call_does_not_buy_fallback(journey):
    import providers
    j = journey(cap='0.01', expensive=True)
    with run_context(j.store, j.payload, j.job):
        draft = providers.text_generate(PROMPT)
    assert draft['reason'] == 'budget_admission_denied'
    view = read_waterfall(j.store, 'owner', SID, j.batch)
    assert view['spending']['spent_units'] == 8448
    assert view['spending']['available_units'] == 1552
    assert view['spending']['held_units'] == 0
    assert j.calls == ['fixture-first-v1'] and not draft.get('text')
    assert j.blob.uploads == 0
