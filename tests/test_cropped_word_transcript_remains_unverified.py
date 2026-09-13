"""A selectable transcript preserves accessibility progress without certifying pixels.

Retaining a cropped graphic avoids losing diagram content, but the real OCR detector
still sees the embedded raster. This is deliberately not a verified 1.4.5 repair.
"""
import io

import pytest
from docx import Document
from lxml import etree as ET

import ocr
from apply_office_image_replacement import NS
from test_remediation_verified_office_image_replacement import (
    TEXT, document, members, mutate,
)


@pytest.mark.skipif(not ocr.is_available(), reason='Tesseract required for independent verification')
def test_append_only_transcript_preserves_cropped_picture_and_remaining_finding(tmp_path):
    def add_crop(root):
        fill = root.xpath('//*[local-name()="blipFill"]')[0]
        ET.SubElement(fill, '{' + NS['a'] + '}srcRect', t='19861')

    original = mutate(document('docx'), 'word/document.xml', add_crop)
    reopened = Document(io.BytesIO(original))
    reopened.add_paragraph(TEXT)
    out = io.BytesIO()
    reopened.save(out)
    augmented = out.getvalue()

    before, after = members(original), members(augmented)
    media = [name for name in before if '/media/' in name]
    assert len(media) == 1
    assert after[media[0]] == before[media[0]], 'The useful graphic must remain unchanged'
    root = ET.fromstring(after['word/document.xml'])
    assert root.xpath('.//a:srcRect/@t', namespaces=NS) == ['19861']
    assert len(Document(io.BytesIO(augmented)).inline_shapes) == 1
    assert Document(io.BytesIO(augmented)).paragraphs[-1].text == TEXT

    source = tmp_path / 'cropped.docx'
    candidate = tmp_path / 'with-transcript.docx'
    source.write_bytes(original)
    candidate.write_bytes(augmented)
    baseline = ocr.images_of_text(source, '.docx')
    measured = ocr.images_of_text(candidate, '.docx')
    assert len(baseline) == 1, 'The fixture must reproduce a real raster-text finding'
    assert measured == baseline, 'Adding text must not be mistaken for removing raster text'
