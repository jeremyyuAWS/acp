"""Explicitly disabled OOXML header properties must change in the saved DOCX."""
import io
import zipfile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from remediate_office import _remediate_docx_structure


def package(header_value):
    doc = Document()
    doc.add_heading('Title', level=1)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'Name'
    table.cell(1, 0).text = 'Patient'
    header = OxmlElement('w:tblHeader')
    if header_value is not None:
        header.set(qn('w:val'), header_value)
    table.rows[0]._tr.get_or_add_trPr().append(header)
    source = io.BytesIO()
    doc.save(source)
    with zipfile.ZipFile(source) as z:
        return {n: z.read(n) for n in z.namelist()}


@pytest.mark.parametrize('disabled', ['0', 'false', 'off'])
def test_disabled_header_is_enabled_after_save_and_reopen(disabled):
    entries = package(disabled)
    original = dict(entries)
    diffs = []
    applied = _remediate_docx_structure(entries, diffs, in_scope=lambda sc: sc == '1.3.1')
    saved = io.BytesIO()
    with zipfile.ZipFile(saved, 'w') as z:
        for name, content in entries.items():
            z.writestr(name, content)
    reopened = Document(saved)
    header = reopened.tables[0].rows[0]._tr.trPr.find(qn('w:tblHeader'))
    assert header is not None and header.get(qn('w:val')) is None
    assert reopened.tables[0].cell(1, 0).text == 'Patient'
    assert any('header' in message for message in applied)
    assert any('tblHeader' in row['after'] for row in diffs)
    assert all(entries[n] == content for n, content in original.items() if n != 'word/document.xml')
    second_diffs = []
    _remediate_docx_structure(entries, second_diffs, in_scope=lambda sc: sc == '1.3.1')
    assert not second_diffs


@pytest.mark.parametrize('enabled', [None, '1', 'true', 'on'])
def test_enabled_header_is_not_reported_as_a_new_fix(enabled):
    entries = package(enabled)
    diffs = []
    _remediate_docx_structure(entries, diffs, in_scope=lambda sc: sc == '1.3.1')
    assert not diffs


def test_excluded_table_criterion_keeps_disabled_header():
    entries = package('false')
    before = entries['word/document.xml']
    diffs = []
    _remediate_docx_structure(entries, diffs, in_scope=lambda sc: False)
    assert entries['word/document.xml'] == before
    assert not diffs
