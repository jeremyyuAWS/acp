"""Actual written Office pixels, caption evidence and rescans govern a bounded retry."""
import io
import json
from hashlib import sha256
from types import SimpleNamespace

import pytest
from docx import Document
from PIL import Image, ImageDraw
from apply_alt import apply_alt_text
from caption_validation import validate_caption
from office_verified_retry import retry_caption, image_target
from proposals import Verification, verify_residual

LOC = 'word/document.xml#Picture 1'
BAD = 'A red square on a white background.'
GOOD = 'A red circle on a white background.'


def source(filename="file.docx"):
    image = Image.new('RGB', (96, 96), 'white')
    ImageDraw.Draw(image).ellipse((20, 20, 75, 75), fill='red')
    pixels = io.BytesIO(); image.save(pixels, format='PNG')
    output = io.BytesIO()
    if filename.endswith('.pptx'):
        from pptx import Presentation
        document = Presentation();slide=document.slides.add_slide(document.slide_layouts[6])
        slide.shapes.add_picture(io.BytesIO(pixels.getvalue()),0,0)
        document.save(output)
    elif filename.endswith('.xlsx'):
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image as SheetImage
        document=Workbook();document.active['A1']='Original body text remains unchanged.'
        document.active.add_image(SheetImage(io.BytesIO(pixels.getvalue())),'C1')
        document.save(output)
    else:
        document = Document(); document.add_paragraph('Original body text remains unchanged.')
        document.add_picture(io.BytesIO(pixels.getvalue()));document.save(output)
    return output.getvalue()



def run(*, filename="file.docx", response=GOOD, current=lambda: True, archive=lambda data:'immutable:' + sha256(data).hexdigest(),
        verify=verify_residual, baseline=None, failed_check=None):
    original = source(filename)
    from formats.office.images import undescribed_images
    import zipfile
    with zipfile.ZipFile(io.BytesIO(original)) as package:
        locator=undescribed_images({name:package.read(name) for name in package.namelist()})[0]['locator']
    failed, applied, unresolved = apply_alt_text(original, {locator: BAD})
    assert len(applied) == 1 and not unresolved
    events, calls = [], []
    def generate(pixels, proof):
        calls.append((pixels, proof))
        return {'text': response, 'settled': True, 'attempts': [{'model':'accepted-next-tier'}]}
    result = retry_caption(original=original, failed=failed, locator=locator, written_caption=BAD,
        baseline=baseline or verify(original, filename), failed_check=failed_check or verify(failed, filename),
        validate=validate_caption, generate=generate, persist=lambda action,proof:events.append((action,proof)),
        archive=archive, current=current, intent={'operation_id':'exact-operation', 'file':filename,
            'verify': lambda data:verify(data,filename)})
    return original, failed, result, events, calls


def test_real_wrong_shape_caption_retry_writes_and_reassesses_actual_office_copy():
    original, failed, result, events, calls = run()
    # Presence passed for the wrong description, so actual pixels supply the failed evidence.
    assert verify_residual(failed, 'file.docx').cleared({'1.1.1'})
    assert validate_caption(BAD, image_target(failed, LOC)[1])['status'] == 'rejected'
    assert result is not None
    fixed, check, changes, proof = result
    assert check.ok and check.cleared({'1.1.1'})
    assert image_target(fixed, LOC) == (GOOD, image_target(original, LOC)[1])
    assert Document(io.BytesIO(fixed)).paragraphs[0].text == Document(io.BytesIO(original)).paragraphs[0].text
    assert image_target(failed, LOC)[0] == BAD and image_target(original, LOC)[0] == ''
    assert len(calls) == 1 and events[0][0] == 'intent' and events[-1][0] == 'validated'
    assert proof['failed_artifact_sha256'] == sha256(failed).hexdigest()
    assert proof['approval_identity'] != proof['writer_identity']


@pytest.mark.parametrize('response', [BAD, 'A patient chart with a diagnosis.', ''])
def test_replacement_not_independently_validated_never_reaches_written_copy(response):
    _, _, result, events, calls = run(response=response)
    assert result is None and len(calls) == 1 and events[-1][0] == 'rejected'


def test_missing_failed_archive_blocks_provider_dispatch():
    _, _, result, events, calls = run(archive=lambda data:None)
    assert result is None and not calls and not events


def test_cancelled_context_blocks_provider_dispatch():
    _, _, result, events, calls = run(current=lambda:False)
    assert result is None and not calls and not events


def test_unknown_verification_does_not_retry():
    _, _, result, events, calls = run(failed_check=Verification(False, {'1.1.1'}, 'scanner unavailable'))
    assert result is None and not calls and not events


def test_reassessment_regression_rejects_model_even_with_good_caption():
    def verifier(data, filename):
        return Verification(True, {'1.4.3'}) if image_target(data,LOC)[0] == GOOD else Verification(True, set())
    _, _, result, events, calls = run(verify=verifier)
    assert result is None and len(calls) == 1 and events[-1][1]['reason'] == 'replacement_reassessment_failed'


def test_supersession_clears_only_exact_old_caption_obligation(isolated_store):
    store = isolated_store
    from unverified_changes import pending_records, blocks_certification
    original, failed, result, _, _ = run()
    fixed, _, changes, proof = result
    digest = sha256(fixed).hexdigest()
    with store._db.cursor() as cur:
        store._db.execute(cur,"INSERT INTO scan_runs(id,status,owner_email) VALUES('scan','done','owner')")
        store._db.execute(cur,"INSERT INTO file_records(scan_id,file,corrected_sha256,remediated_at) VALUES('scan','file.docx',%s,'2026-09-01')",(digest,))
    store.log_decision('system','apply.saved_unverified',scan_id='scan',file='file.docx',rule_id='1.1.1',detail=json.dumps({
        'artifact_sha256':sha256(original).hexdigest(),'source_sha256':'source',
        'requires_semantic_review':True,'assessment_revision':'revision','item_ids':['item'],
        'changes':[{'locator':LOC,'after':BAD}, {'locator':'word/header1.xml#Picture 1','after':'Other caption'}]}))
    saved = {**proof,'artifact_sha256':digest,'item_id':'item','source_revision':'revision',
             'verification':'independent_caption_and_actual_reassessment','original_outcome':'superseded_not_verified'}
    store.log_decision('system','office_retry.saved',scan_id='scan',file='file.docx',rule_id='1.1.1',detail=json.dumps(saved))
    pending = pending_records(store,'scan','file.docx')
    assert len(pending) == 1 and pending[0]['changes'] == [{'locator':'word/header1.xml#Picture 1','after':'Other caption'}]
    assert blocks_certification(store,'scan','file.docx')
    with store._db.cursor() as cur:
        store._db.execute(cur,"INSERT INTO hitl_queue(id,scan_id,file,rule_id,status,applied,proposals) VALUES('item','scan','file.docx','1.1.1','approved',1,'[]')")
    assert not store.mark_file_compliant_if_reviewed('scan','file.docx')
    with store._db.cursor() as cur:
        store._db.execute(cur,"SELECT id,detail FROM decision_log WHERE action='apply.saved_unverified'")
        old=store._db.fetchone(cur);value=json.loads(old['detail']);value['changes']=value['changes'][:1]
        store._db.execute(cur,'UPDATE decision_log SET detail=%s WHERE id=%s',(json.dumps(value),old['id']))
    assert pending_records(store,'scan','file.docx')==[]
    assert store.mark_file_compliant_if_reviewed('scan','file.docx')


def test_store_retry_hash_cas_rejects_changed_source_and_artifact(isolated_store):
    store=isolated_store
    with store._db.cursor() as cur:
        store._db.execute(cur,"INSERT INTO scan_runs(id,status,owner_email) VALUES('scan','done','owner')")
        store._db.execute(cur,"INSERT INTO file_records(scan_id,file,checksum,corrected_sha256) VALUES('scan','file.docx','source','old')")
    with store._db.cursor() as cur:
        store._db.execute(cur,"INSERT INTO stage_executions(execution_id,workflow_id,workflow_revision,scan_id,owner_email,stage,input_snapshot_id,request_fingerprint,state,created_at,updated_at) VALUES('run','scan',1,'scan','owner','remediate','snapshot','request','succeeded','2026-01-01','2026-01-01')")
    args=dict(previous_sha256='old',source_identity={'checksum':'source'},run_id='run',source_revision='snapshot',blob_url='immutable',corrected_sha256='new',corrected_bytes=9)
    assert not store.compare_and_set_retry_artifact('other','scan','file.docx',**args)
    assert not store.compare_and_set_retry_artifact('owner','scan','file.docx',**{**args,'source_identity':{'checksum':'changed'}})
    assert store.compare_and_set_retry_artifact('owner','scan','file.docx',**args)
    assert not store.compare_and_set_retry_artifact('owner','scan','file.docx',**args)


def test_approved_value_worker_restores_exact_parent_context(isolated_store, monkeypatch):
    import office_verified_retry as retry
    from ai_run_policy import optional_current_run_context
    from test_ai_standing_approval import seed, SID, FILE
    parent = seed(isolated_store, monkeypatch)
    assert optional_current_run_context() is None
    observed=[]
    monkeypatch.setattr(retry, '_attempt', lambda store,**kw:observed.append(optional_current_run_context()))
    retry.attempt(isolated_store, scan_id=SID, filename=FILE, tickets=[{'run_id':parent['batch_id']}])
    assert len(observed)==1 and observed[0].run_id==parent['batch_id'] and observed[0].file==FILE
    assert optional_current_run_context() is None


@pytest.mark.parametrize('pointer', ['https://evil.test/remediated/owner/scan/file.docx.retry/{digest}',
    'https://account.test/other/owner/scan/file.docx.retry/{digest}',
    'https://account.test/remediated/other/scan/file.docx.retry/{digest}',
    'https://account.test/remediated/owner/scan/other.docx.retry/{digest}',
    'https://account.test/remediated/owner/scan/%2e%2e/file.docx.retry/{digest}',
    'https://account.test/remediated/owner/scan/file.docx.retry/{digest}?secret=unused'])
def test_blob_retry_pointer_never_fetches_arbitrary_urls(monkeypatch, pointer):
    import blob, core
    data=b'exact immutable copy';digest=sha256(data).hexdigest()
    class Client:
        def __init__(self,key):self.url='https://account.test/remediated/' + key
        def download_blob(self,**kwargs):raise AssertionError('invalid pointer must not download')
    service=SimpleNamespace(get_blob_client=lambda *,container,blob:Client(blob))
    monkeypatch.setattr(blob,'_service_client',lambda:service)
    monkeypatch.setattr(core,'store',SimpleNamespace(get_file_record=lambda *a:{'blob_url':pointer.format(digest=digest),'corrected_sha256':digest},
        get_scan=lambda sid:{'run':{'owner_email':'owner'}}))
    assert blob.download_remediated('owner','scan','file.docx') is None


def test_blob_retry_pointer_reads_exact_copy_and_checks_content_digest(monkeypatch):
    import blob, core
    data=b'exact immutable copy';digest=sha256(data).hexdigest();payload=[data]
    class Client:
        def __init__(self,key):self.url='https://account.test/remediated/' + key
        def download_blob(self,**kwargs):return SimpleNamespace(readall=lambda:payload[0])
    service=SimpleNamespace(get_blob_client=lambda *,container,blob:Client(blob))
    monkeypatch.setattr(blob,'_service_client',lambda:service)
    url='https://account.test/remediated/owner/scan/file.docx.retry/' + digest
    monkeypatch.setattr(core,'store',SimpleNamespace(get_file_record=lambda *a:{'blob_url':url,'corrected_sha256':digest},
        get_scan=lambda sid:{'run':{'owner_email':'owner'}}))
    assert blob.download_remediated('owner','scan','file.docx')==data
    assert blob.download_remediated(None,'scan','file.docx') is None
    payload[0]=b'changed'
    assert blob.download_remediated('owner','scan','file.docx') is None


@pytest.mark.parametrize('filename', ['file.pptx', 'file.xlsx'])
def test_other_office_formats_retry_exact_pixels_and_actual_rescan(filename):
    original, failed, result, events, calls = run(filename=filename)
    assert result is not None and len(calls)==1
    fixed, check, changes, proof=result
    assert check.ok and check.cleared({'1.1.1'})
    assert image_target(fixed,proof['locator']) == (GOOD,image_target(original,proof['locator'])[1])
    assert image_target(failed,proof['locator'])[0] == BAD


def test_completed_run_dispatch_requires_server_retry_capability_and_settled_replay(isolated_store, monkeypatch):
    import time
    from llm_waterfall_provider import StrictTextGenerator, TextModelSpec, managed_generate_attempts
    from office_verified_retry import RetryAuthority
    from ai_spending_budget import BudgetLedger
    from ai_attempt_history import AttemptHistory
    from test_llm_waterfall_provider import FakeProviders, Response, result
    store=isolated_store;ledger=BudgetLedger(store._db);ledger.create_budget('owner','run',100000)
    AttemptHistory(store._db).init_schema()
    names=('small-pinned-v1','large-pinned-v1','third-pinned-v1')
    specs=tuple(TextModelSpec('openai',name,'fixture-price-v1','1','2',8192,128,int(time.time())+3600) for name in names)
    steps=[{'step_id':step,'position':i,'provider':'openai','model':name,'enabled':True,'capabilities':['text']}
           for i,(step,name) in enumerate(zip(('primary','fallback_1','fallback_2'),names))]
    ctx=SimpleNamespace(ledger=ledger,owner_id='owner',run_id='run',scan_id='scan',file='file.docx',enabled=True,deferred=[],policy={'generation_chain':{'version':1,'steps':steps}})
    with store._db.cursor() as cur:
        store._db.execute(cur,"INSERT INTO scan_runs(id,owner_email,status) VALUES('scan','owner','done')")
        store._db.execute(cur,"INSERT INTO stage_executions(execution_id,workflow_id,workflow_revision,scan_id,owner_email,stage,input_snapshot_id,request_fingerprint,state,created_at,updated_at) VALUES('run','scan',1,'scan','owner','remediate','snapshot','request','succeeded','2026-01-01','2026-01-01')")
    calls=[]
    generator=StrictTextGenerator(specs,provider_module=FakeProviders,post=lambda *a,**kw:calls.append(kw) or Response(result(model=names[1],text=GOOD)))
    refused=managed_generate_attempts('Caption',ctx,generator,purpose='review',tier_indices=(2,),operation_id='bounded')
    assert refused['deferred'] and not calls
    refused=managed_generate_attempts('Caption',ctx,generator,purpose='review',tier_indices=(2,),operation_id='bounded',verified_retry={'enabled':True})
    assert refused['deferred'] and not calls
    authority=RetryAuthority('owner','run','scan','file.docx',lambda:True)
    accepted=managed_generate_attempts('Caption',ctx,generator,purpose='review',tier_indices=(2,),operation_id='bounded',verified_retry=authority)
    assert accepted['text']==GOOD and len(calls)==1 and accepted['history_attempt_id']
    replay=managed_generate_attempts('Caption',ctx,generator,purpose='review',tier_indices=(2,),operation_id='bounded',verified_retry=authority)
    assert replay['replayed'] and len(calls)==1
    cancelled=RetryAuthority('owner','run','scan','file.docx',lambda:False)
    assert managed_generate_attempts('Caption',ctx,generator,purpose='review',tier_indices=(2,),operation_id='bounded',verified_retry=cancelled)['deferred']
    assert len(calls)==1


@pytest.mark.parametrize('state,status', [('uncertain','drafted'),('reserved','drafted'),('settled','started'),('settled','unusable_response')])
def test_history_id_without_settled_draft_is_not_write_permission(state,status):
    from office_verified_retry import settled_response
    row={'attempt_id':'attempt','file':'file.docx','model':'next','status':status,
         'spending_state':state,'input_sha256':'input','result':{'text':GOOD}}
    history=SimpleNamespace(list_operation=lambda *a,**kw:[row])
    ctx=SimpleNamespace(owner_id='owner',scan_id='scan',run_id='run',file='file.docx')
    response={'history_attempt_id':'attempt','input_sha256':'input','text':GOOD,'settled':True}
    assert not settled_response(history,ctx,'operation',response,'next')


def test_ordinary_filename_containing_retry_word_reads_canonical(monkeypatch):
    import blob,core
    data=b'ordinary saved document';filename='folder.retry/report.docx'
    class Client:
        def __init__(self,key):self.url='https://account.test/remediated/' + key
        def download_blob(self,**kwargs):return SimpleNamespace(readall=lambda:data)
    service=SimpleNamespace(get_blob_client=lambda *,container,blob:Client(blob))
    monkeypatch.setattr(blob,'_service_client',lambda:service)
    monkeypatch.setattr(core,'store',SimpleNamespace(get_file_record=lambda *a:{'blob_url':'https://account.test/remediated/owner/scan/' + filename}))
    assert blob.download_remediated('owner','scan',filename)==data


@pytest.mark.parametrize('mutation', ['none','delete_review','revoke_consent','missing_contribution'])
def test_live_adapter_uses_restored_frozen_run_and_only_next_settled_model(isolated_store, monkeypatch, mutation):
    import time
    import core, blob, release_artifacts, llm_waterfall_provider as transport
    from ai_run_policy import run_context, optional_current_run_context
    from ai_attempt_history import AttemptHistory
    from ai_standing_approval import approve_file
    from test_llm_waterfall_provider import FakeProviders, Response, result
    from office_verified_retry import attempt
    store=isolated_store;original=source();digest=sha256(original).hexdigest()
    monkeypatch.setattr(core,'store',store)
    monkeypatch.setattr(core,'get_scan_tokens',lambda sid:{'sp':'fixture-token'})
    monkeypatch.setattr(release_artifacts,'require_current_source',lambda *a,**kw:None)
    monkeypatch.setattr(blob,'download_remediated',lambda *a:original)
    archive=[]
    monkeypatch.setattr(blob,'upload_immutable_retry',lambda *args:archive.append(args[3]) or 'immutable:' + sha256(args[3]).hexdigest())
    names=('gpt-4.1-mini-2025-04-14','gpt-4.1-2025-04-14','gpt-4.1-nano-2025-04-14')
    specs=tuple(transport.TextModelSpec('openai',name,'fixture-price-v1','1','2',32768,128,int(time.time())+3600) for name in names)
    calls=[];selected={}
    def post(*a,**kw):
        calls.append(kw);model=kw['json']['model']
        if model==names[1] and mutation=='delete_review':
            with store._db.cursor() as cur:
                store._db.execute(cur,'DELETE FROM hitl_queue WHERE id=%s',(selected['item_id'],))
        if model==names[1] and mutation=='revoke_consent':
            from ai_run_approval_override import read,save
            current=read(store,'owner','scan',selected['run_id'])
            save(store,'owner','scan',selected['run_id'],False,current['revision'],current['source_revision'])
        return Response(result(model=model,text=BAD if model==names[0] else GOOD))
    generator=transport.StrictTextGenerator(specs,provider_module=FakeProviders,post=post)
    monkeypatch.setattr(transport,'configured_generator',lambda:generator)
    steps=[{'step_id':step,'position':i,'provider':'openai','model':name,'enabled':True,'capabilities':['text']}
           for i,(step,name) in enumerate(zip(('primary','fallback_1','fallback_2'),names))]
    policy={'rule_based':2,'ai':1,'ai_budget_usd':'1.00','auto_approve_ai':True,'generation_chain':{'version':1,'steps':steps}}
    with store._db.cursor() as cur:
        store._db.execute(cur,"INSERT INTO scan_runs(id,owner_email,status,source) VALUES('scan','owner','done','sharepoint')")
        store._db.execute(cur,"INSERT INTO file_records(scan_id,file,drive_file_id,source_modified,corrected_sha256,remediated_at) VALUES('scan','file.docx','source','2026-09-01',%s,'2026-09-02')",(digest,))
    batch=store.enqueue_stage_batch('scan','remediate','remediate_file',[{'scan_id':'scan','file':'file.docx','owner':'owner','source':'sharepoint','remediation_impact_policy':policy}],snapshot_id=store.remediation_source_revision('scan'),request_fingerprint='actual-caption')
    parent=store.get_job(batch['job_ids'][0])
    with run_context(store,parent['payload'],parent) as ctx:
        first=transport.managed_generate_attempts('First caption',ctx,generator,purpose='review',tier_indices=(1,),operation_id='first-caption')
        call=store.record_ai_call(surface='synthetic',provider='openai',model=names[0],zone='cloud',latency_ms=0,ok=True,scan_id='scan',file='file.docx')
        history=AttemptHistory(store._db);history.bind_trace('owner','scan',ctx.run_id,'first-caption',call,file='file.docx')
        item_id=store.enqueue_proposals('scan','file.docx','1.1.1',[{'locator':LOC,'before':'','proposed_value':BAD,'source':'AI synthetic','model':names[0],'model_call_id':call,'requires_semantic_review':True}])
        approve_file(store,ctx)
    item=store.get_hitl_item(item_id)
    assert item['status']=='approved' and optional_current_run_context() is None
    with store._db.cursor() as cur:
        store._db.execute(cur,"SELECT id FROM hitl_events WHERE item_id=%s AND action='standing_approve'",(item_id,))
        event_id=store._db.fetchone(cur)['id']
        store._db.execute(cur,"UPDATE stage_executions SET state='succeeded' WHERE execution_id=%s",(batch['batch_id'],))
    ticket={'run_id':batch['batch_id'],'scan_id':'scan','file':'file.docx','rule_id':'1.1.1','item_id':item_id,
            'proposal_id':item['approved_proposal_snapshot_ids'][0],'approval_event_id':event_id,
            'model_call_id':call,'approved_value':BAD,'locator':LOC,'source_revision':item['approved_source_revision'],
            'approved_value_sha256':item['approved_value_sha256'],'actual_source_sha256':digest}
    failed,applied,_=apply_alt_text(original,{LOC:BAD})
    kwargs=dict(scan_id='scan',filename='file.docx',original=original,failed=failed,values={LOC:BAD},applied=applied,
                baseline=verify_residual(original,'file.docx'),failed_check=verify_residual(failed,'file.docx'),tickets=[ticket],verify=lambda data:verify_residual(data,'file.docx'))
    selected.update(item_id=item_id,run_id=batch['batch_id'])
    # Retain one exact immutable baseline edge, then exercise the real queued writer.
    from remediation_contribution import digest as contribution_digest, encoded, proposal_digest, read_contribution
    with store._db.cursor() as cur:
        store._db.execute(cur, '''INSERT INTO remediation_contribution_runs(owner_id,scan_id,run_id,snapshot_id,baseline_json,files_json,created_at) VALUES('owner','scan',%s,%s,%s,%s,'2026-01-01') ON CONFLICT(owner_id,run_id) DO UPDATE SET baseline_json=excluded.baseline_json''',
            (batch['batch_id'], item['approved_source_revision'], encoded([{'finding_id':'finding','file':'file.docx','rule_id':'1.1.1','instance_key':LOC}]),encoded(['file.docx'])))
        store._db.execute(cur, '''INSERT INTO remediation_contribution_proposals
            (owner_id,scan_id,run_id,proposal_id,proposal_sha256,source_sha256,assessment_revision,file,rule_id,item_id,finding_ids_json,origin,attempt_id,operation_id,created_at)
            VALUES('owner','scan',%s,%s,%s,%s,%s,'file.docx','1.1.1',%s,%s,'first_ai',%s,'first-caption','2026-01-01')''',
            (batch['batch_id'],ticket['proposal_id'],proposal_digest(item['proposals'][0]),digest,
             item['approved_source_revision'],item_id,encoded(['finding']),first['history_attempt_id']))
        store._db.execute(cur, "SELECT payload FROM jobs WHERE type='apply_approved_values' AND scan_id='scan'")
        payload=json.loads(store._db.fetchone(cur)['payload'])
    if mutation=='missing_contribution':
        with store._db.cursor() as cur:
            store._db.execute(cur,'DELETE FROM remediation_contribution_proposals WHERE proposal_id=%s',(ticket['proposal_id'],))
        assert attempt(store,**kwargs) is None and len(calls)==1 and not archive
        return
    if mutation=='none':
        import handlers
        monkeypatch.setattr(handlers,'_verify_residual',lambda data,file,**kw:verify_residual(data,file))
        # Storage is a local receipt fixture; actual writer, rescan, SQL CAS and credits run.
        handlers._apply_approved_values(payload,{})
        record=store.get_file_record('scan','file.docx')
        assert record['corrected_sha256']!=digest and record['compliant']==1
        assert release_artifacts.release_ready(record)
        assert release_artifacts.require_current_record(store,'scan','file.docx',record['corrected_sha256'],
            record['remediated_at'],owner='owner')['corrected_sha256']==record['corrected_sha256']
        assert store.get_hitl_item(item_id)['applied']
        from unverified_changes import pending_records
        assert pending_records(store,'scan','file.docx')==[]
        summary=read_contribution(store,'owner','scan',batch['batch_id'])
        assert summary['outcomes']['fixed']==1 and summary['baseline_total']==1
        with store._db.cursor() as cur:
            store._db.execute(cur,"SELECT * FROM ai_validation_outcomes WHERE item_id=%s ORDER BY created_at",(item_id,))
            evidence=store._db.fetchall(cur)
            assert evidence[0]['outcome']=='could_not_verify'
            replacement=evidence[-1]
            assert replacement['outcome']=='verified_cleared' and replacement['proposal_snapshot_id']!=ticket['proposal_id']
            assert replacement['approval_event_id']!=event_id and replacement['artifact_sha256']==record['corrected_sha256']
            store._db.execute(cur,'SELECT attempt_id,model_call_id,proposal_json FROM ai_proposal_snapshots WHERE snapshot_id=%s',(replacement['proposal_snapshot_id'],))
            snapshot=store._db.fetchone(cur);caption=json.loads(snapshot['proposal_json'])
            assert snapshot['attempt_id']!=snapshot['model_call_id'] and snapshot['model_call_id']==replacement['model_call_id']
            assert caption['caption_validation']['approved'] and caption['pixel_sha256']==sha256(image_target(archive[-1],LOC)[1]).hexdigest()
            store._db.execute(cur,"SELECT detail FROM decision_log WHERE action='office_retry.saved'")
            proof=json.loads(store._db.fetchone(cur)['detail'])
        assert proof['original_outcome']=='superseded_not_verified'
        assert image_target(archive[-1],LOC)[0]==GOOD
        assert [c['json']['model'] for c in calls]==[names[0],names[1]]
        handlers._apply_approved_values(payload,{})
        assert len(calls)==2
        return
    retry=attempt(store,**kwargs)
    if mutation!='none':
        assert retry is None and len(calls)==2
        assert store.get_file_record('scan','file.docx')['corrected_sha256']==digest
        with store._db.cursor() as cur:
            store._db.execute(cur,"SELECT detail FROM decision_log WHERE action='office_retry.stopped'")
            assert json.loads(store._db.fetchone(cur)['detail'])['reason']=='authorization_changed_after_generation'
        return
    assert retry is not None and image_target(retry[0],LOC)[0]==GOOD
    assert [c['json']['model'] for c in calls]==[names[0],names[1]]
    assert optional_current_run_context() is None and original in archive and failed in archive
    with pytest.raises(ValueError,match='already_attempted'):
        attempt(store,**kwargs)
    assert len(calls)==2


@pytest.mark.parametrize('edit', ['stretch','crop','rotate','flip','effect','tile','fillRect'])
def test_office_presentation_transform_disables_raw_pixel_retry(edit):
    import zipfile
    from lxml import etree
    data=source()
    with zipfile.ZipFile(io.BytesIO(data)) as package:
        parts={name:package.read(name) for name in package.namelist()}
    root=etree.fromstring(parts['word/document.xml'])
    a='http://schemas.openxmlformats.org/drawingml/2006/main'
    if edit=='stretch':
        extent=root.xpath('//*[local-name()="extent"]')[0];extent.set('cx',str(2*int(extent.get('cx'))))
    elif edit=='crop':
        fill=root.xpath('//*[local-name()="blipFill"]')[0];etree.SubElement(fill,f'{{{a}}}srcRect',l='20000')
    elif edit=='tile':
        stretch=root.xpath('//*[local-name()="stretch"]')[0];stretch.getparent().replace(stretch,etree.Element(f'{{{a}}}tile'))
    elif edit=='fillRect':
        root.xpath('//*[local-name()="fillRect"]')[0].set('l','20000')
    elif edit in {'rotate','flip'}:
        transform=root.xpath('//*[local-name()="xfrm"]')[0];transform.set('rot' if edit=='rotate' else 'flipH','5400000' if edit=='rotate' else '1')
    else:
        blip=root.xpath('//*[local-name()="blip"]')[0];etree.SubElement(blip,f'{{{a}}}grayscl')
    parts['word/document.xml']=etree.tostring(root)
    output=io.BytesIO()
    with zipfile.ZipFile(output,'w') as package:
        for name,content in parts.items():package.writestr(name,content)
    assert image_target(output.getvalue(),LOC) is None


def test_another_undescribed_image_does_not_spend_on_caption_retry():
    document=Document(io.BytesIO(source()))
    pixels=image_target(source(),LOC)[1]
    document.add_picture(io.BytesIO(pixels));output=io.BytesIO();document.save(output)
    original=output.getvalue();failed,_,_=apply_alt_text(original,{LOC:BAD})
    calls=[]
    result=retry_caption(original=original,failed=failed,locator=LOC,written_caption=BAD,
        baseline=verify_residual(original,'file.docx'),failed_check=verify_residual(failed,'file.docx'),
        validate=validate_caption,generate=lambda *a:calls.append('generate'),persist=lambda *a:calls.append('persist'),
        archive=lambda *a:calls.append('archive'),current=lambda:True,
        intent={'operation_id':'other-image','file':'file.docx','verify':lambda data:verify_residual(data,'file.docx')})
    assert result is None and calls==[]


def test_real_two_inch_by_one_inch_docx_circle_is_not_auto_eligible():
    from docx.shared import Inches
    pixels=image_target(source(),LOC)[1]
    document=Document();document.add_picture(io.BytesIO(pixels),width=Inches(2),height=Inches(1))
    output=io.BytesIO();document.save(output)
    assert validate_caption(GOOD,pixels)['approved']
    assert image_target(output.getvalue(),LOC) is None


@pytest.mark.parametrize('filename', ['folder%20/file.docx','folder\\file.docx','../file.docx','a/../file.docx','/file.docx'])
def test_retry_unsafe_filename_stops_before_provider_or_archive(filename, isolated_store, monkeypatch):
    from office_verified_retry import attempt
    import llm_waterfall_provider
    monkeypatch.setattr(llm_waterfall_provider, 'configured_generator', lambda: pytest.fail('unsafe retry dispatched'))
    assert attempt(isolated_store, scan_id='scan', filename=filename, original=b'', failed=b'',
                   values={}, applied=[], baseline=None, failed_check=None, tickets=[], verify=None) is None
