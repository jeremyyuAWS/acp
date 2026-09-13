"""Narrow observable OCR contradiction; no broad semantic certification claim.

The frozen synthetic chart and caption are the actual 2026-09-13 GPU quality
control. Real OCR and a real XLSX image placement exercise the production writer;
the provider generation alone is replaced with its recorded caption.
"""
import re
import zipfile
from pathlib import Path
from types import SimpleNamespace

import pytest

import ai
import ocr
import remediate_office
from release_continuation import eligibility, proposal_identity

CAPTION = ('The image shows a bar chart with three bars, each representing a different category. '
           'The bars are labeled A, B, and C, and they are colored red, blue, and green, respectively. '
           'The x-axis of the chart is labeled "Monthly totals," indicating that the bars represent '
           'the total amounts for each month. The y-axis is labeled "Monthly totals," which is '
           'likely a mistake and should be "Monthly totals," indicating the total amounts for each '
           'month. The chart does not provide any numerical values, so it is not possible to '
           'determine the exact amounts represented by the bars.')
CHART = Path(__file__).parent / 'fixtures/vision/labeled_bar_chart.png'


@pytest.mark.parametrize('caption,text,blocked', [
    (CAPTION, 'Monthly totals\n30\n20', True),
    ('There are no numeric values.', 'Chart\n30', True),
    ('The image has no numbers.', 'Chart\n30', True),
    ('A chart without any numerical values.', 'Chart\n30', True),
    ('The chart does not display any numbers.', 'Chart\n30', True),
    ('The chart shows values 30 and 20.', 'Chart\n30\n20', False),
    ('Three bars labeled A, B, and C.', 'Monthly totals\n30\n20', False),
    ('The tallest bar has value 40.', 'Monthly totals\n30\n20', False),  # OCR omission is not disproof
    ('The chart has no numerical values.', '', False),
    ('The chart has no numerical values.', 'Monthly totals', False),
    ('The x-axis has no numerical values.', 'Monthly totals\n30\n20', False),
    ('There are no numerical values for January.', 'Monthly totals\n30\n20', False),
])
def test_guard_only_detects_explicit_global_numeric_denial(caption, text, blocked):
    assert ai._ocr_numeric_values_denied(caption, text) is blocked


@pytest.mark.skipif(not ocr.is_available(), reason='requires local tesseract OCR')
@pytest.mark.parametrize('validate_policy', [False, True])
def test_real_chart_ocr_blocks_automatic_office_write_but_preserves_draft(tmp_path, monkeypatch, validate_policy):
    from openpyxl import Workbook, load_workbook
    from openpyxl.drawing.image import Image
    import core

    text = ocr.ocr_text(CHART.read_bytes())
    assert 'Monthly totals' in text and re.search(r'\b(?:30|20)\b', text)
    monkeypatch.setattr(ai, '_vision_generate', lambda *a, **kw: CAPTION)
    monkeypatch.setattr(core.store, 'get_auto_apply_validated', lambda: validate_policy)
    validation = []
    def no_validation(*a, **kw):
        validation.append(True)
        raise AssertionError('An objectively contradicted draft must not enter agreement validation')
    monkeypatch.setattr(ai, 'validate_alt_text', no_validation)
    monkeypatch.setattr(ai, '_ALT_VALIDATOR_MODEL', 'fixture-validator')

    source = tmp_path / 'chart.xlsx'
    workbook = Workbook()
    workbook.active['A1'] = 'Preserved data'
    workbook.active['B2'] = '=SUM(1,2)'
    workbook.active.add_image(Image(CHART), 'D4')
    workbook.save(source)
    original = source.read_bytes()
    with zipfile.ZipFile(source) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    before = entries.copy()
    fixes, proposals = [], []
    applied, deferred = remediate_office._fix_image_alt(entries, vision_enabled=True,
        context_file='chart.xlsx', applied_fixes=fixes, proposals=proposals, evidence=[])
    assert applied == [] and deferred == 1 and fixes == [] and validation == []
    assert entries == before and source.read_bytes() == original
    assert load_workbook(source).active['B2'].value == '=SUM(1,2)'
    assert len(proposals) == 1
    proposal = proposals[0]
    assert proposal['proposed_value'] == CAPTION  # optional browse, not discarded
    assert proposal['automatic_write_blocked'] is True
    assert proposal['reason_code'] == 'ocr_numeric_values_denied'
    assert 'OCR read numeric values' in proposal['why_review']
    assert proposal['locator'].startswith('xl/drawings/drawing1.xml#')
    item = dict(id='fixture-item', rule_id='1.1.1', file='chart.xlsx', status='pending',
                finding_count=1, proposals=proposals, proposal_snapshot_ids=['snapshot'], decision_version=0)
    assert 'individual review' in eligibility(item, item['file'])
    from ai_standing_approval import eligible_item
    with pytest.raises(ValueError, match='individual review'):
        eligible_item(SimpleNamespace(_selected_sc=lambda *args: True), 'owner', 'scan', 'run', item)
    identity = proposal_identity(item)
    assert identity['proposals'][0]['automatic_write_blocked'] is True
    assert identity['proposals'][0]['reason_code'] == 'ocr_numeric_values_denied'


@pytest.mark.parametrize('caption,text,grounded', [
    ('Three bars labeled A, B, and C.', 'Monthly totals\n30\n20', True),
    ('The chart has no numerical values.', 'Monthly totals', True),
    ('The chart has no numerical values.', '', False),
])
def test_noncontradicted_structured_results_keep_original_grounding(monkeypatch, caption, text, grounded):
    monkeypatch.setattr(ocr, 'ocr_text', lambda data: text)
    monkeypatch.setattr(ai, '_vision_generate', lambda *a, **kw: caption)
    monkeypatch.setattr(ai, '_escalate_vision', lambda *a, **kw: None)  # no paid provider calls
    result = ai.describe_image_structured(b'fixture-image')
    assert result['alt'] == caption and result['grounded'] is grounded
    assert not result.get('automatic_write_blocked') and 'reason_code' not in result
