"""Saved Office link labels preserve their destination and surrounding content."""
import io
import zipfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from pptx import Presentation
from pptx.util import Inches
import pytest

from apply_link_text import apply_link_text

URL = 'https://example.invalid/patient-rights'


def test_formatted_powerpoint_link_is_replaced_once_and_reopens():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2))
    paragraph = shape.text_frame.paragraphs[0]
    for text, bold in [('click ', True), ('here', False)]:
        run = paragraph.add_run()
        run.text = text
        run.font.bold = bold
        run.hyperlink.address = URL
    paragraph.add_run().text = ' — keep this context'
    separate = shape.text_frame.add_paragraph().add_run()
    separate.text = 'here'
    separate.hyperlink.address = URL
    stream = io.BytesIO()
    presentation.save(stream)
    original = stream.getvalue()
    fixed, applied, unresolved = apply_link_text(original, 'pptx', {URL: 'Patient rights & responsibilities'})
    saved = Presentation(io.BytesIO(fixed)).slides[0].shapes[0].text_frame
    assert saved.paragraphs[0].text == 'Patient rights & responsibilities — keep this context'
    assert saved.paragraphs[1].text == 'Patient rights & responsibilities'
    assert saved.paragraphs[0].runs[0].font.bold is True
    assert saved.paragraphs[0].runs[1].font.bold is False
    assert all(run.hyperlink.address == URL for run in saved.paragraphs[0].runs[:2])
    assert [row['before'] for row in applied] == ['click here', 'here']
    assert len(applied) == 2 and unresolved == []
    with zipfile.ZipFile(io.BytesIO(original)) as old, zipfile.ZipFile(io.BytesIO(fixed)) as new:
        assert all(old.read(name) == new.read(name) for name in old.namelist() if name != 'ppt/slides/slide1.xml')


@pytest.mark.parametrize('kind', ['ins', 'drawing', 'fldChar', 'bookmarkStart'])
def test_complex_word_hyperlink_is_preserved_and_reported_unresolved(kind):
    document = Document()
    paragraph = document.add_paragraph('Keep this paragraph: ')
    hyperlink = OxmlElement('w:hyperlink')
    relationship = document.part.relate_to(URL, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink.set(qn('r:id'), relationship)
    run = OxmlElement('w:r')
    text = OxmlElement('w:t')
    text.text = 'click here'
    run.append(text)
    if kind == 'ins':
        revision = OxmlElement('w:ins')
        revision.set(qn('w:id'), '1')
        revision.append(run)
        hyperlink.append(revision)
    else:
        run.append(OxmlElement('w:' + kind))
        hyperlink.append(run)
    paragraph._p.append(hyperlink)
    stream = io.BytesIO()
    document.save(stream)
    original = stream.getvalue()
    fixed, applied, unresolved = apply_link_text(original, 'docx', {URL: 'Patient rights'})
    assert fixed == original
    assert applied == [] and unresolved == [URL]
    # A real Word package still opens, including the preserved complex content.
    assert Document(io.BytesIO(fixed)).paragraphs[0].text.startswith('Keep this paragraph: ')
