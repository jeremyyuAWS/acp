"""Automatic faithful alt text must leave real Excel packages readable."""
import sys
import re
from xml.sax.saxutils import escape
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest
from PIL import Image
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as ExcelImage

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'api'))
import remediate_office
from formats.office.images import undescribed_images


@pytest.mark.parametrize('prefixed', [False, True])
@pytest.mark.parametrize('description', ['Quarterly revenue chart', 'Sales & operations <overview>'])
def test_automatic_faithful_alt_preserves_actual_drawing_tag_and_workbook(tmp_path, prefixed, description):
    picture = tmp_path / 'picture.png'
    Image.new('RGB', (80, 60), 'blue').save(picture)
    source = tmp_path / 'example.xlsx'
    workbook = Workbook()
    workbook.active['A1'] = 'Preserved data'
    workbook.active['B2'] = '=SUM(1,2)'
    workbook.active.add_image(ExcelImage(picture), 'D4')
    workbook.save(source)
    with zipfile.ZipFile(source) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    drawing_name = 'xl/drawings/drawing1.xml'
    xml = entries[drawing_name].decode()
    xml = xml.replace('descr="Picture"', f'descr="Picture" title="{escape(description)}"')
    if prefixed:
        namespace = 'http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing'
        xml = xml.replace(f'xmlns="{namespace}"', f'xmlns:xdr="{namespace}"')
        for tag in ('wsDr', 'oneCellAnchor', 'from', 'col', 'colOff', 'row', 'rowOff',
                    'ext', 'pic', 'nvPicPr', 'cNvPr', 'cNvPicPr', 'blipFill', 'spPr', 'clientData'):
            xml = re.sub(rf'(<\/?){tag}(?=[\s/>])', rf'\1xdr:{tag}', xml)
    entries[drawing_name] = xml.encode()
    with zipfile.ZipFile(source, 'w') as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
    original = source.read_bytes()
    assert len(undescribed_images(entries)) == 1
    fixed, applied, skipped = remediate_office.remediate_office(source, in_scope=lambda sc: sc == '1.1.1')
    assert fixed is not None and any(description in change for change in applied)
    with zipfile.ZipFile(fixed) as archive:
        saved = {name: archive.read(name) for name in archive.namelist()}
    drawing = ET.fromstring(saved[drawing_name])
    assert next(node for node in drawing.iter() if node.tag.endswith("}cNvPr")).get("descr") == description
    assert undescribed_images(saved) == []
    assert saved['xl/worksheets/sheet1.xml'] == entries['xl/worksheets/sheet1.xml']
    assert saved['xl/media/image1.png'] == entries['xl/media/image1.png']
    reopened = load_workbook(fixed)
    assert reopened.active['A1'].value == 'Preserved data'
    assert reopened.active['B2'].value == '=SUM(1,2)'
    assert len(reopened.active._images) == 1
    assert source.read_bytes() == original
