"""Selected SCs gate execution, not just a score computed after unrelated rules ran."""
import sys
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor
import pytest
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'api'))
from assessment_selection import selection, selected_for_file, filter_findings, enabled
import scanner

HTML = '''<html><head><style>a{outline:none}</style></head><body>
<input required id="x"><video></video><audio></audio><img src="text-image.png">
<a href="x">click here</a><h1>x</h1><h3>y</h3><p style="color:#eee">Text</p>
</body></html>'''

@pytest.mark.parametrize('sc', ['1.1.1','1.2.1','1.2.2','1.2.3','1.3.1','1.3.2','1.3.4','1.3.5',
 '1.4.1','1.4.2','1.4.3','1.4.4','1.4.5','1.4.6','1.4.10','1.4.11','1.4.12',
 '2.4.1','2.4.2','2.4.3','2.4.4','2.4.6','2.4.7','2.4.9','2.5.3','2.5.8','3.1.1','3.1.4','3.3.2','4.1.2'])
def test_html_each_criterion_matches_only_its_unscoped_findings(tmp_path, sc):
    p = tmp_path/'test.html'; p.write_text(HTML)
    unscoped = scanner._analyse_html(p)['issues']
    with selection({sc}):
        assert scanner._analyse_html(p)['issues'] == filter_findings(unscoped)

@pytest.mark.parametrize('ext', ['.pdf','.docx','.pptx','.xlsx','.html'])
def test_title_only_skips_supplemental_detectors(tmp_path, ext, monkeypatch):
    import office_structure, textchecks
    class ForbiddenPath:
        def __getattr__(self, name):
            raise AssertionError('Unselected structural detector opened a file')
    with selection({'2.4.2'}):
        assert office_structure.checks_for(ForbiddenPath(), ext) == []
        assert textchecks.content_findings(object(), object()) == []


def test_selection_is_per_file_and_thread_local():
    scope = {'1.1.1': {'docx'}, '2.4.2': {'pdf'}}
    assert selected_for_file(scope, 'a.pdf') == {'2.4.2'}
    assert selected_for_file(scope, 'a.docx') == {'1.1.1'}
    assert selected_for_file(scope, 'a.html') == {'1.1.1', '2.4.2'}
    assert selected_for_file(scope, 'a.xlsx') == set()
    def run(sc):
        with selection({sc}):
            return enabled(sc), enabled('3.1.1')
    with ThreadPoolExecutor() as ex:
        assert list(ex.map(run, ['1.1.1','2.4.2'])) == [(True,False),(True,False)]
    assert enabled('3.1.1')


def test_scope_read_failure_never_runs_an_engine(monkeypatch, tmp_path):
    import core
    def fail(*args, **kwargs):
        raise RuntimeError('scope unavailable')
    monkeypatch.setattr(core.store, 'get_scan_scope', fail)
    with pytest.raises(RuntimeError, match='scope unavailable'):
        scanner.analyse_and_assess(tmp_path, 'x.pdf', scan_id='scope-test')


def test_pdf_allowlist_skips_unselected_engine_checks(tmp_path, monkeypatch):
    from pypdf import PdfWriter
    sys.path.insert(0, str(scanner.WP))
    from analysers.pdf_analyser import _RULES
    p=tmp_path/'x.pdf'; w=PdfWriter(); w.add_blank_page(width=200, height=200); w.write(p)
    called=[]
    for rule in _RULES:
        def check(*args, rid=rule.rule_id):
            called.append(rid)
            return []
        monkeypatch.setattr(rule, 'check', check)
    with selection({'2.4.2'}):
        result=scanner._analyse_pdf(p)
    assert result['succeeded'], result
    assert called and set(called) <= {'pdf.document-title','pdf.display-doc-title','pdf.missing-bookmarks'}


def test_vision_review_obeys_selection():
    from pdf_vision_review import findings_from_layouts
    with selection({'1.1.1'}):
        issues=findings_from_layouts([{'page':1,'description':'A scanned page'}])
    assert issues and {i['wcag'].split()[0] for i in issues} == {'1.1.1'}


@pytest.mark.parametrize('ext', ['docx', 'pptx', 'xlsx'])
def test_real_office_engine_honors_selected_rules(tmp_path, ext):
    if not scanner.CLI_DLL.exists():
        pytest.skip('Office engine build unavailable')
    p = tmp_path / ('sample.' + ext)
    if ext == 'docx':
        from docx import Document
        doc = Document(); doc.add_paragraph('Example'); doc.save(p)
    elif ext == 'pptx':
        from pptx import Presentation
        doc = Presentation(); doc.slides.add_slide(doc.slide_layouts[6]); doc.save(p)
    else:
        from openpyxl import Workbook
        doc = Workbook(); doc.active['A1'] = 'Example'; doc.save(p)
    with selection({'2.4.2'}):
        result = scanner._analyse_office(tmp_path)[p.name]
        assert result['succeeded'], result
        assert result['issues'] == filter_findings(result['issues'])
    with selection(set()):
        result = scanner._analyse_office(tmp_path)[p.name]
        assert result['succeeded'], result
        assert result['issues'] == []


def test_media_does_not_check_captions_when_only_audio_alternative_selected(monkeypatch):
    from types import SimpleNamespace
    from formats.av.detectors import captions
    monkeypatch.setattr(captions.media, 'probe', lambda p: SimpleNamespace(has_video=True, has_audio=True, has_captions=False))
    def forbidden(*args):
        raise AssertionError('Unselected captions check executed')
    monkeypatch.setattr(captions, '_sidecars', forbidden)
    with selection({'1.2.1'}):
        assert captions.detect('video.mp4') == []


def test_monolithic_scan_resolves_per_file_criterion_rules(tmp_path, monkeypatch):
    import core
    source = tmp_path/'source.html'; source.write_text(HTML)
    def listing(*args, scope_out, **kwargs):
        scope_out['scan_scope'] = {'2.4.2': ['docx']}
        return [{'name':'finance.html', 'path':str(source), 'parent_folder':'Finance'}]
    monkeypatch.setattr(scanner, '_list', listing)
    monkeypatch.setattr(core.store, 'list_scope_rules', lambda **kwargs: [
        {'rule_id':'r','selector':'folder','value':'Finance','codes':['1.1.1'],
         'priority':1,'is_override':True,'enabled':True}])
    monkeypatch.setattr(scanner, '_analyse_office', lambda *args, **kwargs: {})
    report = scanner.run_scan(source='local', ai_enabled=False)
    issues = report['files'][0]['issues']
    assert issues and {i['wcag'].split()[0] for i in issues} == {'1.1.1'}
    assert report['scope']['scope_rules'][0]['codes'] == ['1.1.1']
