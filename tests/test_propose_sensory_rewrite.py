"""WCAG 1.3.3 Sensory Characteristics PROPOSAL — an instruction that relies on shape / colour /
position ("click the large round green button in the bottom-right corner") gets an AI-drafted,
non-sensory rewrite the reviewer approves. Never auto-applied (which green button is a subjective
judgement), so this is always a HITL proposal.

These tests pin the FORMAT WIRING, which is the part that silently rots: the proposer is
format-agnostic (it runs on pii.extract_text's output), so a sensory instruction must surface
whether it sits in docx BODY PROSE or in an xlsx CELL — and a plain instruction must not. They
build real docx/xlsx bytes and drive them through the exact path remediation uses
(extract_text → propose_sensory_rewrite), with the local text model stubbed so the assertions
are deterministic and don't need a model installed.
"""
import io
import sys
import tempfile
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "api"))
import ai  # noqa: E402
import pii  # noqa: E402
import proposals  # noqa: E402

# An instruction that leans on colour / shape / position — the 1.3.3 failure.
SENSORY = ("To approve the request, click the large round green button located in the "
           "bottom-right corner of the screen.")
# The same instruction named by its control, carrying no sensory characteristic — the control.
PLAIN = ("To approve the request, click the Approve button and enter your employee ID.")


@pytest.fixture
def _model_stub(monkeypatch):
    """Make the local text model deterministic and count its calls.

    proposals.propose_sensory_rewrite imports `ai` lazily and calls ai.suggest_fix ONLY for a
    sentence the sensory regex already matched — so the call count is a direct probe of whether
    detection fired, independent of what the (stubbed) model returns.
    """
    calls: list[str] = []

    # Signature must match ai.suggest_fix's real one (rule_id, rule_name, level, filename,
    # detail="", image_bytes=None, style="", guidance="") — ADR 0021 added guidance after this
    # test was first written; propose_sensory_rewrite always passes it by keyword, so a stub
    # missing it raises TypeError, which propose_sensory_rewrite's bare `except Exception`
    # silently swallows into an empty proposal list rather than a visible test failure.
    def _suggest(sc, name, level, filename, detail="", image_bytes=None, style="", guidance=""):
        calls.append(detail)
        return {"suggestion": "Click the Approve button.", "model": "stub"}

    monkeypatch.setattr(ai, "model_is_available", lambda: True)
    monkeypatch.setattr(ai, "suggest_fix", _suggest)
    return calls


def _docx_bytes(paragraph: str) -> bytes:
    import docx as _docx
    d = _docx.Document()
    d.add_paragraph(paragraph)
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def _xlsx_bytes(cell_text: str) -> bytes:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws["A11"] = cell_text   # the sensory instruction lives in a CELL, not a paragraph
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def _proposals_for(data: bytes, ext: str) -> list[dict]:
    """Run the real remediation-side path: write bytes → extract_text → propose_sensory_rewrite."""
    with tempfile.TemporaryDirectory(prefix="acp-sensory-test-") as d:
        p = Path(d) / f"fixture.{ext}"
        p.write_bytes(data)
        text = pii.extract_text(p)
        return proposals.propose_sensory_rewrite(text, filename=p.name, ai_enabled=True)


@pytest.mark.parametrize("ext,build", [("docx", _docx_bytes), ("xlsx", _xlsx_bytes)])
def test_sensory_instruction_yields_a_proposal(ext, build, _model_stub):
    """A sensory instruction in docx body text AND in an xlsx cell each yields a 1.3.3 proposal."""
    props = _proposals_for(build(SENSORY), ext)
    assert len(props) >= 1, f"{ext}: expected a sensory-rewrite proposal, got none"
    assert _model_stub, f"{ext}: the sensory regex never matched — detection didn't reach the model"
    prop = props[0]
    assert prop["proposed_value"] == "Click the Approve button."
    # It is a PROPOSAL for human approval, never a validated/auto-applied fix.
    assert prop.get("validated") is not True
    assert "green button" in prop["before"]      # the offending prose is carried for review


@pytest.mark.parametrize("ext,build", [("docx", _docx_bytes), ("xlsx", _xlsx_bytes)])
def test_plain_instruction_yields_no_proposal(ext, build, _model_stub):
    """A plain instruction (no shape/colour/position) yields nothing — and never even calls the
    model, proving the gate is the sensory detector, not a proposer that always fires."""
    props = _proposals_for(build(PLAIN), ext)
    assert props == [], f"{ext}: a non-sensory instruction must not produce a proposal"
    assert _model_stub == [], f"{ext}: the model was consulted for non-sensory text"


def test_disabled_ai_yields_no_proposal(_model_stub):
    """With AI off the proposer degrades to a plain HITL deferral (no proposal, no model call)."""
    assert _proposals_for(_xlsx_bytes(SENSORY), "xlsx") and True  # sanity: stub is wired
    _model_stub.clear()
    props = proposals.propose_sensory_rewrite(SENSORY, filename="f.xlsx", ai_enabled=False)
    assert props == []
    assert _model_stub == []


@pytest.mark.parametrize('advance_consent', [True, False])
@pytest.mark.parametrize('call_field', ['ai_call_id', 'model_call_id', 'call_id', None])
def test_sensory_draft_retains_real_call_for_frozen_run_approval(isolated_store, monkeypatch, advance_consent, call_field):
    """A recorded successful draft enters the normal durable approval path.

    Exact model/call provenance is retained; consent-off stays pending, and
    approval alone neither writes an artifact nor grants verified credit.
    """
    from test_ai_standing_approval import seed, OWNER, SID, FILE, apply_jobs
    from ai_run_policy import run_context
    from ai_standing_approval import approve_file
    s = isolated_store
    job = seed(s, monkeypatch, enabled=advance_consent)
    with run_context(s, job['payload'], job) as ctx:
        call_id = s.record_ai_call(surface='synthetic', provider='synthetic', model='fixture-text',
            zone='local', latency_ms=0, ok=True, scan_id=SID, file=FILE)
        monkeypatch.setattr(ai, 'model_is_available', lambda: True)
        monkeypatch.setattr(ai, 'suggest_fix', lambda *a, **k: {
            'suggestion': 'Click the Approve button.', 'model': 'fixture-text',
            **({call_field: call_id} if call_field else {})})
        draft = proposals.propose_sensory_rewrite(SENSORY, filename=FILE)
        item_id = s.enqueue_proposals(SID, FILE, '1.3.3', draft)
        approve_file(s, ctx)
    row = s.get_hitl_item(item_id)
    can_approve = advance_consent and call_field is not None
    assert row['status'] == ('approved' if can_approve else 'pending')
    assert row['proposals'][0]['model'] == 'fixture-text'
    assert row['proposals'][0].get('model_call_id') == (call_id if call_field else None)
    assert not row['applied']
    assert len(apply_jobs(s)) == int(can_approve)
    with s._db.cursor() as cur:
        s._db.execute(cur, 'SELECT model_call_id,proposal_json FROM ai_proposal_snapshots WHERE item_id=%s', (item_id,))
        retained = s._db.fetchone(cur)
    assert retained['model_call_id'] == (call_id if call_field else None)
    assert retained['proposal_json'] and 'fixture-text' in retained['proposal_json']
