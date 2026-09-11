"""Real DOCX writer: retain deterministic edits, add AI edit, deliver exact saved bytes."""
import io
import json
from hashlib import sha256

import pytest
from docx import Document

import handlers
import document_wide_workflow
import document_wide_provider
from ai_run_policy import run_context
from ai_standing_approval import approve_file
from experiments.document_wide_ai.fixtures.make_fixtures import make_docx
from experiments.document_wide_ai.contracts.v1 import EditResponseEnvelope, ProposedEdit, CONTRACT_VERSION
from remediation_contribution import SOURCE, bind_assessed_input
from test_apply_approved_values import _Blob
import test_ai_standing_approval as base


def prepare(store, monkeypatch):
    monkeypatch.setattr(base, 'FILE', 'file.docx')
    original = make_docx(image_count=1)
    doc=Document(io.BytesIO(original)); doc.core_properties.title='Deterministic title'
    out=io.BytesIO(); doc.save(out); candidate=out.getvalue()
    job=base.seed(store,monkeypatch)
    sid,file=base.SID,base.FILE
    with store._db.cursor() as cur:
        store._db.execute(cur,'UPDATE file_records SET corrected_sha256=%s WHERE scan_id=%s', (sha256(candidate).hexdigest(),sid))
        store._db.execute(cur,'SELECT execution_id,input_snapshot_id FROM stage_executions WHERE scan_id=%s',(sid,))
        stage=store._db.fetchone(cur)
        baseline={'run_id':stage['execution_id'],'snapshot_id':stage['input_snapshot_id']}
        rows=[{'finding_id':'exact-assessed-image','file':file,'rule_id':'1.1.1','instance_key':'docx:drawing:1:paragraph:1'}]
        from remediation_contribution import freeze_baseline
        freeze_baseline(store._db,cur,base.OWNER,sid,baseline['run_id'],baseline['snapshot_id'],rows,[file])
    monkeypatch.setattr(store,'get_scan_scope',lambda *a:{'1.1.1':frozenset({'docx'})})
    monkeypatch.setattr(store,'scope_for_file',lambda s,f,scope:scope)
    monkeypatch.setattr(store,'list_finding_dispositions',lambda *a:rows)
    import sys
    blob=_Blob(candidate);monkeypatch.setitem(sys.modules,'blob',blob)
    return job,original,candidate,blob,baseline


def test_document_proposals_apply_to_current_candidate_and_release_exact_bytes(isolated_store,monkeypatch):
    store=isolated_store
    job,original,candidate,blob,baseline=prepare(store,monkeypatch)
    from proposals import Verification
    def verify(data,file,**kwargs):
        doc=Document(io.BytesIO(data))
        values=doc.element.xpath('//*[local-name()="docPr"]')
        return Verification(True,set() if values[0].get('descr')=='A green square' else {'1.1.1'})
    monkeypatch.setattr(handlers,'_verify_residual',verify)
    token=SOURCE.set(None)
    try:
        with run_context(store,job['payload'],job) as ctx:
            object.__setattr__(ctx,'policy',{**ctx.policy,'document_wide_ai':True})
            with store._db.cursor() as cur:
                store._db.execute(cur,'UPDATE ai_spending_run_policies SET policy_json=%s WHERE run_id=%s',(json.dumps(dict(ctx.policy)),ctx.run_id))
            assert bind_assessed_input(base.SID,base.FILE,original,original)
            def generate(request,**kwargs):
                f=request.manifest.findings[0]
                p=base.proposal(store,rule='1.1.1',locator='word/document.xml#Picture 1',value='A green square')
                with store._db.cursor() as cur:
                    store._db.execute(cur,"UPDATE ai_attempt_history SET status='drafted' WHERE run_id=%s",(ctx.run_id,))
                return {'envelope':EditResponseEnvelope(CONTRACT_VERSION,request.request_id,request.manifest.source_sha256,
                    (ProposedEdit('edit',(f.finding_id,),f.locator,'set_office_image_alt_text','A green square',''),),()),
                    'model':p['model'],'model_call_id':p['model_call_id']}
            monkeypatch.setattr(document_wide_provider,'generate_document',generate)
            document_wide_workflow.process_file(store,ctx)
            assert SOURCE.get()[2]==sha256(original).hexdigest()
            contributions=base.rows(store,'remediation_contribution_proposals')
            assert len(contributions)==1
            assert contributions[0]['source_sha256']==sha256(candidate).hexdigest()
            assert contributions[0]['assessment_revision']==baseline['snapshot_id']
            assert json.loads(contributions[0]['finding_ids_json'])==['exact-assessed-image']
            approve_file(store,ctx)
        jobs=base.apply_jobs(store)
        assert len(jobs)==1
        handlers._apply_approved_values(json.loads(jobs[0]['payload']),{})
        result=Document(io.BytesIO(blob.data))
        assert result.core_properties.title=='Deterministic title'
        assert result.element.xpath('//*[local-name()="docPr"]')[0].get('descr')=='A green square'
        assert blob.data!=candidate and candidate!=original
        from remediation_delivery import load_artifact
        delivered=load_artifact(owner=base.OWNER,scan_id=base.SID,file=base.FILE,
            expected_digest=sha256(blob.data).hexdigest(),download=blob.download_remediated)
        assert delivered==blob.data
        from unverified_changes import record_verification, blocks_certification
        assert record_verification(store,base.SID,base.FILE,blob.data,Verification(True,set()))==0
        assert blocks_certification(store,base.SID,base.FILE)
        receipts=base.rows(store,'ai_validation_outcomes')
        assert any(r.get('actual_source_sha256')==sha256(candidate).hexdigest() and r['outcome']=='could_not_verify' for r in receipts)
    finally:
        SOURCE.reset(token)


@pytest.mark.parametrize('tamper',['missing_receipt','owner','run','request','value','locator','finding','artifact','revision','cancelled'])
def test_candidate_metadata_cannot_spoof_server_lineage(isolated_store,monkeypatch,tamper):
    from copy import deepcopy
    from remediation_contribution import _document_candidate_hash
    store=isolated_store
    job,original,candidate,blob,baseline=prepare(store,monkeypatch)
    p={'document_wide_request_id':'request','source_sha256':sha256(candidate).hexdigest(),
       'assessment_revision':baseline['snapshot_id'],'locator':'word/document.xml#Picture 1',
       'proposed_value':'A green square','baseline_finding_ids':['exact-assessed-image']}
    receipt={'owner_id':base.OWNER,'run_id':baseline['run_id'],'request_id':'request','proposals':{'1.1.1':[deepcopy(p)]}}
    if tamper in {'owner','run','request'}:
        receipt[{'owner':'owner_id','run':'run_id','request':'request_id'}[tamper]]='other'
    if tamper=='value':p['proposed_value']='Different description'
    if tamper=='locator':p['locator']='word/document.xml#Picture 2'
    if tamper=='finding':p['baseline_finding_ids']=['another-finding']
    if tamper=='revision':p['assessment_revision']='other'
    if tamper!='missing_receipt':
        store.log_decision('system','document_wide.generated',scan_id=base.SID,file=base.FILE,detail=json.dumps(receipt))
    with store._db.cursor() as cur:
        if tamper=='artifact':store._db.execute(cur,'UPDATE file_records SET corrected_sha256=%s WHERE scan_id=%s',('other',base.SID))
        if tamper=='cancelled':store._db.execute(cur,"UPDATE stage_executions SET cancel_requested_at='now' WHERE scan_id=%s",(base.SID,))
        assert _document_candidate_hash(store._db,cur,owner=base.OWNER,scan_id=base.SID,
            run_id=baseline['run_id'],file=base.FILE,rule_id='1.1.1',proposal=p,revision=baseline['snapshot_id']) is None
