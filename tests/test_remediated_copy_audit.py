"""Exact byte/property evidence from actual generated Office/PDF packages."""
import hashlib
import io
import json
from pathlib import Path
import subprocess
import sys
import xml.etree.ElementTree as ET
import zipfile

import pytest
from remediated_copy_audit import audit, inspect

ROOT = Path(__file__).resolve().parents[1]


def package(extension):
    stream = io.BytesIO()
    # A real image establishes the real package namespaces and nonvisual property location.
    from PIL import Image
    image = io.BytesIO()
    Image.new('RGB', (2, 2), 'red').save(image, format='PNG')
    image.seek(0)
    if extension == 'docx':
        from docx import Document
        document = Document()
        document.add_paragraph('Source document')
        document.add_picture(image)
        document.save(stream)
    elif extension == 'pptx':
        from pptx import Presentation
        from pptx.util import Inches
        document = Presentation()
        document.slides.add_slide(document.slide_layouts[6]).shapes.add_picture(image, Inches(1), Inches(1))
        document.save(stream)
    else:
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image as SheetImage
        document = Workbook()
        document.active['A1'] = 'Source document'
        document.active.add_image(SheetImage(image), 'B2')
        document.save(stream)
    return stream.getvalue()


def rewrite(data, transform):
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as original, zipfile.ZipFile(output, 'w') as corrected:
        for entry in original.infolist():
            corrected.writestr(entry, transform(entry.filename, original.read(entry)))
    return output.getvalue()


@pytest.mark.parametrize('extension', ['docx', 'pptx', 'xlsx'])
def test_real_office_saved_alt_text_exact_properties(extension):
    original = package(extension)
    before = inspect(original, extension)
    assert before['readable'] and before['images']
    row = next(row for row in before['images'] if row['id'] is not None)
    identity = {key: row[key] for key in ('part', 'tag', 'id')}
    expected = {**identity, 'description': 'A red square'}

    def change(name, data):
        if name != row['part']:
            return data
        root = ET.fromstring(data)
        node = next(node for node in root.iter() if node.tag == row['tag'] and node.get('id') == row['id'])
        node.set('descr', expected['description'])
        return ET.tostring(root, encoding='utf-8', xml_declaration=True)

    corrected = rewrite(original, change)
    report = audit(original, corrected, extension,
                   expected_sha256=hashlib.sha256(corrected).hexdigest(), expected_images=[expected])
    assert report['bytes_changed'] and report['supplied_claims_verified']
    assert report['all_recorded_changes_verified'] == 'not established'
    assert report['accessibility_checker_pass'] == 'not established'
    assert not audit(original, original, extension, expected_images=[expected])['supplied_claims_verified']
    assert not audit(original, corrected, extension, expected_sha256='0' * 64)['supplied_claims_verified']
    assert not audit(original, corrected, extension)['supplied_claims_verified']
    assert not audit(original, corrected, extension, expected_images=[{**expected, 'id': 'missing'}])['supplied_claims_verified']
    assert not audit(original, corrected, extension, expected_images=[identity])['supplied_claims_verified']


def test_invalid_xml_missing_parts_and_ambiguous_property_identity():
    original = package('docx')
    row = inspect(original, 'docx')['images'][0]
    expected = {key: row[key] for key in ('part', 'tag', 'id', 'description')}
    invalid = rewrite(original, lambda name, data: b'<broken' if name == 'word/document.xml' else data)
    assert not inspect(invalid, 'docx')['readable']
    assert not audit(original, invalid, 'docx', expected_images=[expected])['supplied_claims_verified']

    def duplicate(name, data):
        if name != row['part']:
            return data
        root = ET.fromstring(data)
        node = next(node for node in root.iter() if node.tag == row['tag'])
        root.append(ET.fromstring(ET.tostring(node)))
        return ET.tostring(root)

    ambiguous = rewrite(original, duplicate)
    assert inspect(ambiguous, 'docx')['readable']
    assert not audit(original, ambiguous, 'docx', expected_images=[expected])['supplied_claims_verified']
    expected.pop('id')
    assert not audit(original, original, 'docx', expected_images=[expected])['supplied_claims_verified']
    output = io.BytesIO()
    with zipfile.ZipFile(output, 'w') as archive:
        archive.writestr('[Content_Types].xml', '<Types/>')
    assert 'missing' in inspect(output.getvalue(), 'docx')['errors'][0]


@pytest.mark.parametrize('encoding', ['utf-8', 'utf-16', 'utf-32'])
def test_xml_entity_declarations_rejected_before_parsing(encoding):
    original = package('docx')
    dangerous = '<?xml version="1.0"?><!DOCTYPE root [<!ENTITY value "expanded">]><root>&value;</root>'
    modified = rewrite(original, lambda name, data: dangerous.encode(encoding) if name == 'word/document.xml' else data)
    result = inspect(modified, 'docx')
    assert not result['readable']
    assert result['errors'] == ['XML entity declarations are not supported.']


def pdf(encrypted=False):
    from pypdf import PdfWriter
    document = PdfWriter()
    document.add_blank_page(width=100, height=100)
    document.add_metadata({'/Title': 'Saved corrected copy'})
    if encrypted:
        document.encrypt('test-password')
    stream = io.BytesIO()
    document.write(stream)
    return stream.getvalue()


def test_real_pdf_readability_is_not_accessibility_checker_pass():
    data = pdf()
    result = inspect(data, 'pdf')
    assert result['readable'] and result['pages'] == 1
    assert result['accessibility_checker_pass'] == 'not established'
    assert audit(data, data, 'pdf', expected_sha256=hashlib.sha256(data).hexdigest())['supplied_claims_verified']
    assert not inspect(pdf(encrypted=True), 'pdf')['readable']
    assert not inspect(b'not a PDF', 'pdf')['readable']


def test_cli_writes_conservative_report_and_failure_exit_for_unestablished_claim(tmp_path):
    original, corrected = tmp_path / 'original.docx', tmp_path / 'corrected.docx'
    data = package('docx')
    original.write_bytes(data)
    corrected.write_bytes(data)
    output, expected = tmp_path / 'report.json', tmp_path / 'expected.json'
    expected.write_text(json.dumps({'corrected_sha256': hashlib.sha256(data).hexdigest()}))
    command = [sys.executable, str(ROOT / 'scripts/audit_remediated_copy.py'), str(original), str(corrected), '--output', str(output)]
    result = subprocess.run([*command, '--expected', str(expected)], capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
    report = json.loads(output.read_text())
    assert report['supplied_claims_verified'] and not report['bytes_changed']
    assert report['all_recorded_changes_verified'] == 'not established'
    result = subprocess.run(command, capture_output=True, text=True)
    assert result.returncode == 1
    assert not json.loads(output.read_text())['supplied_claims_verified']


def test_valid_xml_with_wrong_office_root_and_duplicate_zip_parts_are_not_readable():
    original = package('docx')
    wrong = rewrite(original, lambda name, data: b'<root/>' if name == 'word/document.xml' else data)
    assert not inspect(wrong, 'docx')['readable']
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(original)) as source, zipfile.ZipFile(output, 'w') as target:
        for entry in source.infolist():
            target.writestr(entry, source.read(entry))
        with pytest.warns(UserWarning, match='Duplicate'):
            target.writestr('word/document.xml', source.read('word/document.xml'))
    assert 'Duplicate package entries' in inspect(output.getvalue(), 'docx')['errors'][0]
