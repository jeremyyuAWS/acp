"""New releases reassess the exact saved candidate, without replacing source findings."""
import hashlib
import io
import json
import zipfile

import pytest
from docx import Document
from PIL import Image

from release_artifacts import ReleaseArtifactError
from release_candidate_assessment import assess_candidate

SID, OWNER, FILE = 'saved-assess', 'owner', 'actual.docx'


@pytest.fixture
def candidate(isolated_store, monkeypatch):
    import core
    import blob
    store = isolated_store
    store.init_scan_run(SID, 'local', 1, '2026-09-12T00:00:00Z', 'rubric', 'hash', owner=OWNER)
    store.save_file_result(SID, {'file': FILE, 'engine': 'office', 'status': 'analysed',
        'score': 100, 'compliant': 1, 'issues': [], 'errors': [], 'skipped_rules': 0},
        '2026-09-12T00:00:00Z')
    monkeypatch.setattr(core, 'store', store)
    monkeypatch.setattr(store, 'get_scan_scope', lambda *args, **kwargs: {'1.1.1': ['docx']})
    monkeypatch.setattr(store, 'scope_for_file', lambda sid, name, scope: scope)
    image = io.BytesIO()
    Image.new('RGB', (80, 60), 'blue').save(image, format='PNG')
    image.seek(0)
    document = Document()
    document.add_paragraph('Original content')
    document.add_picture(image)
    output = io.BytesIO()
    document.save(output)
    state = {'bytes': output.getvalue()}
    monkeypatch.setattr(blob, 'download_remediated', lambda *args: state['bytes'])
    def save(data):
        state['bytes'] = data
        digest = hashlib.sha256(data).hexdigest()
        at = store.record_remediation(SID, FILE, blob_url='blob://corrected', corrected_sha256=digest)
        return digest, at
    return store, state, save


def assess(candidate, *, remaining=False):
    store, state, save = candidate
    digest, at = save(state['bytes'])
    return assess_candidate(store, SID, FILE, OWNER, digest, at,
                            allow_remaining_issues=remaining)


def test_real_saved_docx_missing_alt_is_blocked_by_fresh_assessment_despite_old_clear_record(candidate):
    store, _, _ = candidate
    before = store.get_scan(SID, owner=OWNER)['files']
    with pytest.raises(ReleaseArtifactError) as error:
        assess(candidate)
    assert error.value.category == 'release_assessment_remaining'
    assert store.get_scan(SID, owner=OWNER)['files'][0]['issues'] == before[0]['issues'] == []
    from release_candidate_assessment import saved_assessment
    evidence = saved_assessment(store, SID, OWNER, FILE, store.get_file_record(SID, FILE)['corrected_sha256'])
    assert evidence['remaining_criteria'] == ['1.1.1']
    assert evidence['remaining_issues'] and evidence['assessment_ok'] is True


def test_real_saved_described_docx_passes_selected_checks_and_evidence_matches_bytes(candidate):
    store, state, _ = candidate
    with zipfile.ZipFile(io.BytesIO(state['bytes'])) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    entries['word/document.xml'] = entries['word/document.xml'].replace(
        b'<wp:docPr ', b'<wp:docPr descr="Quarterly revenue chart" ')
    output = io.BytesIO()
    with zipfile.ZipFile(output, 'w') as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
    state['bytes'] = output.getvalue()
    evidence = assess(candidate)
    assert evidence['assessment_ok'] is True and evidence['remaining_criteria'] == []
    assert evidence['artifact_sha256'] == hashlib.sha256(state['bytes']).hexdigest()
    assert evidence['assessment_scope'] == {'1.1.1': ['docx']}


def test_explicit_remaining_keeps_real_findings_without_claiming_clear(candidate):
    evidence = assess(candidate, remaining=True)
    assert evidence['assessment_ok'] is True and evidence['remaining_criteria'] == ['1.1.1']
    assert evidence['allow_remaining_issues'] is True and evidence['remaining_issues']


def test_actual_corrupted_saved_archive_is_unknown_and_strict_fails(candidate):
    _, state, _ = candidate
    state['bytes'] = b'not an Office file'
    with pytest.raises(ReleaseArtifactError) as error:
        assess(candidate)
    assert error.value.category == 'corrected_copy_unreadable'
    with pytest.raises(ReleaseArtifactError) as remaining_error:
        assess(candidate, remaining=True)
    assert remaining_error.value.category == 'corrected_copy_unreadable'


def test_hash_change_and_owner_are_rejected_before_scanning(candidate):
    store, state, save = candidate
    digest, at = save(state['bytes'])
    state['bytes'] = b'changed blob'
    with pytest.raises(ReleaseArtifactError, match='bytes changed'):
        assess_candidate(store, SID, FILE, OWNER, digest, at)
    with pytest.raises(ReleaseArtifactError, match='owner'):
        assess_candidate(store, SID, FILE, 'other-owner', digest, at)


def test_valid_copy_unknown_checker_is_strict_blocker_and_honest_explicit_remaining(candidate, monkeypatch):
    import proposals
    monkeypatch.setattr(proposals, 'verify_residual', lambda *args, **kwargs:
        proposals.Verification(False, reason='engine unavailable'))
    with pytest.raises(ReleaseArtifactError) as error:
        assess(candidate)
    assert error.value.category == 'release_assessment_unavailable'
    evidence = assess(candidate, remaining=True)
    assert evidence['assessment_ok'] is False and evidence['remaining_issues'] is None
    assert evidence['reason'] == 'engine unavailable'


def test_publication_route_never_records_a_publish_when_real_saved_pdf_has_findings(candidate, monkeypatch):
    import core
    import blob
    import publish
    from routes import scans
    from test_synchronous_release_stage import _RouteStore, _request
    _, state, _ = candidate
    route_store = _RouteStore('local')
    digest = hashlib.sha256(state['bytes']).hexdigest()
    old_record = route_store.get_file_record
    route_store.get_file_record = lambda sid, filename: {**old_record(sid, filename),
        'corrected_sha256': digest}
    route_store.get_scan_scope = lambda *args, **kwargs: {'1.1.1': ['docx']}
    route_store.scope_for_file = lambda sid, name, scope: scope
    route_store.log_decision = lambda *args, **kwargs: route_store.events.append(('assessment', kwargs))
    monkeypatch.setattr(core, 'store', route_store)
    monkeypatch.setattr(publish, 'remediated_content_digest', lambda *args: digest)
    # Store fixture's stable name is one.pdf; use a real PDF candidate with a
    # missing document language/title to prove the local release boundary itself.
    from pypdf import PdfWriter
    pdf = PdfWriter()
    pdf.add_blank_page(width=200, height=200)
    output = io.BytesIO()
    pdf.write(output)
    state['bytes'] = output.getvalue()
    digest = hashlib.sha256(state['bytes']).hexdigest()
    route_store.get_scan_scope = lambda *args, **kwargs: {'3.1.1': ['pdf'], '2.4.2': ['pdf']}
    response = scans.publish_files('scan-1', _request(), {'files': ['one.pdf']})
    assert response['published'][0]['status'] == 'failed'
    assert response['published'][0]['failure_category'] in {'release_assessment_remaining', 'release_assessment_unavailable'}
    assert 'publish' not in route_store.events
    assert not any(isinstance(event, tuple) and event[0] == 'receipt' for event in route_store.events)


def test_unapplied_approved_values_are_never_bypassed_by_remaining_mode(candidate, monkeypatch):
    store, _, _ = candidate
    monkeypatch.setattr(store, 'count_unapplied_approved_values', lambda *args: 1)
    with pytest.raises(ReleaseArtifactError) as error:
        assess(candidate, remaining=True)
    assert error.value.category == 'approved_changes_unapplied'


def test_saved_assessment_requires_owned_exact_artifact(candidate):
    from release_candidate_assessment import saved_assessment
    store, _, _ = candidate
    evidence = assess(candidate, remaining=True)
    assert saved_assessment(store, SID, OWNER, FILE, evidence['artifact_sha256']) == evidence
    assert saved_assessment(store, SID, 'other', FILE, evidence['artifact_sha256']) is None
    assert saved_assessment(store, SID, OWNER, FILE, '0' * 64) is None


def test_unmapped_raw_finding_blocks_strict_and_remains_in_partial_evidence(candidate, monkeypatch):
    import proposals
    raw = {'status': 'analysed', 'issues': [{'wcag': 'unknown', 'detail': 'Unmapped recorded issue'}]}
    monkeypatch.setattr(proposals, 'verify_residual', lambda *args, **kwargs:
        proposals.Verification(True, assessment=raw))
    with pytest.raises(ReleaseArtifactError) as error:
        assess(candidate)
    assert error.value.category == 'release_assessment_remaining'
    evidence = assess(candidate, remaining=True)
    assert evidence['remaining_criteria'] == [] and evidence['remaining_issues'] == raw['issues']


def test_missing_actual_assessment_cannot_attest_strict_clear(candidate, monkeypatch):
    import proposals
    monkeypatch.setattr(proposals, 'verify_residual', lambda *args, **kwargs:
        proposals.Verification(True))
    with pytest.raises(ReleaseArtifactError) as error:
        assess(candidate)
    assert error.value.category == 'release_assessment_unavailable'
