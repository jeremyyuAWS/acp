"""Model requests for missing input must never become document instructions."""
import io
import pytest
import ai
import proposals
from apply_text_values import apply_sensory_rewrite
from ai_run_policy import run_context
from ai_standing_approval import approve_file
import test_ai_standing_approval as standing
from sensory_rewrite_output import sensory_non_answer_reason, NON_ANSWER_REASON

INSTRUCTION='Refer to the cold therapy unit instructions shown below.'
REFUSAL="I don't see the original cold therapy unit instruction in your message. Please provide the step text that needs rewriting, and I'll revise it to remove sensory characteristic references."


def word(text):
    from docx import Document
    doc=Document();doc.add_paragraph(text)
    out=io.BytesIO();doc.save(out);return out.getvalue()


def office(text,ext):
    if ext=='docx':
        return word(text)
    out=io.BytesIO()
    if ext=='xlsx':
        from openpyxl import Workbook
        workbook=Workbook();workbook.active['A1']=text;workbook.save(out)
    else:
        from pptx import Presentation
        from pptx.util import Inches
        deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1),Inches(1),Inches(6),Inches(2)).text=text
        deck.save(out)
    return out.getvalue()


def test_proposer_does_not_offer_exact_observed_non_answer(monkeypatch):
    monkeypatch.setattr(ai,'model_is_available',lambda:True)
    monkeypatch.setattr(ai,'suggest_fix',lambda *a,**k:{'suggestion':REFUSAL,'model':'synthetic','model_call_id':'call'})
    assert proposals.propose_sensory_rewrite(INSTRUCTION,filename='instruction.docx')==[]


@pytest.mark.parametrize('ext',['docx','pptx','xlsx'])
@pytest.mark.parametrize('value',[REFUSAL,'"'+REFUSAL+'"','```text\n'+REFUSAL+'\n```'])
def test_writer_preserves_original_bytes_when_approved_value_is_model_non_answer(ext,value):
    original=office(INSTRUCTION,ext)
    result,applied,unresolved=apply_sensory_rewrite(original,ext,{INSTRUCTION:value})
    assert result==original and not applied and unresolved==[INSTRUCTION]


def test_standing_approval_does_not_admit_model_non_answer(isolated_store,monkeypatch):
    store=isolated_store;job=standing.seed(store,monkeypatch)
    with run_context(store,job['payload'],job) as ctx:
        p=standing.proposal(store,rule='1.3.3',locator=INSTRUCTION[:60],value=REFUSAL,review=None)
        item=store.enqueue_proposals(standing.SID,standing.FILE,'1.3.3',[p])
        approve_file(store,ctx)
    assert store.get_hitl_item(item)['status']=='pending'
    assert not standing.apply_jobs(store)


STEPS='Refer to the cold therapy unit instructions shown below (steps 1.1 to 1.6).'
REWRITE='Refer to steps 1.1 through 1.6 in the cold therapy unit instructions.'


def test_proposer_supplies_full_decimal_step_reference(monkeypatch):
    seen=[]
    monkeypatch.setattr(ai,'model_is_available',lambda:True)
    def suggest(*a,**k):
        seen.append(k['detail']);return {'suggestion':REWRITE,'model':'synthetic'}
    monkeypatch.setattr(ai,'suggest_fix',suggest)
    drafts=proposals.propose_sensory_rewrite('Clinical intro. '+STEPS+' Keep the dressing dry.')
    assert seen==[STEPS] and len(drafts)==1
    assert drafts[0]['before']==STEPS


def test_writer_replaces_full_step_sentence_across_runs_preserving_neighbors():
    from docx import Document
    doc=Document();paragraph=doc.add_paragraph()
    paragraph.add_run('Clinical intro. ').italic=True
    paragraph.add_run(STEPS[:55]).bold=True
    paragraph.add_run(STEPS[55:])
    paragraph.add_run(' Keep the dressing dry.').italic=True
    output=io.BytesIO();doc.save(output);original=output.getvalue()
    result,applied,unresolved=apply_sensory_rewrite(original,'docx',{STEPS[:60]:REWRITE})
    saved=Document(io.BytesIO(result)).paragraphs[0]
    assert saved.text=='Clinical intro. '+REWRITE+' Keep the dressing dry.'
    assert not unresolved and applied[0]['before']==STEPS
    assert saved.runs[0].text=='Clinical intro. ' and saved.runs[0].italic
    assert saved.runs[-1].text==' Keep the dressing dry.' and saved.runs[-1].italic


@pytest.mark.parametrize('value',[REFUSAL, "I don’t see the original instruction in your message.",
    "I'm sorry, but I can't assist with rewriting this instruction.",
    'As an AI language model, I cannot rewrite this.',
    'Please provide the original text.', 'Could you share your source instruction?'])
def test_obvious_model_non_answers_are_rejected(value):
    assert sensory_non_answer_reason(value)==NON_ANSWER_REASON


@pytest.mark.parametrize('value',['Please provide your patient ID to the nurse.',
    'Select the field labeled "Please provide the original text" and enter the referral.',
    'If you cannot see the screen, ask the nurse for assistance.', REWRITE])
def test_legitimate_instructions_are_not_banned(value):
    assert sensory_non_answer_reason(value) is None


@pytest.mark.parametrize('wrapper',[lambda text:'"'+text+'"',lambda text:"'"+text+"'",
    lambda text:'“'+text+'”',lambda text:'```text\n'+text+'\n```',
    lambda text:'```\n'+text+'\n```',lambda text:'~~~plaintext\n'+text+'\n~~~',
    lambda text:'```text\n"'+text+'"\n```'])
def test_whole_transport_wrappers_cannot_hide_non_answer(wrapper):
    assert sensory_non_answer_reason(wrapper(REFUSAL))==NON_ANSWER_REASON
    assert sensory_non_answer_reason(wrapper('Please provide your patient ID to the nurse.')) is None
    assert sensory_non_answer_reason(wrapper('Select the field labeled "Please provide the original text".')) is None


def test_valid_quoted_instruction_is_preserved_exactly_by_writer():
    original=word(INSTRUCTION)
    valid='"Please provide your patient ID to the nurse."'
    result,applied,unresolved=apply_sensory_rewrite(original,'docx',{INSTRUCTION:valid})
    from docx import Document
    assert Document(io.BytesIO(result)).paragraphs[0].text==valid
    assert len(applied)==1 and not unresolved


def test_mixed_writer_batch_never_inserts_refusal_but_retains_valid_rewrite():
    original=word(INSTRUCTION+' Press the round button to continue.')
    result,applied,unresolved=apply_sensory_rewrite(original,'docx',{
        INSTRUCTION:REFUSAL,'Press the round button to continue.':'Press Continue.'})
    from docx import Document
    assert Document(io.BytesIO(result)).paragraphs[0].text==INSTRUCTION+' Press Continue.'
    assert len(applied)==1 and unresolved==[INSTRUCTION]
