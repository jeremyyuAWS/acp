"""The ACP stamp is visible in Office's ordinary Summary Comments properties."""
from datetime import datetime, timezone
from io import BytesIO

import pytest

from output_provenance import stamp_output


def office_document(ext, comments):
    if ext == "docx":
        from docx import Document
        document = Document()
        document.add_paragraph("Original document text")
        props = document.core_properties
        props.comments = comments
    elif ext == "xlsx":
        from openpyxl import Workbook
        document = Workbook()
        document.active["A1"] = "Original document text"
        props = document.properties
        props.description = comments
    else:
        from pptx import Presentation
        document = Presentation()
        slide = document.slides.add_slide(document.slide_layouts[1])
        slide.shapes.title.text = "Original document text"
        props = document.core_properties
        props.comments = comments
    props.title = "Customer title"
    if ext == "xlsx":
        props.creator = "Customer author"
    else:
        props.author = "Customer author"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def read_properties(ext, data):
    if ext == "docx":
        from docx import Document
        document = Document(BytesIO(data))
        props = document.core_properties
        return props.comments, props.title, props.author, document.paragraphs[0].text
    if ext == "xlsx":
        from openpyxl import load_workbook
        document = load_workbook(BytesIO(data))
        props = document.properties
        return props.description, props.title, props.creator, document.active["A1"].value
    from pptx import Presentation
    document = Presentation(BytesIO(data))
    props = document.core_properties
    return props.comments, props.title, props.author, document.slides[0].shapes.title.text


@pytest.mark.parametrize("ext", ["docx", "xlsx", "pptx"])
@pytest.mark.parametrize("comments", ["", "generated using python-pptx\nCustomer's existing notes"])
def test_visible_stamp_preserves_summary_and_content_when_restamped(ext, comments, monkeypatch):
    monkeypatch.setenv("ACP_BUILD_VERSION", "2026.9.10.3")
    first = stamp_output(office_document(ext, comments), "output." + ext,
                         now=datetime(2026, 9, 10, 10, 11, 12, tzinfo=timezone.utc))
    first_comment, title, author, content = read_properties(ext, first)
    first_line = "Mova-io ACP 2026.9.10.3 · 2026-09-10T10:11:12Z"
    assert first_comment.splitlines()[0] == first_line
    assert first_comment[len(first_line):].strip() == comments
    assert (title, author, content) == ("Customer title", "Customer author", "Original document text")

    monkeypatch.setenv("ACP_BUILD_VERSION", "2026.9.10.4")
    second = stamp_output(first, "output." + ext,
                          now=datetime(2026, 9, 11, 12, 13, 14, tzinfo=timezone.utc))
    second_comment, title, author, content = read_properties(ext, second)
    second_line = "Mova-io ACP 2026.9.10.4 · 2026-09-11T12:13:14Z"
    assert second_comment.splitlines()[0] == second_line
    assert second_comment[len(second_line):].strip() == comments
    assert first_line not in second_comment
    assert (title, author, content) == ("Customer title", "Customer author", "Original document text")
