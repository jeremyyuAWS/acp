"""Real DOCX representations that the detector and writer previously disagreed on."""
import zipfile
from pathlib import Path
import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from remediate_office import _remediate_docx_structure


def fixture(path, *, outline=True, wrapped=True):
    doc = Document()
    heading = doc.add_paragraph('Policy title', 'Heading 2')
    if outline:
        level = OxmlElement('w:outlineLvl'); level.set(qn('w:val'), '1')
        heading._p.get_or_add_pPr().append(level)
    table = doc.add_table(rows=2, cols=2)
    for n,row in enumerate(table.rows):
        row.cells[0].text='Column' if n==0 else 'Value'
    if wrapped:
        for row in list(table._tbl.findall(qn('w:tr'))):
            table._tbl.remove(row)
            control=OxmlElement('w:sdt'); content=OxmlElement('w:sdtContent')
            content.append(row);control.append(content);table._tbl.append(control)
    doc.save(path)


def test_outline_override_and_wrapped_table_rows_are_actually_written(tmp_path):
    source=tmp_path/'source.docx';fixture(source)
    with zipfile.ZipFile(source) as z: entries={n:z.read(n) for n in z.namelist()}
    changes=[]
    _remediate_docx_structure(entries, changes, in_scope=lambda sc: sc == '1.3.1')
    root=etree.fromstring(entries['word/document.xml'])
    assert root.find('.//'+qn('w:outlineLvl')).get(qn('w:val'))=='0'
    assert root.find('.//'+qn('w:tblHeader')) is not None
    assert len([d for d in changes if d['rule_id']=='1.3.1'])==2


def scan(path):
    import scanner
    if not Path(scanner.CLI_DLL).exists():
        pytest.skip('Build the Office CLI to run the real detector round trip')
    result=scanner._analyse_office(path.parent)[path.name]
    assert result['succeeded'] and not result['errors']
    return [i for i in result['issues'] if i['wcag'] in {'SC_1_3_1','1.3.1'}]


def test_real_detector_clears_heading_override_and_content_control_table(tmp_path):
    source=tmp_path/'source.docx';fixture(source)
    before=scan(source)
    assert {i['detail'] for i in before} >= {'No H1 heading found','Table has no header row'}
    with zipfile.ZipFile(source) as z: entries={n:z.read(n) for n in z.namelist()}
    changes=[]
    _remediate_docx_structure(entries, changes, in_scope=lambda sc: sc=='1.3.1')
    output=tmp_path/'fixed.docx'
    with zipfile.ZipFile(output,'w') as z:
        for name,data in entries.items():z.writestr(name,data)
    assert scan(output)==[]


def test_nested_table_rows_are_not_counted_as_outer_layout_rows(tmp_path):
    doc=Document();outer=doc.add_table(rows=1,cols=1)
    inner=outer.cell(0,0).add_table(rows=2,cols=2)
    inner.cell(0,0).text='Heading';inner.cell(1,0).text='Data'
    source=tmp_path/'nested.docx';doc.save(source)
    findings=scan(source)
    assert len([i for i in findings if i['detail']=='Table has no header row'])==1


def test_excluded_structure_sc_does_not_write_heading_or_wrapped_table(tmp_path):
    source=tmp_path/'source.docx';fixture(source)
    with zipfile.ZipFile(source) as z: entries={n:z.read(n) for n in z.namelist()}
    original=etree.fromstring(entries['word/document.xml'])
    _remediate_docx_structure(entries, in_scope=lambda sc: False)
    output=etree.fromstring(entries['word/document.xml'])
    assert etree.tostring(output)==etree.tostring(original)
