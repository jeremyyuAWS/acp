"""Document generation must enter normal review without claiming applied findings."""
import hashlib
import sys
from contextlib import nullcontext
from types import SimpleNamespace

import document_wide_workflow as workflow
from experiments.document_wide_ai.contracts.v1 import (
    CONTRACT_VERSION, DocumentContextManifest, DocumentFormat, Finding, Locator,
    AllowedOperation, ProposedEdit, EditResponseEnvelope,
)


def fixture(monkeypatch):
    data = b'current corrected file'
    digest = hashlib.sha256(data).hexdigest()
    locator = Locator(DocumentFormat.PDF, 0, None, 'pdf:field:0:0', 'fingerprint')
    manifest = DocumentContextManifest(CONTRACT_VERSION, '1', '1', 'file', DocumentFormat.PDF,
        digest, 'revision', ('4.1.2',), (Finding('real-finding', 'pdf.form', '4.1.2', locator),),
        (AllowedOperation('set_pdf_field_accessible_name', DocumentFormat.PDF),), 'Form context')
    return data, digest, locator, manifest


def setup(monkeypatch, *, status='pending', stale=False):
    data, digest, locator, manifest = fixture(monkeypatch)
    context = SimpleNamespace(file='file.pdf', owner_id='owner', scan_id='scan', run_id='run',
                              enabled=True, policy={'document_wide_ai': True})
    calls, logs, queued = [], [], []
    current = {'corrected_sha256': digest}
    store = SimpleNamespace(
        remediation_source_revision=lambda sid: 'revision',
        get_file_record=lambda *a: current,
        transaction=lambda: nullcontext(),
        list_finding_dispositions=lambda *a:[{'finding_id':'real-finding','file':'file.pdf','rule_id':'4.1.2'}],
        list_hitl_queue=lambda **kw: [{'file':'file.pdf', 'rule_id':'4.1.2', 'status':status}],
        enqueue_proposals=lambda *a, **kw: queued.append((a,kw)),
        log_decision=lambda *a, **kw: logs.append((a,kw)),
    )
    def generate(request, **kwargs):
        calls.append(request)
        if stale:
            current['corrected_sha256'] = 'different'
        return {'envelope': EditResponseEnvelope(CONTRACT_VERSION, request.request_id, digest,
            (ProposedEdit('edit', ('real-finding',), locator, 'set_pdf_field_accessible_name',
                          'Patient name', '', 'Visible label'),), ()),
            'model': 'model', 'model_call_id': 'actual-call'}
    monkeypatch.setitem(sys.modules,'document_wide_manifest',SimpleNamespace(build_manifest=lambda *a:manifest, package_images=lambda *a:{}))
    monkeypatch.setitem(sys.modules,'document_wide_provider',SimpleNamespace(generate_document=generate))
    monkeypatch.setitem(sys.modules,'blob',SimpleNamespace(download_remediated=lambda *a:data))
    monkeypatch.setattr(workflow, '_saved', lambda *a:None)
    return store,context,calls,logs,queued


def test_default_and_unsupported_are_off():
    assert not workflow.enabled(None,'file.pdf')
    assert not workflow.enabled(SimpleNamespace(policy={}),'file.pdf')
    assert not workflow.enabled(SimpleNamespace(policy={'document_wide_ai':True}),'file.xlsx')


def test_local_or_zero_budget_never_calls_provider(monkeypatch):
    store,ctx,calls,logs,queued=setup(monkeypatch)
    ctx.enabled=False
    workflow.process_file(store,ctx)
    assert not calls and not queued
    assert 'positive spending' in logs[0][1]['detail']


def test_valid_proposal_preserves_lineage_and_is_not_verified(monkeypatch):
    store,ctx,calls,logs,queued=setup(monkeypatch)
    workflow.process_file(store,ctx)
    assert len(calls)==len(queued)==1
    args,kwargs=queued[0]
    assert args[3][0]['finding_ids']==['real-finding']
    assert args[3][0]['model_call_id']=='actual-call'
    assert kwargs=={'validated':False,'finding_count':1}
    assert args[3][0]['baseline_finding_ids']==['real-finding']


def test_stale_artifact_does_not_enqueue(monkeypatch):
    store,ctx,calls,logs,queued=setup(monkeypatch,stale=True)
    workflow.process_file(store,ctx)
    assert calls and not queued
    assert 'changed during generation' in logs[-1][1]['detail']


def test_existing_decision_is_never_overwritten(monkeypatch):
    store,ctx,calls,logs,queued=setup(monkeypatch,status='approved')
    workflow.process_file(store,ctx)
    assert not queued
    assert 'decision was preserved' in logs[-1][1]['detail']


def test_cached_response_does_not_spend_again(monkeypatch):
    store,ctx,calls,logs,queued=setup(monkeypatch)
    monkeypatch.setattr(workflow,'_saved',lambda *a:{'proposals':{}})
    workflow.process_file(store,ctx)
    assert not calls


def test_hook_runs_after_durable_phase_before_standing_approval(monkeypatch):
    import ai_run_policy
    import ai_standing_approval
    import handlers
    store,ctx,calls,logs,queued=setup(monkeypatch)
    ctx.deferred=[]
    ctx.policy['auto_approve_ai']=True
    order=[]
    monkeypatch.setattr(handlers.core,'store',store)
    monkeypatch.setattr(ai_run_policy,'run_context',lambda *a:nullcontext(ctx))
    monkeypatch.setattr(handlers,'_remediate_file_with_policy',lambda *a:order.append('durable'))
    monkeypatch.setattr(workflow,'process_file',lambda *a:order.append('generate'))
    monkeypatch.setattr(ai_standing_approval,'approve_file',lambda *a:order.append('approve'))
    handlers._remediate_file({}, {})
    assert order==['durable','generate','approve']


def test_document_mode_suppresses_only_supported_criterion(monkeypatch):
    import ai_run_policy
    import handlers
    from assessment_selection import enabled as selected
    ctx=SimpleNamespace(policy={'document_wide_ai':True})
    store=SimpleNamespace(get_scan_scope=lambda *a:{}, scope_for_file=lambda *a:{'1.1.1':['docx'], '1.3.1':['docx']})
    monkeypatch.setattr(handlers.core,'store',store)
    monkeypatch.setattr(ai_run_policy,'optional_current_run_context',lambda:ctx)
    checks=[]
    monkeypatch.setattr(handlers,'_propose_text_findings_selected',lambda *a:checks.extend([selected('1.1.1'),selected('1.3.1')]))
    handlers._propose_text_findings('scan','file.docx',b'',True)
    assert checks==[False,True]


def test_ai_only_document_still_stores_working_copy(monkeypatch):
    import io
    import pytest
    import handlers
    import ai_run_policy
    import activity
    import output_provenance
    import remediate_office
    from docx import Document
    doc=Document(); doc.add_paragraph('Original content')
    stream=io.BytesIO(); doc.save(stream); original=stream.getvalue()
    ctx=SimpleNamespace(policy={'document_wide_ai':True})
    store=SimpleNamespace(is_shadowed_output=lambda *a:False,get_ai_enabled=lambda:True,
        list_auto_fail_rules=lambda *a:[],get_scan=lambda *a:{'run':{'owner_email':'owner'}},
        attach_hitl_evidence=lambda *a:None)
    monkeypatch.setattr(handlers.core,'store',store)
    monkeypatch.setattr(ai_run_policy,'optional_current_run_context',lambda:ctx)
    monkeypatch.setattr(handlers,'_phase',lambda *a:None)
    monkeypatch.setattr(activity,'record',lambda *a,**kw:None)
    monkeypatch.setattr(handlers,'_remediation_source_bytes',lambda *a:(original,None))
    monkeypatch.setattr(handlers,'_propose_text_findings',lambda *a:None)
    monkeypatch.setattr(handlers,'_propose_form_fields',lambda *a:None)
    monkeypatch.setattr(handlers,'_record_applied_fixes',lambda *a:None)
    monkeypatch.setattr(handlers,'_rem_event',lambda *a,**kw:None)
    monkeypatch.setattr(handlers,'_remediation_scope',lambda *a:lambda sc:False)
    monkeypatch.setattr(remediate_office,'remediate_office',lambda *a,**kw:(None,[],[]))
    monkeypatch.setattr(output_provenance,'stamp_output',lambda b,f:b)
    class Stored(Exception): pass
    def upload(owner,sid,file,data,mime):
        assert data==original
        assert Document(io.BytesIO(data)).paragraphs[0].text=='Original content'
        raise Stored()
    monkeypatch.setitem(sys.modules,'blob',SimpleNamespace(upload_remediated=upload))
    with pytest.raises(Stored):
        handlers._remediate_file_with_policy({'scan_id':'scan','file':'file.docx','source':'local',
            'remediation_impact_policy':{'rule_based':0,'ai':1},'remediation_impact_allowed_rules':[]},{})
