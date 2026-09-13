"""Real Word correction persists structural positions without inventing pages."""
import re
import zipfile
from docx import Document
from docx.shared import RGBColor, Pt
from lxml import etree
from remediate_office import remediate_office
from release_reports import _location

W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'


def corrected(tmp_path, document):
    source=tmp_path/'locations.docx';document.save(source);diffs=[]
    path,applied,_=remediate_office(source,ai_enabled=False,diffs=diffs,in_scope=lambda sc: sc in {'1.3.1','1.4.3'})
    assert path and applied
    with zipfile.ZipFile(path) as z: root=etree.fromstring(z.read('word/document.xml'))
    return root,diffs


def test_heading_tables_and_second_run_locations_follow_the_saved_docx(tmp_path):
    doc=Document();doc.add_paragraph('Introduction text')
    doc.add_paragraph('First section',style='Heading 2')
    p=doc.add_paragraph();p.add_run('Ordinary black text ')
    p.add_run('Light text').font.color.rgb=RGBColor.from_string('CCCCCC')
    for heading in ['First table','Second table']:
        table=doc.add_table(rows=2,cols=2)
        table.cell(0,0).text=heading;table.cell(0,1).text='Amount'
        table.cell(1,0).text='Value';table.cell(1,1).text='10'
    root,diffs=corrected(tmp_path,doc)
    header_diffs=[d for d in diffs if '<w:tblHeader/>' in d['after']]
    assert [re.search(r'\[location:([^]]+)\]',d['note']).group(1) for d in header_diffs] == ['word:table:1:row:1','word:table:2:row:1']
    tables=list(root.iter(f'{{{W}}}tbl'))
    assert all(table.find(f'{{{W}}}tr/{{{W}}}trPr/{{{W}}}tblHeader') is not None for table in tables)
    heading=next(d for d in diffs if d['after']=='top heading promoted to Heading 1')
    assert heading['note'].endswith('[location:word:p:2]')
    paragraphs=list(root.iter(f'{{{W}}}p'))
    assert paragraphs[1].find(f'{{{W}}}pPr/{{{W}}}pStyle').get(f'{{{W}}}val') == 'Heading1'
    contrast=next(d for d in diffs if d['rule_id']=='1.4.3')
    assert contrast['note'].endswith('[location:word:p:3:run:2]')
    runs=list(paragraphs[2].iter(f'{{{W}}}r'))
    assert runs[1].find(f'{{{W}}}rPr/{{{W}}}color').get(f'{{{W}}}val') != 'CCCCCC'
    assert all('page' not in d['note'].lower() for d in [heading,contrast,*header_diffs])
    # Feed exactly the persisted before/after/note shape into the actual report consumer.
    assert [_location(d) for d in header_diffs] == ['Table 1; Row 1','Table 2; Row 1']
    assert _location(heading) == 'Paragraph 2'
    assert _location(contrast) == 'Paragraph 3; Run 2'


def test_document_level_heading_normalization_is_not_misattributed_to_one_paragraph(tmp_path):
    doc=Document();doc.add_paragraph('Title',style='Heading 1');doc.add_paragraph('Section',style='Heading 1')
    root,diffs=corrected(tmp_path,doc)
    diff=next(d for d in diffs if 'demoted 1 to Heading 2' in d['after'])
    assert diff['note'].endswith('[location:word:document:outline]')
    assert _location(diff) == 'Document heading outline'
    assert list(root.iter(f'{{{W}}}p'))[1].find(f'{{{W}}}pPr/{{{W}}}pStyle').get(f'{{{W}}}val') == 'Heading2'


def test_pseudoheading_has_exact_global_paragraph_location(tmp_path):
    doc=Document();doc.add_paragraph('Body paragraph with normal sized content')
    p=doc.add_paragraph();run=p.add_run('Section Overview');run.bold=True;run.font.size=Pt(22)
    root,diffs=corrected(tmp_path,doc)
    diff=next(d for d in diffs if 'promoted to Heading' in d['after'])
    assert diff['note'].endswith('[location:word:p:2]')
    assert _location(diff) == 'Paragraph 2'
    assert list(root.iter(f'{{{W}}}p'))[1].find(f'{{{W}}}pPr/{{{W}}}pStyle') is not None


def test_nested_table_location_uses_its_own_row_not_the_outer_table_row(tmp_path):
    doc=Document();outer=doc.add_table(rows=2,cols=1)
    outer.cell(0,0).text='Outer heading';outer.cell(1,0).text='Outer value'
    inner=outer.cell(1,0).add_table(rows=2,cols=1)
    inner.cell(0,0).text='Inner heading';inner.cell(1,0).text='Inner value'
    root,diffs=corrected(tmp_path,doc)
    headers=[d for d in diffs if '<w:tblHeader/>' in d['after']]
    assert [_location(d) for d in headers] == ['Table 1; Row 1','Table 2; Row 1']
    tables=list(root.iter(f'{{{W}}}tbl'))
    assert len(tables)==2
    assert all(table.find(f'{{{W}}}tr/{{{W}}}trPr/{{{W}}}tblHeader') is not None for table in tables)


def test_literal_location_text_in_a_run_is_not_reported_as_a_structural_target(tmp_path):
    doc=Document();p=doc.add_paragraph();p.add_run('[location:word:p:999]').font.color.rgb=RGBColor.from_string('CCCCCC')
    _,diffs=corrected(tmp_path,doc)
    contrast=next(d for d in diffs if d['rule_id']=='1.4.3')
    assert '[location:word:p:999]' in contrast['note']
    assert contrast['note'].endswith('[location:word:p:1:run:1]')
    assert _location(contrast)=='Paragraph 1; Run 1'
