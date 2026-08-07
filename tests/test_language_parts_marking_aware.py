"""3.1.2 Language of Parts must be able to tell a marked passage from an unmarked one.

Until this landed, `textchecks.detect_language_parts` read extracted TEXT and nothing else.
Two consequences, both bad and both invisible while the suite was green:

  * It could never certify a pass — a perfectly marked multilingual document failed.
  * NO write could ever clear it. apply_text_values writes w:lang on the runs of an approved
    passage, the criterion kept firing, and `_apply_one_value_kind` therefore credited
    nothing: the value reached the document and the reviewer's approval still went nowhere.

The second is what makes this a remediation bug rather than a detector nicety, and it is why
the round-trip assertions below matter more than the unit ones.
"""
from __future__ import annotations
import io
import sys
import zipfile
from pathlib import Path

import pytest

ACP = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ACP / "api"))

import textchecks as tc  # noqa: E402
import office_structure as osx  # noqa: E402
from apply_text_values import apply_language_parts  # noqa: E402

FR = ("Nous avons publié le rapport trimestriel détaillé pour toutes les régions "
      "concernées par cette réorganisation, et les résultats financiers confirment la tendance.")
EN = ("The quarterly report has now been published for every region in which we currently "
      "operate, and the financial results confirm the trend we observed last year.")

pytestmark = pytest.mark.skipif(not tc._langdetect_available(),
                                reason="langdetect not installed")


def _docx(paragraphs, *, lang_on_every_run: str | None = None) -> bytes:
    docx = pytest.importorskip("docx")
    d = docx.Document()
    for t in paragraphs:
        d.add_paragraph(t)
    buf = io.BytesIO()
    d.save(buf)
    data = buf.getvalue()
    if lang_on_every_run is None:
        return data
    # Word and PowerPoint really do stamp the default language on nearly every run; this
    # reproduces that so the "marked" signal can be shown not to be fooled by it.
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        entries = {n: z.read(n) for n in z.namelist()}
    xml = entries["word/document.xml"].decode()
    xml = xml.replace("<w:r>", f'<w:r><w:rPr><w:lang w:val="{lang_on_every_run}"/></w:rPr>')
    entries["word/document.xml"] = xml.encode()
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        for n, v in entries.items():
            z.writestr(n, v)
    return out.getvalue()


def _text(data: bytes) -> str:
    docx = pytest.importorskip("docx")
    return "\n".join(p.text for p in docx.Document(io.BytesIO(data)).paragraphs)


def _scs(data: bytes, tmp_path) -> list[str]:
    """What the residual re-scan sees: extracted text plus the file's own language marks."""
    p = tmp_path / "f.docx"
    p.write_bytes(data)
    marked = osx.language_marked_spans(p, ".docx")
    return sorted({f["wcag"].split()[0] for f in tc.content_findings(_text(data), marked)})


# ── the detector ──────────────────────────────────────────────────────────────

def test_unmarked_foreign_passage_still_fails(tmp_path):
    assert "3.1.2" in _scs(_docx([EN, FR, EN, FR]), tmp_path)


def test_the_default_language_stamped_everywhere_does_not_excuse_the_foreign_passage(tmp_path):
    """The regression that would silently retire this check.

    Marking the PRIMARY language everywhere is what Word does by default. If marked passages
    were simply dropped before the "does this document mix languages" test, the only language
    left would be the unmarked French -- one language, no finding, and the detector quietly
    stops working on precisely the documents it exists to catch.
    """
    data = _docx([EN, FR, EN, FR], lang_on_every_run="en-US")
    assert osx.language_marked_spans(tmp_path / "x", ".docx") == {}     # nothing read yet
    assert "3.1.2" in _scs(data, tmp_path)


def test_a_passage_marked_as_its_own_language_passes(tmp_path):
    marked = {"fr": FR}
    assert not [f for f in tc.detect_language_parts("\n".join([EN, FR, EN, FR]), marked)]


def test_a_single_language_document_is_never_flagged(tmp_path):
    assert "3.1.2" not in _scs(_docx([EN, EN, EN]), tmp_path)


def test_omitting_the_marks_preserves_the_previous_behaviour():
    """Callers that cannot supply markup (PDF, xlsx, plain text) must be unaffected."""
    text = "\n".join([EN, FR, EN, FR])
    assert tc.detect_language_parts(text) == tc.detect_language_parts(text, {})
    assert tc.detect_language_parts(text)


# ── the round trip that decides whether an approval earns credit ──────────────

def test_writing_the_approved_mark_clears_the_criterion(tmp_path):
    """The whole point: write -> re-scan -> criterion gone -> the lane can credit it."""
    data = _docx([EN, FR, EN, FR])
    assert "3.1.2" in _scs(data, tmp_path)
    out, applied, unresolved = apply_language_parts(data, "docx", {FR[:60]: "fr"})
    assert applied and not unresolved
    assert "3.1.2" not in _scs(out, tmp_path)


def test_marking_the_wrong_language_does_not_clear_it(tmp_path):
    """A mark only satisfies the criterion for the language the passage actually is."""
    data = _docx([EN, FR, EN, FR])
    out, applied, _ = apply_language_parts(data, "docx", {FR[:60]: "de"})
    assert applied                                   # the write happened
    assert "3.1.2" in _scs(out, tmp_path)            # and did not excuse the French


# ── the format that structurally cannot carry the answer ─────────────────────

def test_xlsx_has_no_language_marks_to_read(tmp_path):
    """SpreadsheetML's rich-text run properties (CT_RPrElt) have no language element, so
    there is nowhere to record this and no write can ever clear 3.1.2 there. Absent by
    construction, asserted so a future port does not quietly assume otherwise."""
    openpyxl = pytest.importorskip("openpyxl")
    from openpyxl.cell.text import InlineFont
    assert not [e for e in InlineFont.__elements__ if "lang" in e.lower()]
    wb = openpyxl.Workbook()
    wb.active["A1"] = FR
    p = tmp_path / "f.xlsx"
    wb.save(p)
    assert osx.language_marked_spans(p, ".xlsx") == {}
