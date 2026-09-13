import io
import zipfile
import pytest
from PIL import Image
from docx import Document
from lxml import etree as ET
from office_visible_image import visible_word_image, NS


def _document(crop='19861', flip=False, shared=False):
    image = Image.new('RGB', (100, 100), 'white')
    for y in range(19):
        for x in range(100):
            image.putpixel((x, y), (255, 0, 0))
    raster = io.BytesIO()
    image.save(raster, format='PNG')
    doc = Document()
    doc.add_picture(io.BytesIO(raster.getvalue()))
    if shared:
        doc.add_picture(io.BytesIO(raster.getvalue()))
    output = io.BytesIO()
    doc.save(output)
    changed = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(output.getvalue())) as source, zipfile.ZipFile(changed, 'w') as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == 'word/document.xml':
                root = ET.fromstring(data)
                fill = root.xpath('.//a:blip', namespaces=NS)[0].getparent()
                ET.SubElement(fill, '{'+NS['a']+'}srcRect', t=crop)
                if flip:
                    root.xpath('.//a:xfrm', namespaces=NS)[0].set('flipH', '1')
                data = ET.tostring(root)
            target.writestr(entry, data)
    return changed.getvalue()


def test_visible_crop_excludes_hidden_pixels_and_binds_identity():
    original = _document()
    visible = visible_word_image(original, 'image 1')
    assert visible['crop'] == {'l':0, 't':19861, 'r':0, 'b':0}
    image = Image.open(io.BytesIO(visible['image_bytes']))
    assert image.size == (100,80)
    assert image.getpixel((0,0)) == (255,255,255)
    assert len(visible['source_image_sha256']) == 64
    assert len(visible['visible_image_sha256']) == 64
    assert visible_word_image(original, 'image 1') == visible
    assert visible_word_image(_document('30000'), 'image 1')['visible_image_sha256'] != visible['visible_image_sha256']


@pytest.mark.parametrize('options', [{'crop':'-1'}, {'crop':'100000'}, {'crop':'oops'}, {'flip':True}, {'shared':True}])
def test_ambiguous_or_unsupported_placements_do_not_supply_crop_evidence(options):
    assert visible_word_image(_document(**options), 'image 1') is None


def test_proposer_ocr_and_thumbnail_receive_visible_crop(tmp_path, monkeypatch):
    import ocr
    import proposals
    path = tmp_path / 'cropped.docx'
    path.write_bytes(_document())
    monkeypatch.setattr(ocr, 'is_available', lambda: True)
    monkeypatch.setattr(proposals, 'criteria_enabled', lambda sc: sc == '1.4.5')
    seen = []
    def read(image, *args, **kwargs):
        seen.append(Image.open(io.BytesIO(image)).size)
        return 'Visible words only here in this cropped instruction paragraph today safely'
    monkeypatch.setattr(ocr, 'ocr_text', read)
    result = proposals.propose_images_of_text(path, '.docx', ai_enabled=False)
    assert len(result) == 1
    assert all(size == (100,80) for size in seen)
    assert result[0]['visible_crop']['requires_visual_confirmation'] is True
    assert 'Automatic picture replacement remains unavailable' in result[0]['rationale']
    assert 'model_call_id' not in result[0] or not result[0]['model_call_id']


@pytest.mark.parametrize('options', [{'shared':True}, {'crop':'oops'}, {'flip':True}])
def test_unsupported_crop_never_falls_back_to_full_image_draft(tmp_path, monkeypatch, options):
    import ocr
    import proposals
    path = tmp_path / 'unsupported.docx'
    path.write_bytes(_document(**options))
    monkeypatch.setattr(ocr, 'is_available', lambda: True)
    monkeypatch.setattr(ocr, 'ocr_text', lambda *args, **kwargs: pytest.fail('full raster OCR must not run'))
    assert proposals.propose_images_of_text(path, '.docx', ai_enabled=False) == []
