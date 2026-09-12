#!/usr/bin/env python3
"""Build isolated synthetic native-checker controls. Does not open apps or cloud files.

Use the project test environment with python-docx, Pillow, WeasyPrint and pikepdf.
Open output DOCX files manually in Word Review > Check Accessibility. The generated
JSON is a file manifest, not a substitute for actual native checker observations.
"""
from pathlib import Path
import argparse
import io
import json
import runpy
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'api'))


def build_controls(output):
    from docx import Document
    from docx.shared import Inches
    from docx.oxml import OxmlElement
    from PIL import Image, ImageDraw
    from weasyprint import HTML
    import pikepdf
    from apply_alt import apply_alt_text
    from remediate_office import remediate_office
    from remediate_pdf import apply_pdf_approved
    from pdf_structure_repairs import propose_tagged_repairs
    import hashlib

    output = Path(output); output.mkdir(parents=True, exist_ok=True)
    image = Image.new('RGB', (320, 160), 'white'); draw = ImageDraw.Draw(image)
    draw.rectangle((45, 40, 110, 120), fill='#326ba8')
    draw.rectangle((170, 15, 235, 120), fill='#326ba8')
    image.save(output / 'bars.png')
    doc = Document(); doc.core_properties.title = 'Accessibility repair validation'; doc.core_properties.language = 'en-US'
    doc.add_heading('Accessibility repair validation', 0)
    doc.add_paragraph('This document checks a corrected copy with the Word Accessibility Checker.')
    doc.add_heading('Quarterly comparison', 1)
    doc.add_paragraph('The first bar shows a lower value than the second bar.')
    picture = doc.add_picture(str(output / 'bars.png'), width=Inches(3.2))
    locator = 'word/document.xml#' + picture._inline.docPr.get('name')
    doc.add_heading('Summary', 1)
    table = doc.add_table(rows=3, cols=2); table.style = 'Light Shading Accent 1'
    for row, values in zip(table.rows, [['Quarter', 'Result'], ['First', 'Lower'], ['Second', 'Higher']]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    doc.save(output / 'word-missing-alt.docx')
    data = (output / 'word-missing-alt.docx').read_bytes()
    description = 'Two blue bars with the second taller than the first'
    corrected, applied, unresolved = apply_alt_text(data, {locator: description})
    if len(applied) != 1 or unresolved:
        raise RuntimeError('Supported approved-alt writer did not write the control')
    (output / 'word-approved-alt-corrected.docx').write_bytes(corrected)
    picture._inline.docPr.set('title', description)
    doc.save(output / 'word-title-only.docx')
    fixed, changes, skipped = remediate_office(output / 'word-title-only.docx', ai_enabled=False)
    (output / 'word-automatic-corrected.docx').write_bytes(Path(fixed).read_bytes() if fixed else (output / 'word-title-only.docx').read_bytes())
    html = '<html lang="en"><head><title>Accessible PDF control</title></head><body><h1>Accessible PDF control</h1><p>This tagged document checks the independent PDF validator.</p></body></html>'
    HTML(string=html).write_pdf(output / 'pdf-good-control.pdf', pdf_variant='pdf/ua-1')
    HTML(string=html).write_pdf(output / 'pdf-bad-untagged-control.pdf')
    # Reuse the exact preservation fixture tested by the writer rather than inventing
    # another text-to-tag association for independent validator comparison.
    fixture = runpy.run_path(str(ROOT / 'tests/test_pdf_structure_repairs.py'))['fixture']
    original = fixture(); (output / 'pdf-tagged-original.pdf').write_bytes(original)
    with pikepdf.open(io.BytesIO(original)) as pdf:
        plans = propose_tagged_repairs(pdf, [('Introduction', 0, 20)])
    corrected, applied, unresolved = apply_pdf_approved(original, {p['locator']: p['proposed_value'] for p in plans})
    if len(applied) != 3 or unresolved:
        raise RuntimeError('Supported tagged PDF writers did not write the controls')
    (output / 'pdf-tagged-corrected.pdf').write_bytes(corrected)
    files = [output / name for name in ('word-missing-alt.docx', 'word-approved-alt-corrected.docx',
        'word-title-only.docx', 'word-automatic-corrected.docx', 'pdf-good-control.pdf',
        'pdf-bad-untagged-control.pdf', 'pdf-tagged-original.pdf', 'pdf-tagged-corrected.pdf')]
    manifest = {'scope': 'Synthetic controls; native checker must actually be run',
                'word_automatic_changes': changes, 'word_automatic_skips': skipped,
                'files': [{'file': p.name, 'sha256': hashlib.sha256(p.read_bytes()).hexdigest()} for p in files]}
    (output / 'control-manifest.json').write_text(json.dumps(manifest, indent=2) + '\n')
    return manifest


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('output_directory')
    args = parser.parse_args()
    print(json.dumps(build_controls(args.output_directory), indent=2))
