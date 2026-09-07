#!/usr/bin/env python3
"""Generate (or --check) config/vpat-2.5rev.json — the ITI VPAT® template's HEADINGS only.

WHAT IS REPRODUCED HERE, AND WHY IT IS ONLY THIS. The owner's decision on 2026-09-07, recorded in
`_meta.reproduction_scope` and in ADR 0053: the template's section headings, table titles and
column headers may be reproduced in this repository; its instructional and boilerplate prose may
not. That prose is the Essential Requirements, Best Practices, Terms and Legal Disclaimer sections,
and none of it is in the catalog this script writes — the extractor takes paragraph text only from
paragraphs styled as headings, and cell text only from each table's first row.

The decision answers ADR 0053's Q2 (copyright, reproduction). It does NOT answer Q1 (the service
mark), so nothing here licenses calling ACP's output a VPAT®, and no renderer consuming this
catalog may use the mark. What this buys is the SHAPE: a reader who knows the template can find
their way around ACP's report, which is what acceptance row 13 asks for.

  scripts/gen_wcag_catalog.py         vendors config/wcag22-source.html
  scripts/gen_section_508_catalog.py  vendors config/section-508-source.xml
  scripts/gen_en_301_549_catalog.py   VENDORS NOTHING (the standard's text may not be reproduced)
  this one                            VENDORS NOTHING (the template file may not be redistributed)

So `--check` verifies internal consistency rather than re-parsing a committed copy, exactly as the
EN 301 549 generator does, and with the same weakness stated the same way: it cannot detect a
heading that was wrong at generation time. Regenerating fetches.

SOURCE. The four editions are published separately, each as a .docx, linked from

  https://www.itic.org/policy/accessibility/vpat

Their asset URLs are opaque GUIDs that change when ITI republishes, so they live in EDITIONS below
and are part of what a regeneration re-checks. A 404 there means ITI has published a new revision
and this catalog describes a template that is no longer current — which is the drift this file
exists to make visible, and it is a failure worth reading rather than routing around.

WHY A CATALOG RATHER THAN CONSTANTS IN THE RENDERER. Three reasons, in increasing order of how
much they cost when ignored. The headings differ per edition in ways that are easy to guess wrong
(the 508 edition's WCAG section is headed `WCAG 2.0 Report`, not `WCAG 2.x Report`, because that
edition incorporates WCAG 2.0). The provenance — which revision, fetched when, from where — is
what ADR 0053's Option B asks for and Option C benefits from equally. And a future VPAT 2.6 shows
up as a diff in one JSON file rather than as a renderer nobody re-read.

  python scripts/gen_vpat_structure.py            # regenerate (fetches all four editions)
  python scripts/gen_vpat_structure.py --check    # CI: verify the committed catalog; no fetch
  python scripts/gen_vpat_structure.py --docx-dir <d>   # parse local copies instead of fetching
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import sys
import urllib.request
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "config" / "vpat-2.5rev.json"

PAGE = "https://www.itic.org/policy/accessibility/vpat"
ASSET = "https://www.itic.org/dotAsset/{}.docx"

# Edition -> ITI asset GUID, read off PAGE on 2026-09-07 (all four dated "April 2025" there).
EDITIONS = {
    "VPAT 2.5Rev WCAG": "67270ffc-91e8-4812-9481-8067621f24fd",
    "VPAT 2.5Rev 508": "c9497d16-4684-480c-bae2-99256283731c",
    "VPAT 2.5Rev EU": "3c3891f7-0f54-4b70-a0b6-3f1e3fbf034c",
    "VPAT 2.5Rev INT": "2434a080-87fe-4db1-815e-1e032bf7ac09",
}

# A report section is one of these three, keyed by the requirement set it obliges. The heading text
# is the template's, per edition; the key is ACP's.
SECTION_SETS = {
    "WCAG": "wcag-2.2-aa",
    "Revised Section 508": "section-508",
    "EN 301 549": "en-301-549",
}

# Headings that introduce ACP's own content or ITI's prose, not a requirement table. Recorded so a
# reader of the catalog can see what was deliberately dropped rather than wonder what was missed.
NON_TABLE_HEADINGS = (
    "About This Document",
    "Essential Requirements and Best Practices for Information & Communications Technology "
    "(ICT) Vendors",
    "Getting Started",
    "Essential Requirements for Authors",
    "Best Practices for Authors",
    "Posting the Final Document",
    "Table Information for VPAT® Readers",
    "Terms",
    "Legal Disclaimer (Company)",
)

_REPORT = re.compile(r"^(.*?) Report$")


def _fetch(edition: str, dest: Path) -> Path:
    url = ASSET.format(EDITIONS[edition])
    with urllib.request.urlopen(url, timeout=120) as r:
        if r.status != 200:                                    # pragma: no cover - network
            raise SystemExit(f"{edition}: {url} answered {r.status}")
        dest.write_bytes(r.read())
    return dest


def _headings(path: Path) -> list[tuple[int, str]]:
    """(level, text) for every paragraph styled as a heading, in document order."""
    import docx

    out = []
    for p in docx.Document(str(path)).paragraphs:
        name = p.style.name
        if name.startswith("Heading") and p.text.strip():
            out.append((int(name.rsplit(" ", 1)[1]), " ".join(p.text.split())))
    return out


def _columns(path: Path) -> list[list[str]]:
    """Header row of every table, duplicates from merged cells removed."""
    import docx

    out = []
    for t in docx.Document(str(path)).tables:
        seen, cols = set(), []
        for c in t.rows[0].cells:
            x = " ".join(c.text.split())
            if x and x not in seen:
                seen.add(x)
                cols.append(x)
        out.append(cols)
    return out


def _edition_structure(path: Path) -> dict:
    """Report sections and their sub-headings, from the headings alone."""
    sections, current = [], None
    for level, text in _headings(path):
        if level == 2:
            m = _REPORT.match(text)
            if m:
                key = next((k for k in SECTION_SETS if m.group(1).startswith(k)), None)
                if key is None:
                    raise SystemExit(f"unrecognised report section {text!r} in {path.name}")
                current = {"heading": text, "requirement_set": SECTION_SETS[key],
                           "subsections": []}
                sections.append(current)
            else:
                current = None
        elif level == 3 and current is not None:
            current["subsections"].append(text)
    return {"report_sections": sections}


def build(docx_dir: Path | None) -> dict:
    editions, columns = {}, []
    for name in EDITIONS:
        if docx_dir:
            path = docx_dir / f"{name.replace(' ', '-')}.docx"
            if not path.exists():
                raise SystemExit(f"{path} not found")
        else:
            path = Path(f"/tmp/{name.replace(' ', '-')}.docx")
            _fetch(name, path)
        editions[name] = _edition_structure(path)
        columns.append(_columns(path))

    criteria_cols = {tuple(c) for table in columns for c in table
                     if c and c[0] == "Criteria"}
    if len(criteria_cols) != 1:
        raise SystemExit(f"criteria tables disagree on columns: {criteria_cols}")
    standards_cols = {tuple(c) for table in columns for c in table
                      if c and c[0] == "Standard/Guideline"}
    if len(standards_cols) != 1:
        raise SystemExit(f"standards tables disagree on columns: {standards_cols}")

    return {
        "_meta": {
            "template": "ITI VPAT® 2.5Rev",
            "revision": "April 2025",
            "publisher": "Information Technology Industry Council (ITI)",
            "source_page": PAGE,
            "retrieved": dt.date.today().isoformat(),
            "generated_by": "scripts/gen_vpat_structure.py",
            "reproduction_scope": (
                "Section HEADINGS, table titles and column headers only. The owner decided on "
                "2026-09-07 (ADR 0053, Q2) that these may be reproduced in this repository and "
                "that the template's instructional and boilerplate prose may not. No such prose "
                "is present in this file, and the template .docx is not vendored. ADR 0053's Q1 "
                "— whether output may be CALLED a VPAT(R) — is NOT answered by that decision, so "
                "no renderer consuming this catalog may use the mark."
            ),
            "not_reproduced": list(NON_TABLE_HEADINGS),
        },
        "criteria_table_columns": list(next(iter(criteria_cols))),
        "standards_table_columns": list(next(iter(standards_cols))),
        "editions": editions,
    }


def check() -> int:
    if not OUT.exists():
        print(f"{OUT} is missing", file=sys.stderr)
        return 1
    d = json.loads(OUT.read_text())
    problems: list[str] = []

    if sorted(d["editions"]) != sorted(EDITIONS):
        problems.append(f"editions {sorted(d['editions'])} != {sorted(EDITIONS)}")

    if d["criteria_table_columns"] != ["Criteria", "Conformance Level", "Remarks and Explanations"]:
        problems.append(f"criteria columns changed: {d['criteria_table_columns']}")
    if d["standards_table_columns"] != ["Standard/Guideline", "Included In Report"]:
        problems.append(f"standards columns changed: {d['standards_table_columns']}")

    # The point of the catalog: each edition carries exactly the requirement sets it obliges, and
    # they are the same sets acr_catalog builds a matrix from.
    expected = {
        "VPAT 2.5Rev WCAG": ["wcag-2.2-aa"],
        "VPAT 2.5Rev 508": ["wcag-2.2-aa", "section-508"],
        "VPAT 2.5Rev EU": ["wcag-2.2-aa", "en-301-549"],
        "VPAT 2.5Rev INT": ["wcag-2.2-aa", "section-508", "en-301-549"],
    }
    for name, want in expected.items():
        got = [s["requirement_set"] for s in d["editions"].get(name, {}).get("report_sections", [])]
        if got != want:
            problems.append(f"{name}: requirement sets {got} != {want}")

    for name, ed in d["editions"].items():
        for sec in ed["report_sections"]:
            if not sec["heading"].endswith(" Report"):
                problems.append(f"{name}: section heading {sec['heading']!r} is not a report")
            if sec["requirement_set"] == "wcag-2.2-aa" and len(sec["subsections"]) != 3:
                problems.append(f"{name}: WCAG section has {len(sec['subsections'])} level tables")
            for sub in sec["subsections"]:
                # A heading is a title. Prose is what the decision does not permit, and the two are
                # told apart the same way the EN generator tells them apart: by length and by a
                # sentence's full stop.
                if len(sub) > 80 or sub.rstrip().endswith("."):
                    problems.append(f"{name}: {sub!r} reads as prose, not a heading")

    for banned in NON_TABLE_HEADINGS:
        if banned in json.dumps(d["editions"]):
            problems.append(f"non-table heading {banned!r} leaked into the editions")

    if problems:
        for p in problems:
            print(f"vpat-2.5rev: {p}", file=sys.stderr)
        return 1
    total = sum(len(s["subsections"]) for e in d["editions"].values()
                for s in e["report_sections"])
    print(f"vpat-2.5rev structure OK — {len(d['editions'])} editions, "
          f"{total} table headings, no prose")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--check", action="store_true", help="verify the committed catalog; no fetch")
    ap.add_argument("--docx-dir", type=Path, help="parse local .docx copies instead of fetching")
    args = ap.parse_args()

    if args.check:
        return check()

    d = build(args.docx_dir)
    OUT.write_text(json.dumps(d, indent=2, ensure_ascii=False) + "\n")
    print(f"wrote {OUT.relative_to(ROOT)}")
    return check()


if __name__ == "__main__":
    raise SystemExit(main())
